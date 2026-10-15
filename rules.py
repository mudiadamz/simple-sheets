"""Data validation (drop-down lists, rejecting invalid input) and conditional formatting (colour rules,
colour scales, data bars) read from .xlsx files. The rules themselves stay in the file untouched; ranges
are kept in on-disk coordinates and looked up through Sheet.orig_pos, so inserted rows simply have none."""
import math

from formatting import color_hex, parse_date, parse_number
from formula import FormulaError, Parser, Range, XLError, serial, shift

OPS = {
    "between": lambda v, a, b: min(a, b) <= v <= max(a, b),
    "notBetween": lambda v, a, b: not min(a, b) <= v <= max(a, b),
    "equal": lambda v, a, b: v == a, "notEqual": lambda v, a, b: v != a,
    "greaterThan": lambda v, a, b: v > a, "lessThan": lambda v, a, b: v < a,
    "greaterThanOrEqual": lambda v, a, b: v >= a, "lessThanOrEqual": lambda v, a, b: v <= a,
}
OP_WORDS = {"between": "between {a} and {b}", "notBetween": "not between {a} and {b}", "equal": "equal to {a}",
            "notEqual": "not equal to {a}", "greaterThan": "greater than {a}", "lessThan": "less than {a}",
            "greaterThanOrEqual": "at least {a}", "lessThanOrEqual": "at most {a}"}


def _ranges(sqref):
    return [(r.min_row - 1, r.min_col - 1, r.max_row - 1, r.max_col - 1) for r in sqref.ranges]


def _inside(ranges, o):
    r, c = o
    return r is not None and c is not None and any(r1 <= r <= r2 and c1 <= c <= c2 for r1, c1, r2, c2 in ranges)


# ---------- reading ----------

def read_validations(ws):
    out = []
    for i, dv in enumerate(ws.data_validations.dataValidation):
        if not dv.type or not dv.sqref:
            continue
        out.append({"id": i, "ranges": _ranges(dv.sqref), "type": dv.type, "op": dv.operator or "between",
                    "f1": dv.formula1, "f2": dv.formula2, "arrow": not dv.showDropDown,   # Excel's flag is inverted
                    "show_error": bool(dv.showErrorMessage), "style": dv.errorStyle or "stop",
                    "title": dv.errorTitle or "", "error": dv.error or ""})
    return out


def _color(c, theme):
    if c is None or (c.type == "rgb" and c.rgb in ("00000000", None)):
        return None
    return color_hex(c, theme)


def _dxf_style(dxf, theme):
    st = {}
    if dxf is None:
        return st
    f = dxf.font
    if f is not None:
        for attr, name in (("b", "bold"), ("i", "italic"), ("u", "underline"), ("strike", "strike")):
            if getattr(f, attr, None):
                st[name] = True
        fg = _color(f.color, theme)
        if fg:
            st["fg"] = fg
    fill = dxf.fill
    if fill is not None and hasattr(fill, "bgColor"):
        bg = _color(fill.bgColor, theme) or _color(fill.fgColor, theme)   # dxf solid fills use bgColor
        if bg:
            st["bg"] = bg
    return st


def read_conditional(ws, theme):
    out, n = [], 0
    for cf in ws.conditional_formatting:
        ranges = _ranges(cf.sqref)
        for rule in cf.rules:
            n += 1
            if not ranges:
                continue
            d = {"id": n - 1, "ranges": ranges, "anchor": ranges[0][:2], "type": rule.type, "op": rule.operator,
                 "f": list(rule.formula or []), "priority": rule.priority or 0, "stop": bool(rule.stopIfTrue),
                 "style": _dxf_style(rule.dxf, theme), "text": rule.text, "rank": rule.rank,
                 "percent": bool(rule.percent), "bottom": bool(rule.bottom),
                 "above": rule.aboveAverage is not False, "equal": bool(rule.equalAverage)}
            if rule.colorScale is not None:
                cs = rule.colorScale
                d["scale"] = [(v.type, v.val, _color(col, theme) or "#ffffff") for v, col in zip(cs.cfvo, cs.color)]
            if rule.dataBar is not None:
                db = rule.dataBar
                d["bar"] = [(v.type, v.val) for v in db.cfvo], _color(db.color, theme) or "#638ec6"
            out.append(d)
    out.sort(key=lambda d: d["priority"])
    return out


# ---------- shared helpers ----------

class Rules:
    """Per-app helper. Range statistics for conditional formats are cached until formulas recalculate."""

    def __init__(self, app):
        self.app = app
        self.stats, self.memo = {}, None

    @property
    def eng(self):
        return self.app.engine

    def name(self, sheet):
        return self.eng.ids.get(id(sheet))

    def evaluate(self, sheet, text, dr=0, dc=0):
        """A rule formula (no '=') evaluated for the cell (dr, dc) away from the rule's top-left."""
        name = self.name(sheet)
        if name is None or text is None:
            return None
        text = str(text)
        if dr or dc:
            try:
                text = shift("=" + text, dr, dc)[1:]
            except ValueError:
                return None
        try:
            v = self.eng.ev(Parser(text).parse(), name)
        except FormulaError as e:
            return e.err
        except Exception:
            return None
        return v

    def number(self, sheet, text, dr=0, dc=0):
        v = self.evaluate(sheet, text, dr, dc)
        if isinstance(v, Range):
            v = v.at(0, 0)
        if isinstance(v, bool):
            return float(v)
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, str):
            n = parse_number(v)
            if n is not None:
                return n
            d = parse_date(v)
            if d is not None:
                return serial(d)
        return None

    def model_cells(self, sheet, ranges, model=False):
        """Model (r, c) cells inside on-disk ranges (or model ranges), clipped to the data."""
        rows = sheet.rows
        maxc = max((len(r) for r in rows), default=0)
        inv_r = (lambda i: i) if model else sheet._inverse(sheet.rowmap, sheet.rowtail)
        inv_c = (lambda i: i) if model else sheet._inverse(sheet.colmap, sheet.coltail)
        out = []
        for r1, c1, r2, c2 in ranges:
            rs = [inv_r(o) for o in range(r1, min(r2, r1 + len(rows) + len(sheet.rowmap)) + 1)]
            cs = [inv_c(o) for o in range(c1, min(c2, c1 + maxc + len(sheet.colmap)) + 1)]
            out += [(r, c) for r in rs if r is not None and r < len(rows)
                    for c in cs if c is not None and c < len(rows[r])]
        return out

    def cell(self, sheet, r, c):
        name = self.name(sheet)
        return self.eng.cell(name, r, c) if name is not None else None

    # ---------- data validation ----------

    def validation_at(self, sheet, r, c):
        rules = getattr(sheet, "validations", None)
        if not rules:
            return None
        o = sheet.orig_pos(r, c)
        return next((d for d in rules if _inside(d["ranges"], (r, c) if d.get("model") else o)), None)

    def list_options(self, sheet, d):
        f = (d.get("f1") or "").strip()
        if f.startswith('"') and f.endswith('"'):
            return [x.strip() for x in f[1:-1].split(",") if x.strip()]
        v = self.evaluate(sheet, f)
        vals = list(v.flat()) if isinstance(v, Range) else [v]
        out = []
        for x in vals:
            if x is None or isinstance(x, XLError):
                continue
            t = (str(int(x)) if isinstance(x, float) and x.is_integer() else str(x)).strip()
            if t and t not in out:
                out.append(t)
        return out

    def dropdown(self, sheet, r, c):
        """Options for a cell's drop-down arrow, or None."""
        d = self.validation_at(sheet, r, c)
        if d and d["type"] == "list" and d["arrow"]:
            return self.list_options(sheet, d)
        return None

    def check(self, sheet, r, c, text):
        """None if `text` may go into the cell, else (style, title, message)."""
        d = self.validation_at(sheet, r, c)
        if not d or not d["show_error"] or text == "" or text.startswith("="):
            return None
        ok, why = self._valid(sheet, d, text)
        if ok:
            return None
        return (d["style"], d["title"] or "Invalid value",
                d["error"] or f"“{text}” isn't allowed here: {why}.")

    def _valid(self, sheet, d, text):
        kind = d["type"]
        if kind == "list":
            opts = self.list_options(sheet, d)
            n = parse_number(text)
            ok = any(o.lower() == text.strip().lower() or (n is not None and parse_number(o) == n) for o in opts)
            return ok, "choose one of " + ", ".join(opts[:8]) + ("…" if len(opts) > 8 else "")
        if kind == "custom" or kind == "any":
            return True, ""
        if kind == "textLength":
            v, what = float(len(text)), "the length must be "
        elif kind in ("date", "time"):
            v = parse_number(text)
            if v is None:
                dte = parse_date(text)
                v = serial(dte) if dte is not None else None
            what = "the date must be " if kind == "date" else "the time must be "
            if v is None:
                return False, "it must be a " + kind
        else:
            v = parse_number(text)
            what = "the number must be "
            if v is None:
                return False, "it must be a number"
            if kind == "whole" and not float(v).is_integer():
                return False, "it must be a whole number"
        a, b = self.number(sheet, d["f1"]), self.number(sheet, d["f2"])
        op = d["op"]
        if a is None or op not in OPS or (op in ("between", "notBetween") and b is None):
            return True, ""
        fmt = lambda x: str(int(x)) if float(x).is_integer() else f"{x:g}"   # noqa: E731
        return OPS[op](v, a, b), what + OP_WORDS[op].format(a=fmt(a), b=fmt(b) if b is not None else "")

    # ---------- conditional formatting ----------

    def fresh(self):
        if self.memo is not self.eng.memo:      # formulas recalculated: statistics may have changed
            self.memo, self.stats = self.eng.memo, {}

    def numbers(self, sheet, d):
        key = (id(sheet), id(d), "nums")
        if key not in self.stats:
            vals = [self.cell(sheet, r, c) for r, c in self.model_cells(sheet, d["ranges"], d.get("model"))]
            self.stats[key] = sorted(float(v) for v in vals
                                     if isinstance(v, (int, float)) and not isinstance(v, bool))
        return self.stats[key]

    def counts(self, sheet, d):
        key = (id(sheet), id(d), "counts")
        if key not in self.stats:
            counts = {}
            for r, c in self.model_cells(sheet, d["ranges"], d.get("model")):
                v = self.cell(sheet, r, c)
                if v is not None and v != "":
                    k = v.lower() if isinstance(v, str) else v
                    counts[k] = counts.get(k, 0) + 1
            self.stats[key] = counts
        return self.stats[key]

    def threshold(self, sheet, d, kind, val, nums):
        """A colour-scale / data-bar point (min, max, number, percent, percentile, formula) -> number."""
        if not nums:
            return None
        if kind == "min":
            return nums[0]
        if kind == "max":
            return nums[-1]
        x = self.number(sheet, val) if val is not None else None
        if x is None:
            return None
        if kind == "percent":
            return nums[0] + (nums[-1] - nums[0]) * x / 100
        if kind == "percentile":
            pos = x / 100 * (len(nums) - 1)
            lo = math.floor(pos)
            hi = min(lo + 1, len(nums) - 1)
            return nums[lo] + (nums[hi] - nums[lo]) * (pos - lo)
        return x                                   # num / formula

    def conditional(self, sheet, r, c):
        """Extra style for a cell from its colour rules: {'bg', 'fg', 'bold', …, 'bar': (fraction, colour)}."""
        rules = getattr(sheet, "cond", None)
        if not rules:
            return None
        o = sheet.orig_pos(r, c)
        out = {}
        for d in rules:
            od = (r, c) if d.get("model") else o     # rules made here use the sheet's own positions
            if not _inside(d["ranges"], od):
                continue
            self.fresh()
            hit = self._apply(sheet, d, r, c, od)
            if hit is None:
                continue
            for k, v in hit.items():
                out.setdefault(k, v)
            if d["stop"]:
                break
        return out or None

    def _apply(self, sheet, d, r, c, o):
        t = d["type"]
        v = self.cell(sheet, r, c)
        isnum = isinstance(v, (int, float)) and not isinstance(v, bool)
        if t == "colorScale" and d.get("scale"):
            if not isnum:
                return None
            nums = self.numbers(sheet, d)
            pts = [(self.threshold(sheet, d, k, val, nums), col) for k, val, col in d["scale"]]
            pts = [(x, col) for x, col in pts if x is not None]
            if len(pts) < 2:
                return None
            return {"bg": scale_color(float(v), pts)}
        if t == "dataBar" and d.get("bar"):
            if not isnum:
                return None
            nums = self.numbers(sheet, d)
            (k1, v1), (k2, v2) = (d["bar"][0] + [("min", None), ("max", None)])[:2]
            lo, hi = self.threshold(sheet, d, k1, v1, nums), self.threshold(sheet, d, k2, v2, nums)
            if lo is None or hi is None:
                return None
            frac = 1.0 if hi == lo else (float(v) - lo) / (hi - lo)
            return {"bar": (max(0.1, min(1.0, frac)), d["bar"][1])}
        if not self._matches(sheet, d, v, isnum, o):
            return None
        return d["style"]

    def _matches(self, sheet, d, v, isnum, o):
        t, dr, dc = d["type"], o[0] - d["anchor"][0], o[1] - d["anchor"][1]
        txt = "" if v is None else str(v).lower() if not isinstance(v, float) else (
            str(int(v)) if v.is_integer() else f"{v:.15g}")
        if t == "cellIs":
            if d["op"] not in OPS or not d["f"]:
                return False
            a = self.evaluate(sheet, d["f"][0], dr, dc)
            b = self.evaluate(sheet, d["f"][1], dr, dc) if len(d["f"]) > 1 else None
            a, b = (x.at(0, 0) if isinstance(x, Range) else x for x in (a, b))
            if isinstance(a, XLError) or isinstance(v, XLError):
                return False
            if isnum and isinstance(a, (int, float)) and (b is None or isinstance(b, (int, float))):
                return OPS[d["op"]](float(v), float(a), float(b) if b is not None else None)
            if isinstance(a, str) and not isnum:
                return OPS[d["op"]](txt, a.lower(), (b or "").lower() if isinstance(b, str) else b) \
                    if d["op"] in ("equal", "notEqual") else False
            if v is None and isinstance(a, (int, float)):
                return OPS[d["op"]](0.0, float(a), float(b) if isinstance(b, (int, float)) else None)
            return False
        if t == "expression":
            if not d["f"]:
                return False
            res = self.evaluate(sheet, d["f"][0], dr, dc)
            if isinstance(res, Range):
                res = res.at(0, 0)
            return bool(res) and not isinstance(res, XLError) and res != ""
        if t in ("containsText", "notContainsText", "beginsWith", "endsWith"):
            needle = (d["text"] or "").lower()
            return {"containsText": needle in txt, "notContainsText": needle not in txt,
                    "beginsWith": txt.startswith(needle), "endsWith": txt.endswith(needle)}[t]
        if t == "containsBlanks":
            return txt.strip() == ""
        if t == "notContainsBlanks":
            return txt.strip() != ""
        if t == "containsErrors":
            return isinstance(v, XLError)
        if t == "notContainsErrors":
            return not isinstance(v, XLError)
        if t in ("duplicateValues", "uniqueValues"):
            if v is None or v == "":
                return False
            n = self.counts(sheet, d).get(v.lower() if isinstance(v, str) else v, 0)
            return n > 1 if t == "duplicateValues" else n == 1
        if t == "top10":
            if not isnum:
                return False
            nums = self.numbers(sheet, d)
            k = d["rank"] or 10
            if d["percent"]:
                k = max(1, int(len(nums) * k / 100))
            k = min(k, len(nums))
            if not k:
                return False
            return float(v) <= nums[k - 1] if d["bottom"] else float(v) >= nums[-k]
        if t == "aboveAverage":
            if not isnum:
                return False
            nums = self.numbers(sheet, d)
            avg = sum(nums) / len(nums) if nums else 0
            x = float(v)
            if d["above"]:
                return x > avg or (d["equal"] and x == avg)
            return x < avg or (d["equal"] and x == avg)
        return False


def scale_color(x, pts):
    pts = sorted(pts)
    if x <= pts[0][0]:
        return pts[0][1]
    if x >= pts[-1][0]:
        return pts[-1][1]
    for (x1, c1), (x2, c2) in zip(pts, pts[1:]):
        if x1 <= x <= x2:
            t = 0 if x2 == x1 else (x - x1) / (x2 - x1)
            a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
            b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
            return "#%02x%02x%02x" % tuple(round(p + (q - p) * t) for p, q in zip(a, b))
    return pts[-1][1]



# ---------- rules made / removed in the app -> file ----------

def _ref(ranges):
    from formula import col_letters
    return " ".join(f"{col_letters(c1)}{r1 + 1}" + ("" if (r1, c1) == (r2, c2) else f":{col_letters(c2)}{r2 + 1}")
                    for r1, c1, r2, c2 in ranges)


def _argb(hex6):
    return "FF" + hex6[1:].upper()


def build_cf(d):
    """In-app rule dict -> openpyxl Rule."""
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle
    st = d.get("style") or {}
    font = Font(b=st.get("bold") or None, i=st.get("italic") or None,
                color=_argb(st["fg"]) if st.get("fg") else None)
    fill = PatternFill(bgColor=_argb(st["bg"]), fill_type="solid") if st.get("bg") else None
    dxf = DifferentialStyle(font=font, fill=fill)
    t = d["type"]
    if t == "colorScale":
        pts = d["scale"]
        if len(pts) == 3:
            return ColorScaleRule(start_type="min", start_color=_argb(pts[0][2]), mid_type="percentile", mid_value=50,
                                  mid_color=_argb(pts[1][2]), end_type="max", end_color=_argb(pts[2][2]))
        return ColorScaleRule(start_type="min", start_color=_argb(pts[0][2]), end_type="max", end_color=_argb(pts[-1][2]))
    if t == "dataBar":
        return DataBarRule(start_type="min", end_type="max", color=_argb(d["bar"][1]))
    first = _ref(d["ranges"]).split(":")[0].split(" ")[0]
    rule = Rule(type=t, dxf=dxf, stopIfTrue=d.get("stop") or None)
    if t == "cellIs":
        rule.operator = d["op"]
        rule.formula = [f for f in d["f"] if f not in (None, "")]
    elif t == "expression":
        rule.formula = [d["f"][0]]
    elif t in ("containsText", "notContainsText", "beginsWith", "endsWith"):
        rule.operator, rule.text = t, d["text"]
        q = d["text"].replace('"', '""')
        rule.formula = [{"containsText": f'NOT(ISERROR(SEARCH("{q}",{first})))',
                         "notContainsText": f'ISERROR(SEARCH("{q}",{first}))',
                         "beginsWith": f'LEFT({first},LEN("{q}"))="{q}"',
                         "endsWith": f'RIGHT({first},LEN("{q}"))="{q}"'}[t]]
    elif t == "containsBlanks":
        rule.formula = [f"LEN(TRIM({first}))=0"]
    elif t == "notContainsBlanks":
        rule.formula = [f"LEN(TRIM({first}))>0"]
    elif t == "top10":
        rule.rank, rule.bottom, rule.percent = int(d.get("rank") or 10), d.get("bottom") or None, d.get("percent") or None
    elif t == "aboveAverage":
        rule.aboveAverage = None if d.get("above", True) else False
    return rule


def build_dv(d):
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type=d["type"], operator=d.get("op") if d["type"] not in ("list", "custom") else None,
                        formula1=d.get("f1"), formula2=d.get("f2"), allow_blank=True,
                        showDropDown=not d.get("arrow", True), showErrorMessage=d.get("show_error", True),
                        errorStyle=d.get("style", "stop"), error=d.get("error") or None,
                        errorTitle=d.get("title") or None)
    for part in _ref(d["ranges"]).split():
        dv.add(part)
    return dv


def changed(sheet):
    """True if rules were added / removed / edited here since the file was read."""
    return getattr(sheet, "cond_disk", None) is not None and (
        [id(d) for d in sheet.cond] != sheet.cond_disk or [id(d) for d in sheet.validations] != sheet.dv_disk)


def remember(sheet):
    """The rules as they are in the file now (after loading / saving)."""
    for d in sheet.cond + sheet.validations:
        d.pop("model", None)
    sheet.cond_disk = [id(d) for d in sheet.cond]
    sheet.dv_disk = [id(d) for d in sheet.validations]


def save_fn(name, sheet):
    """fn(workbook) writing the sheet's current rules, or None if nothing changed. Rules from the file
    are kept as openpyxl read them (full detail); only new ones are built from the app's dicts."""
    if not changed(sheet):
        return None
    cond, dvs = list(sheet.cond), list(sheet.validations)

    def fn(wb):
        from openpyxl.formatting.formatting import ConditionalFormattingList
        ws = wb[name]
        old = []
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                old.append((cf.sqref, rule))
        new = ConditionalFormattingList()
        prio = 1
        for d in cond:
            if "id" in d and d["id"] < len(old):
                sqref, rule = old[d["id"]]
                rule.priority = prio
                new.add(str(sqref), rule)
            else:
                rule = build_cf(d)
                rule.priority = prio
                new.add(_ref(d["ranges"]), rule)
            prio += 1
        ws.conditional_formatting = new
        old_dv = list(ws.data_validations.dataValidation)
        ws.data_validations.dataValidation = []
        for d in dvs:
            if "id" in d and d["id"] < len(old_dv):
                ws.data_validations.dataValidation.append(old_dv[d["id"]])
            else:
                ws.data_validations.dataValidation.append(build_dv(d))
    return fn


def after_save(sheet):
    """Saved: the rules' order in the file is now the app's order."""
    for i, d in enumerate(sheet.cond):
        d["id"] = i
    for i, d in enumerate(sheet.validations):
        d["id"] = i
    remember(sheet)


def write_all(ws, sheet):
    """A brand-new file: every rule is written from the app's dicts (positions = the sheet's)."""
    prio = 1
    for d in sheet.cond:
        try:
            rule = build_cf(d)
        except Exception:
            continue
        rule.priority = prio
        prio += 1
        ws.conditional_formatting.add(_ref(d["ranges"]), rule)
    for d in sheet.validations:
        try:
            ws.add_data_validation(build_dv(d))
        except Exception:
            pass
