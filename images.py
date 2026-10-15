"""Pictures in a file, drawn over the grid at their anchor cell (PNG / GIF natively; other formats show
as a labelled box). They stay in the file when saving (openpyxl keeps them)."""
import base64
import tkinter as tk


def ratio(want, have, limit=4000):
    """want / have as zoom p / subsample q (small whole numbers, Tk has no free scaling): -> (p, q)."""
    r = want / max(have, 1)
    best = (1, 1)
    for q in range(1, 13):
        p = max(1, round(r * q))
        if p > 12 or have * p > limit:
            continue
        if abs(p / q - r) < abs(best[0] / best[1] - r) - 1e-9:
            best = (p, q)
    return best


def scaled(photo, want_w, want_h):
    """A PhotoImage made close to want_w x want_h pixels (each side on its own)."""
    k = max(1, min(photo.width() // max(want_w, 1), photo.height() // max(want_h, 1)))
    if k > 1:                                   # much bigger: halve / third it first (cheap)
        photo = photo.subsample(k, k)
    (px, qx), (py, qy) = ratio(want_w, photo.width()), ratio(want_h, photo.height())
    if px > 1 or py > 1:
        photo = photo.zoom(px, py)
    if qx > 1 or qy > 1:
        photo = photo.subsample(qx, qy)
    return photo


def _photo(g, im):
    """Tk image for a picture at its size in the sheet (cached)."""
    cache = g.__dict__.setdefault("image_cache", {})
    key = (id(im), round(g.zoom, 2), im.get("width"), im.get("height"))
    if key in cache:
        return cache[key]
    photo = None
    if im.get("data") and im.get("format") in ("png", "gif"):
        try:
            photo = tk.PhotoImage(master=g, data=base64.b64encode(im["data"]))
            want_w = max(g.px(im.get("width") or photo.width()), 1)
            want_h = max(g.px(im.get("height") or photo.height()), 1)
            if (want_w, want_h) != (photo.width(), photo.height()):
                photo = scaled(photo, want_w, want_h)
        except tk.TclError:
            photo = None
    cache[key] = photo
    return photo


def draw_all(g):
    from theme import C
    g.image_boxes, g.image_items = [], []
    images = g.sheet.extra.get("images") if g.sheet else None
    if not images:
        return
    W, H = g.cv.winfo_width(), g.cv.winfo_height()
    for im in images:
        r, c = im["cell"]
        x, y = g.col_x(c), g.row_y(r)
        if x > W or y > H:
            continue
        photo = _photo(g, im)
        if photo is not None:
            g.cv.create_image(x, y, image=photo, anchor="nw", tags="image")
            w, h = photo.width(), photo.height()
        else:
            w, h = g.px(im.get("width") or 120), g.px(im.get("height") or 80)
            g.cv.create_rectangle(x, y, x + w, y + h, fill=C["header_bg"], outline=C["header_line"], tags="image")
            g.cv.create_text(x + w / 2, y + h / 2, text=f"picture ({im.get('format') or '?'})", fill=C["muted"],
                             tags="image")
        g.image_boxes.append((x, y, x + w, y + h))
        g.image_items.append(((x, y, x + w, y + h), im))
        if getattr(g, "selected_image", None) is im:     # picked: dashed frame + resize handle
            g.cv.create_rectangle(x - 2, y - 2, x + w + 2, y + h + 2, outline=C["accent"], width=2, dash=(4, 3),
                                  tags="image")
            g.cv.create_rectangle(x + w - 5, y + h - 5, x + w + 5, y + h + 5, fill=C["accent"],
                                  outline=C["cell_bg"], tags="image")


# ---------- picking, inserting, moving, resizing, deleting ----------

HANDLE = 8   # px around the bottom-right corner that grab the resize handle


def on_corner(box, x, y):
    return abs(x - box[2]) <= HANDLE and abs(y - box[3]) <= HANDLE


def hit(g, x, y):
    """The picture under (x, y) - its resize handle (which sticks out of the corner) counts too."""
    for box, im in reversed(getattr(g, "image_items", [])):
        x1, y1, x2, y2 = box
        if x1 <= x <= x2 and y1 <= y <= y2 or on_corner(box, x, y):
            return im
    return None


def cursor(g, x, y):
    """Mouse pointer over a picture / chart: resize arrows on the corner handle, move cross on the body."""
    for items in (getattr(g, "image_items", []), getattr(g, "chart_boxes", [])):
        for box, _obj in reversed(items):
            if on_corner(box, x, y):
                return "size_nw_se"
            if box[0] <= x <= box[2] and box[1] <= y <= box[3]:
                return "fleur"
    return None


def images_of(sheet):
    return sheet.extra.get("images") or []


def changed(sheet):
    return [id(i) for i in images_of(sheet)] != getattr(sheet, "images_disk", [id(i) for i in images_of(sheet)])


def remember(sheet):
    sheet.images_disk = [id(i) for i in images_of(sheet)]


def replace(app, old, new):
    imgs = [new if i is old else i for i in images_of(app.current)]
    if new is None:
        imgs = [i for i in imgs if i is not None]
    app.apply_mixed([("im", 0, 0, imgs)])
    app.table.selected_image = new
    app.table.redraw()
    return new


def delete(app, im):
    return replace(app, im, None)


def insert(app, path=None, data=None):
    """Insert → Picture: a PNG / GIF / JPEG / BMP file at the active cell, at its own size."""
    import pillowless
    if data is None:
        if path is None:
            from tkinter import filedialog
            path = filedialog.askopenfilename(filetypes=[("Pictures", "*.png *.gif *.jpg *.jpeg *.bmp"),
                                                         ("All files", "*.*")], parent=app)
            if not path:
                return None
        with open(path, "rb") as f:
            data = f.read()
    try:
        fmt, (w, h) = pillowless.sniff(data)
    except OSError:
        from tkinter import messagebox
        messagebox.showwarning("Picture", "That file isn't a picture Simple Sheets can insert (PNG, GIF, JPEG, BMP).",
                               parent=app)
        return None
    g = app.table
    im = {"cell": tuple(g.cur), "end": None, "width": w, "height": h, "format": fmt.lower(), "data": data}
    app.apply_mixed([("im", 0, 0, images_of(app.current) + [im])])
    g.selected_image = im
    g.redraw()
    return im


def nudge(app, im, dr, dc, resize=False):
    """Arrows move a picked picture one cell; Shift+arrows make it 10 % bigger / smaller."""
    if resize:
        k = 1.1 ** (dr + dc)
        return replace(app, im, {**im, "width": max(int(im["width"] * k), 8), "height": max(int(im["height"] * k), 8),
                                 "end": None, "moved": True})
    r, c = im["cell"]
    return replace(app, im, {**im, "cell": (max(r + dr, 0), max(c + dc, 0)), "end": None, "moved": True})


def drop(g, im, mode, x0, y0, x1, y1):
    """End of a mouse drag: 'move' to another cell, or 'resize' (keeps the proportions)."""
    if mode == "resize":
        k = resize_factor(g, im, x1 - x0, y1 - y0)
        return replace(g.app, im, {**im, "width": max(int(im["width"] * k), 8),
                                   "height": max(int(im["height"] * k), 8), "end": None, "moved": True})
    r_from, c_from = g.hit(x0, y0)
    r_to, c_to = g.hit(max(x1, g.hdr_w() + 1), max(y1, g.hdr_h() + 1))
    if min(r_from, c_from, r_to, c_to) < 0 or (r_to, c_to) == (r_from, c_from):
        return None
    return nudge(g.app, im, r_to - r_from, c_to - c_from)


def resize_factor(g, im, dx, dy):
    """Dragging the corner by (dx, dy) screen px: the scale (proportions kept; the bigger move wins)."""
    w, h = max(g.px(im["width"]), 1), max(g.px(im["height"]), 1)
    k = max((w + dx) / w, (h + dy) / h) if dx + dy > 0 else min((w + dx) / w, (h + dy) / h)
    return max(k, 0.05)


def save_fn(name, sheet):
    if not changed(sheet):
        return None
    images = list(images_of(sheet))

    def fn(wb):
        from io import BytesIO

        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.utils import get_column_letter
        ws = wb[name]
        old = list(getattr(ws, "_images", []))
        ws._images = []
        for im in images:
            if "id" in im and im["id"] < len(old):
                img = old[im["id"]]
            elif im.get("data"):
                img = XLImage(BytesIO(im["data"]))
            else:
                continue
            if "id" not in im or im.get("moved"):
                r, c = im["cell"]
                img.anchor = f"{get_column_letter(c + 1)}{r + 1}"
                img.width, img.height = im["width"], im["height"]
            ws._images.append(img)
    return fn


def write_all(ws, sheet):
    """A brand-new file: every picture from its bytes."""
    from io import BytesIO

    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
    for im in images_of(sheet):
        if not im.get("data"):
            continue
        try:
            img = XLImage(BytesIO(im["data"]))
        except Exception:
            continue
        r, c = im["cell"]
        img.anchor = f"{get_column_letter(c + 1)}{r + 1}"
        img.width, img.height = im["width"], im["height"]
        ws.add_image(img)


def after_save(sheet):
    for i, im in enumerate(images_of(sheet)):
        im["id"] = i
        im.pop("moved", None)
    remember(sheet)
