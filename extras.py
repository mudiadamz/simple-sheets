"""Things in an .xlsx besides cells: Excel tables, images, charts, pivot tables — read for display /
formulas — and parts that saving through openpyxl would drop (shapes, text boxes, form controls…),
so the app can warn before they are lost."""
import re
import zipfile


def _anchor(a):
    """openpyxl anchor -> (row, col) of its top-left cell (0-based), and (row, col) bottom-right or None."""
    fr = getattr(a, "_from", None)
    if fr is None:
        return (0, 0), None
    to = getattr(a, "to", None)
    return (fr.row, fr.col), ((to.row, to.col) if to is not None else None)


EMU = 9525   # per pixel


def _col_px(ws, c):
    from openpyxl.utils import get_column_letter
    d = ws.column_dimensions.get(get_column_letter(c + 1))
    w = d.width if d is not None and d.width else 8.43
    return 0 if d is not None and d.hidden else w * 7 + 5


def _row_px(ws, r):
    d = ws.row_dimensions.get(r + 1)
    h = d.height if d is not None and d.height else 15
    return 0 if d is not None and d.hidden else h * 96 / 72


def anchor_size(ws, a):
    """A drawing anchor -> (width, height) in pixels at 100 %, or None when it doesn't say."""
    ext = getattr(a, "ext", None)
    if ext is not None and getattr(ext, "width", None):          # one-cell anchor: size given
        return round(ext.width / EMU), round(ext.height / EMU)
    fr, to = getattr(a, "_from", None), getattr(a, "to", None)
    if fr is None or to is None:                                 # absolute anchor / a cell name
        return None
    w = sum(_col_px(ws, c) for c in range(fr.col, to.col)) - fr.colOff / EMU + to.colOff / EMU
    h = sum(_row_px(ws, r) for r in range(fr.row, to.row)) - fr.rowOff / EMU + to.rowOff / EMU
    return (round(w), round(h)) if w > 0 and h > 0 else None


def read(ws, path):
    out = {}
    tables = []
    for t in ws.tables.values():
        m = re.fullmatch(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", t.ref or "")
        if not m:
            continue
        from formula import col_index
        r1, c1, r2, c2 = int(m.group(2)) - 1, col_index(m.group(1)), int(m.group(4)) - 1, col_index(m.group(3))
        info = t.tableStyleInfo
        tables.append({"name": t.displayName or t.name, "range": (r1, c1, r2, c2),
                       "columns": [c.name for c in t.tableColumns],
                       "header": (t.headerRowCount if t.headerRowCount is not None else 1) > 0,
                       "totals": bool(t.totalsRowCount),
                       "style": info.name if info is not None else None,
                       "banded_rows": bool(info.showRowStripes) if info is not None else False,
                       "banded_cols": bool(info.showColumnStripes) if info is not None else False,
                       "filter": t.autoFilter is not None})
    if tables:
        out["tables"] = tables
    images = []
    for i, img in enumerate(getattr(ws, "_images", [])):
        (r, c), end = _anchor(img.anchor)
        try:
            data = img._data()
        except Exception:
            data = None
        w, h = anchor_size(ws, img.anchor) or (img.width, img.height)   # its size in the sheet, not the file's
        images.append({"id": i, "cell": (r, c), "end": end, "width": w, "height": h,
                       "format": (img.format or "").lower(), "data": data})
    if images:
        out["images"] = images
    charts = []
    for ch in getattr(ws, "_charts", []):
        (r, c), end = _anchor(ch.anchor)
        series = []
        for s in ch.series:
            val = s.val.numRef.f if s.val is not None and s.val.numRef is not None else None
            cat = None
            if s.cat is not None:
                ref = s.cat.numRef or s.cat.strRef
                cat = ref.f if ref is not None else None
            name = None
            if s.tx is not None and s.tx.strRef is not None:
                name = s.tx.strRef.f
            series.append({"values": val, "categories": cat, "name": name})
        charts.append({"kind": type(ch).__name__.replace("Chart", "").lower() or "bar", "cell": (r, c), "end": end,
                       "title": chart_title(ch), "series": series,
                       "horizontal": getattr(ch, "type", None) == "bar"})
    if charts:
        out["charts"] = charts
    pivots = getattr(ws, "_pivots", [])
    if pivots:
        out["pivots"] = len(pivots)
    return out


def chart_title(ch):
    t = ch.title
    if t is None:
        return ""
    try:
        return "".join(r.t for p in t.tx.rich.p for r in (p.r or []))
    except AttributeError:
        return ""


def _unkeepable_shapes(xml):
    """Shapes that keep_parts can't copy back (they point at other parts, e.g. a picture fill)."""
    m = re.search(rb'xmlns:(\w+)="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"', xml)
    if not m:
        return False
    p = re.escape(m.group(1))
    anchor = re.compile(rb"<" + p + rb":(twoCellAnchor|oneCellAnchor|absoluteAnchor)\b.*?</" + p + rb":\1>", re.S)
    for a in anchor.finditer(xml):
        body = a.group(0)
        if re.search(rb"<" + p + rb":(sp|grpSp|cxnSp)\b", body) and re.search(rb'\br:(embed|link|id)="', body):
            return True
    return False


def lost_parts(path):
    """What saving through openpyxl would drop from this file (a list of short descriptions)."""
    found = []
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            for n in names:
                if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
                    if "shapes with pictures inside" not in found and _unkeepable_shapes(z.read(n)):
                        found.append("shapes with pictures inside")
            if any(n.startswith("xl/ctrlProps/") for n in names):
                found.append("form controls (buttons, check boxes…)")
            if any(n.startswith("xl/slicers/") or n.startswith("xl/slicerCaches/") for n in names):
                found.append("slicers")
            if any(n.startswith("xl/timelines/") for n in names):
                found.append("timelines")
    except (OSError, zipfile.BadZipFile, KeyError):
        pass
    return found
