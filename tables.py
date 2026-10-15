"""Excel tables (Insert → Table in Excel): banded colours from the table style, ▼ filter buttons in the
header row, and the table grows when you type in the row right below it (saved in the file).

Table dicts come from extras.read: {'name', 'range': (r1, c1, r2, c2), 'columns', 'header', 'totals',
'style', 'banded_rows', 'banded_cols', 'filter'}."""
import re

from formatting import tint

ACCENTS = {1: "#000000", 2: "#4472c4", 3: "#ed7d31", 4: "#a5a5a5", 5: "#ffc000", 6: "#5b9bd5", 0: "#70ad47"}


def colors(style):
    """Table style name -> (header fill, header text, band fill)."""
    m = re.match(r"TableStyle(Light|Medium|Dark)(\d+)", style or "")
    if not m:
        return None
    kind, n = m.group(1), int(m.group(2))
    accent = ACCENTS[n % 7]
    if kind == "Light":
        return None, accent if accent != "#000000" else "#000000", tint(accent[1:], 0.8)
    if kind == "Dark":
        return accent, "#ffffff", tint(accent[1:], 0.4)
    return accent if accent != "#000000" else "#404040", "#ffffff", tint(accent[1:], 0.8)


def tables(sheet):
    return sheet.extra.get("tables") or []


def table_at(sheet, r, c):
    return next((t for t in tables(sheet) if t["range"][0] <= r <= t["range"][2] and t["range"][1] <= c <= t["range"][3]),
                None)


def look(sheet, r, c):
    """Display style a cell gets from its table (under the cell's own formatting)."""
    t = table_at(sheet, r, c)
    if t is None:
        return None
    cs = colors(t.get("style"))
    if cs is None:
        return None
    head_bg, head_fg, band = cs
    r1, c1, r2, _ = t["range"]
    if t["header"] and r == r1:
        st = {"bold": True}
        if head_bg:
            st["bg"] = head_bg
        st["fg"] = head_fg
        return st
    if t["totals"] and r == r2:
        return {"bold": True, "border": {"top": ["double", head_bg or head_fg]}}
    data_row = r - (r1 + 1 if t["header"] else r1)
    if t.get("banded_rows") and data_row % 2 == 0:
        return {"bg": band}
    if t.get("banded_cols") and (c - c1) % 2 == 0:
        return {"bg": band}
    return None


def header_buttons(sheet):
    """[(table, row, (c1, c2))] for tables whose header shows ▼ (and isn't the sheet's live filter)."""
    out = []
    f = sheet.filter
    for t in tables(sheet):
        if t.get("filter") and t["header"] and not (f and f.get("table") == t["name"]):
            r1, c1, _, c2 = t["range"]
            out.append((t, r1, (c1, c2)))
    return out


def start_filter(app, t):
    """First click on a table's ▼: the table's range becomes the live filter (saved into the table)."""
    r1, c1, r2, c2 = t["range"]
    app.apply_mixed([("flt", 0, 0, {"range": (r1, c1, r2, c2), "criteria": {}, "table": t["name"]})])


def grow(app, changes):
    """Typing in the row right under a table (and inside its columns) makes the table one row longer.
    Returns the extra change for the same undo step, or None."""
    s = app.current
    ts = tables(s)
    if not ts:
        return None
    new = None
    for i, t in enumerate(ts):
        r1, c1, r2, c2 = t["range"]
        if t["totals"]:
            continue
        if any(r == r2 + 1 and c1 <= c <= c2 and v for r, c, v in changes):
            new = new or [dict(x) for x in ts]
            new[i]["range"] = (r1, c1, r2 + 1, c2)
    return ("tb", 0, 0, new) if new else None


def changed(sheet):
    return [t["range"] for t in tables(sheet)] != getattr(sheet, "tables_disk", [t["range"] for t in tables(sheet)])


def remember(sheet):
    sheet.tables_disk = [t["range"] for t in tables(sheet)]


def save_fn(name, sheet):
    if not changed(sheet):
        return None
    ranges = {t["name"]: t["range"] for t in tables(sheet)}

    def fn(wb):
        from openpyxl.utils import get_column_letter
        ws = wb[name]
        for t in ws.tables.values():
            rng = ranges.get(t.displayName or t.name)
            if rng:
                r1, c1, r2, c2 = rng
                t.ref = f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2 + 1)}{r2 + 1}"
                if t.autoFilter is not None:
                    t.autoFilter.ref = t.ref
    return fn
