"""Data → What-if: scenario manager (saved in .xlsx), data tables (one / two inputs) and a simple solver.

Scenarios live in Sheet.props["scenarios"]: [{"name", "cells": {"B1": "5", ...}, "comment"}].
"""
import math
import re
import tkinter as tk
from tkinter import messagebox, simpledialog, ttk

from formatting import parse_number
from formula import col_index, col_letters


def ref(r, c):
    return f"{col_letters(c)}{r + 1}"


def cells_of(text):
    """'B1, B3:C4' -> [(r, c), ...] in order (None if it doesn't read)."""
    out = []
    for part in re.split(r"[\s,;]+", (text or "").strip().upper().replace("$", "")):
        if not part:
            continue
        m = re.fullmatch(r"([A-Z]{1,3})(\d+)(?::([A-Z]{1,3})(\d+))?", part)
        if not m:
            return None
        r1, c1 = int(m.group(2)) - 1, col_index(m.group(1))
        r2, c2 = (int(m.group(4)) - 1, col_index(m.group(3))) if m.group(3) else (r1, c1)
        out += [(r, c) for r in range(min(r1, r2), max(r1, r2) + 1) for c in range(min(c1, c2), max(c1, c2) + 1)]
    return out or None


def _text(x):
    if isinstance(x, float):
        x = round(x, 10)
        return str(int(x)) if x.is_integer() and abs(x) < 1e15 else f"{x:.10g}"
    return "" if x is None else str(x)


class WhatIf:
    """Try values in cells without touching undo: with WhatIf(app) as w: w.set({(r, c): x}); w.value(r, c)."""
    def __init__(self, app):
        self.app, self.name, self.sheet = app, app.sheet, app.current
        self.saved = {}

    def __enter__(self):
        return self

    def set(self, values):
        rows = self.sheet.rows
        for (r, c), x in values.items():
            while len(rows) <= r:
                rows.append([])
            while len(rows[r]) <= c:
                rows[r].append("")
            self.saved.setdefault((r, c), rows[r][c])
            rows[r][c] = _text(x)
        self.app.engine.clear()

    def value(self, r, c):
        return self.app.engine.cell(self.name, r, c)

    def __exit__(self, *exc):
        for (r, c), old in self.saved.items():
            self.sheet.rows[r][c] = old
        self.app.engine.clear()


# ---------- scenarios ----------

def scenarios(app):
    return [dict(s) for s in app.current.props.get("scenarios") or []]


def _set_scenarios(app, lst):
    app.apply_mixed([("sp", 0, "scenarios", lst or None)])


def add_scenario(app, name, cells, values=None):
    """cells: [(r, c)]; values: their texts (default: what the cells hold now). Replaces a same-named one."""
    g = app.table
    values = values if values is not None else [g.value(r, c) for r, c in cells]
    lst = [s for s in scenarios(app) if s["name"].casefold() != name.casefold()]
    lst.append({"name": name, "cells": {ref(r, c): str(v) for (r, c), v in zip(cells, values)}, "comment": ""})
    _set_scenarios(app, lst)


def delete_scenario(app, name):
    _set_scenarios(app, [s for s in scenarios(app) if s["name"] != name])


def show_scenario(app, name):
    s = next(s for s in scenarios(app) if s["name"] == name)
    app.set_cells([(*cells_of(k)[0], v) for k, v in s["cells"].items()])


def summary(app, result_cells):
    """A new sheet: changing cells and result cells for 'Current values' and each scenario."""
    import data_more
    lst = scenarios(app)
    changing = []
    for s in lst:
        changing += [k for k in s["cells"] if k not in changing]
    g = app.table
    head = ["", "Current values"] + [s["name"] for s in lst]
    rows = [head, ["Changing cells:"] + [""] * (len(head) - 1)]
    for k in changing:
        r, c = cells_of(k)[0]
        rows.append([k, g.value(r, c)] + [s["cells"].get(k, g.value(r, c)) for s in lst])
    rows.append(["Result cells:"] + [""] * (len(head) - 1))
    results = {k: [] for k in result_cells}
    for s in lst:
        with WhatIf(app) as w:
            w.set({cells_of(k)[0]: v for k, v in s["cells"].items()})
            for rc in result_cells:
                results[rc].append(_text(w.value(*rc)))
    for rc in result_cells:
        rows.append([ref(*rc), _text(app.engine.cell(app.sheet, *rc))] + results[rc])
    styles = {(0, c): {"bold": True, "bg": "#dde5f3"} for c in range(len(head))}
    styles.update({(r, 0): {"bold": True} for r, row in enumerate(rows) if row[0].endswith(":")})
    return data_more.add_table_sheet(app, "Scenario summary", rows, styles)


class ScenarioDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Scenario manager")
        self.transient(app)
        f = ttk.Frame(self, padding=10)
        f.pack(fill="both", expand=True)
        self.list = tk.Listbox(f, height=8, width=34, exportselection=False)
        self.list.grid(row=0, column=0, rowspan=6, sticky="nsew")
        self.list.bind("<<ListboxSelect>>", lambda e: self.describe())
        self.list.bind("<Double-Button-1>", lambda e: self.show())
        for i, (text, cmd) in enumerate([("Add…", self.add), ("Show", self.show), ("Delete", self.delete),
                                         ("Summary…", self.summary), ("Close", self.destroy)]):
            ttk.Button(f, text=text, command=cmd).grid(row=i, column=1, sticky="ew", padx=(8, 0), pady=1)
        self.info = ttk.Label(f, text="", wraplength=320, justify="left")
        self.info.grid(row=6, column=0, columnspan=2, sticky="w", pady=(6, 0))
        self.fill()

    def fill(self):
        self.list.delete(0, "end")
        for s in scenarios(self.app):
            self.list.insert("end", s["name"])
        if self.list.size():
            self.list.selection_set(0)
        self.describe()

    def chosen(self):
        sel = self.list.curselection()
        return self.list.get(sel[0]) if sel else None

    def describe(self):
        name = self.chosen()
        s = next((s for s in scenarios(self.app) if s["name"] == name), None)
        self.info["text"] = ", ".join(f"{k} = {v}" for k, v in s["cells"].items()) if s else "No scenarios yet."

    def add(self, name=None, cells=None):
        g = self.app.table
        name = name or simpledialog.askstring("Add scenario", "Scenario name:", parent=self)
        if not name:
            return None
        r1, c1, r2, c2 = g.sel_cells()
        cells = cells or simpledialog.askstring("Add scenario", "Changing cells (they keep their current values):",
                                                initialvalue=f"{ref(r1, c1)}:{ref(r2, c2)}", parent=self)
        rc = cells_of(cells) if cells else None
        if not rc:
            return None
        add_scenario(self.app, name, rc)
        self.fill()
        return name

    def show(self):
        if self.chosen():
            show_scenario(self.app, self.chosen())

    def delete(self):
        if self.chosen():
            delete_scenario(self.app, self.chosen())
            self.fill()

    def summary(self, results=None):
        results = results or simpledialog.askstring("Scenario summary", "Result cells (formulas), e.g. B5 or B5:B7:",
                                                    parent=self)
        rc = cells_of(results) if results else None
        if not rc or not scenarios(self.app):
            return None
        self.destroy()
        return summary(self.app, rc)


def scenario_manager(app):
    return ScenarioDialog(app)


# ---------- data tables ----------

def data_table(app, rect, row_input=None, col_input=None):
    """Excel's What-if data table, filled with values (not live).
    One input in a column: rect's first column (below the top) holds input values, its top row (right of the
    corner) holds formulas; col_input is the cell those values go into. One input in a row: transposed, with
    row_input. Two inputs: the corner holds the formula, the top row values go to row_input and the first
    column values to col_input. Returns the number of cells filled."""
    r1, c1, r2, c2 = rect
    g = app.table
    out = []
    with WhatIf(app) as w:
        for r in range(r1 + 1, r2 + 1):
            for c in range(c1 + 1, c2 + 1):
                values = {}
                if col_input:
                    values[col_input] = g.value(r, c1)
                if row_input:
                    values[row_input] = g.value(r1, c)
                w.set(values)
                if row_input and col_input:
                    formula_at = (r1, c1)
                elif col_input:
                    formula_at = (r1, c)
                else:
                    formula_at = (r, c1)
                out.append((r, c, _text(w.value(*formula_at))))
    app.set_cells(out)
    return len(out)


def data_table_dialog(app):
    g = app.table
    rect = g.sel_cells()
    if not rect or rect[0] == rect[2] or rect[1] == rect[3]:
        messagebox.showinfo("Data table", "Select the table first: input values down the first column and/or "
                            "across the top row, formulas in the top row / first column (or the corner).", parent=app)
        return None
    row_in = simpledialog.askstring("Data table", "Row input cell (leave empty if none):", parent=app)
    col_in = simpledialog.askstring("Data table", "Column input cell (leave empty if none):", parent=app)
    row_rc, col_rc = cells_of(row_in), cells_of(col_in)
    if not row_rc and not col_rc:
        return None
    return data_table(app, rect, row_rc[0] if row_rc else None, col_rc[0] if col_rc else None)


# ---------- solver ----------

CONSTRAINT = re.compile(r"^\s*([A-Za-z]{1,3}\d+(?::[A-Za-z]{1,3}\d+)?)\s*(<=|>=|=)\s*(\S+)\s*$")


def parse_constraints(lines):
    """['B1>=0', 'B2:B3<=10', 'B4=int'] -> [((r, c), op, number-or-(r, c) or 'int')] (bad lines -> ValueError)."""
    out = []
    for line in lines:
        if not line.strip():
            continue
        m = CONSTRAINT.match(line)
        if not m:
            raise ValueError(line)
        rhs = m.group(3)
        if rhs.lower() in ("int", "integer"):
            bound = "int"
        elif parse_number(rhs) is not None:
            bound = parse_number(rhs)
        elif cells_of(rhs):
            bound = cells_of(rhs)[0]
        else:
            raise ValueError(line)
        out += [(rc, m.group(2), bound) for rc in cells_of(m.group(1))]
    return out


def nelder_mead(f, x0, step=1.0, tol=1e-10, max_iter=4000):
    n = len(x0)
    pts = [list(x0)] + [[x + step if i == j else x for j, x in enumerate(x0)] for i in range(n)]
    vals = [f(p) for p in pts]
    for _ in range(max_iter):
        order = sorted(range(n + 1), key=lambda i: vals[i])
        pts, vals = [pts[i] for i in order], [vals[i] for i in order]
        if abs(vals[-1] - vals[0]) <= tol * (1 + abs(vals[0])) and max(
                max(abs(a - b) for a, b in zip(p, pts[0])) for p in pts[1:]) < 1e-9:
            break
        cen = [sum(p[j] for p in pts[:-1]) / n for j in range(n)]
        ref_ = [c + (c - w) for c, w in zip(cen, pts[-1])]
        fr = f(ref_)
        if fr < vals[0]:
            exp = [c + 2 * (c - w) for c, w in zip(cen, pts[-1])]
            fe = f(exp)
            pts[-1], vals[-1] = (exp, fe) if fe < fr else (ref_, fr)
        elif fr < vals[-2]:
            pts[-1], vals[-1] = ref_, fr
        else:
            con = [c + 0.5 * (w - c) for c, w in zip(cen, pts[-1])]
            fc = f(con)
            if fc < vals[-1]:
                pts[-1], vals[-1] = con, fc
            else:
                pts = [pts[0]] + [[b + 0.5 * (x - b) for x, b in zip(p, pts[0])] for p in pts[1:]]
                vals = [vals[0]] + [f(p) for p in pts[1:]]
    best = min(range(n + 1), key=lambda i: vals[i])
    return pts[best], vals[best]


def solve(app, target, goal, changing, constraints=(), value=None):
    """Simple solver: goal 'max' | 'min' | 'value' for formula cell target by changing cells (list of (r, c)),
    with constraints from parse_constraints. Nelder-Mead on a penalty function, a few restarts.
    Returns (values, objective, ok) without changing the sheet."""
    g = app.table
    x0 = [parse_number(g.value(r, c)) or 0.0 for r, c in changing]
    ints = {rc for rc, op, b in constraints if b == "int"}
    lo, hi = {}, {}            # number bounds on the changing cells themselves: kept exactly by clipping
    for rc, op, b in constraints:
        if rc in changing and isinstance(b, float | int) and not isinstance(b, bool):
            if op in (">=", "="):
                lo[rc] = max(lo.get(rc, -math.inf), b)
            if op in ("<=", "="):
                hi[rc] = min(hi.get(rc, math.inf), b)

    def clip(x):
        return [min(max(v, lo.get(rc, -math.inf)), hi.get(rc, math.inf)) for rc, v in zip(changing, x)]

    with WhatIf(app) as w:
        def at(x):
            x = clip(x)
            w.set({rc: (round(v) if rc in ints else v) for rc, v in zip(changing, x)})
            obj = w.value(*target)
            if not isinstance(obj, (int, float)) or isinstance(obj, bool) or math.isnan(obj):
                return None, math.inf
            pen = 0.0
            for rc, op, b in constraints:
                if b == "int":
                    continue
                lhs = w.value(*rc)
                rhs = w.value(*b) if isinstance(b, tuple) else b
                if not isinstance(lhs, (int, float)) or not isinstance(rhs, (int, float)):
                    pen += 1e6
                    continue
                gap = {"<=": max(lhs - rhs, 0), ">=": max(rhs - lhs, 0), "=": abs(lhs - rhs)}[op]
                pen += gap * gap
            return obj, pen

        def cost(x, weight):
            obj, pen = at(x)
            if obj is None:
                return 1e300
            base = {"max": -obj, "min": obj, "value": (obj - (value or 0)) ** 2}[goal]
            return base + weight * pen

        x = x0
        for weight in (1e2, 1e4, 1e6, 1e8):
            for _ in range(3):
                x, _ = nelder_mead(lambda p: cost(p, weight), x, step=max(1.0, max((abs(v) for v in x), default=1) * 0.2))
        x = [round(v) if rc in ints else v for rc, v in zip(changing, clip(x))]
        x = [round(v, 9) for v in x]
        obj, pen = at(x)
    ok = obj is not None and pen < 1e-6 and (goal != "value" or abs(obj - (value or 0)) < 1e-6 * (1 + abs(value or 0)))
    return x, obj, ok


class SolverDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Solver")
        self.transient(app)
        f = ttk.Frame(self, padding=10)
        f.pack(fill="both", expand=True)
        ttk.Label(f, text="Set objective").grid(row=0, column=0, sticky="w")
        self.target = ttk.Entry(f, width=12)
        self.target.insert(0, ref(*app.table.cur))
        self.target.grid(row=0, column=1, sticky="w")
        self.goal = tk.StringVar(master=self, value="max")
        g = ttk.Frame(f)
        g.grid(row=1, column=0, columnspan=2, sticky="w", pady=4)
        for text, v in (("Max", "max"), ("Min", "min"), ("Value of", "value")):
            ttk.Radiobutton(g, text=text, value=v, variable=self.goal).pack(side="left")
        self.value = ttk.Entry(g, width=8)
        self.value.pack(side="left")
        ttk.Label(f, text="By changing cells").grid(row=2, column=0, sticky="w")
        self.changing = ttk.Entry(f, width=20)
        self.changing.grid(row=2, column=1, sticky="w")
        ttk.Label(f, text="Constraints (one per line: B1>=0, B2<=C1, B3=int)").grid(row=3, column=0, columnspan=2,
                                                                                    sticky="w", pady=(6, 0))
        self.cons = tk.Text(f, width=34, height=5)
        self.cons.grid(row=4, column=0, columnspan=2, sticky="ew")
        b = ttk.Frame(f)
        b.grid(row=5, column=0, columnspan=2, sticky="e", pady=(8, 0))
        ttk.Button(b, text="Solve", command=self.ok).pack(side="left", padx=4)
        ttk.Button(b, text="Close", command=self.destroy).pack(side="left")

    def ok(self):
        target, changing = cells_of(self.target.get()), cells_of(self.changing.get())
        try:
            cons = parse_constraints(self.cons.get("1.0", "end").splitlines())
            value = float(self.value.get()) if self.goal.get() == "value" else None
        except ValueError as e:
            messagebox.showwarning("Solver", f"Can't read: {e}", parent=self)
            return None
        if not target or not changing:
            messagebox.showwarning("Solver", "Type the objective cell and the cells to change.", parent=self)
            return None
        x, obj, ok = solve(self.app, target[0], self.goal.get(), changing, cons, value)
        self.app.set_cells([(r, c, _text(v)) for (r, c), v in zip(changing, x)])
        msg = (f"Solved: {ref(*target[0])} = {_text(obj)}" if ok
               else "No solution meets every constraint; the closest one found was put in.")
        self.app.status["text"] = "Solver: " + msg
        messagebox.showinfo("Solver", msg, parent=self)
        return ok


def solver(app):
    return SolverDialog(app)


# ---------- saving: Excel <scenarios> ----------

def read(ws):
    sl = ws.scenarios
    if sl is None or not sl.scenario:
        return None
    return [{"name": s.name or f"Scenario {i + 1}", "comment": s.comment or "",
             "cells": {c.r.replace("$", ""): str(c.val if c.val is not None else "") for c in s.inputCells}}
            for i, s in enumerate(sl.scenario)]


def write(ws, lst):
    from openpyxl.worksheet.scenario import InputCells, Scenario, ScenarioList
    if not lst:
        ws.scenarios = None
        return
    ws.scenarios = ScenarioList(scenario=[
        Scenario(inputCells=[InputCells(r=k, val=v) for k, v in s["cells"].items()], name=s["name"],
                 count=len(s["cells"]), comment=s.get("comment") or None) for s in lst])
