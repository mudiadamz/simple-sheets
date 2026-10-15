"""Editing charts (type, title, data, move, resize) and pictures (insert, move, resize, delete);
pictures are kept even without Pillow (the exe doesn't ship it)."""
import base64
import zipfile

import openpyxl

import charts
import images
import pillowless
import workbooks
from conftest import key, open_in
from test_p3_data import chart_file

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAFklEQVR4nGP8z8DAwMDAxMDAwMDAAAANHQEDasKb6QAA"
                       "AABJRU5ErkJggg==")


def test_sniff_formats():
    assert pillowless.sniff(PNG) == ("PNG", (2, 2))
    assert pillowless.sniff(b"GIF89a" + b"\x05\x00\x03\x00" + b"\0" * 20) == ("GIF", (5, 3))
    jpeg = b"\xff\xd8\xff\xe0\x00\x04ab\xff\xc0\x00\x11\x08\x00\x07\x00\x09" + b"\0" * 20
    assert pillowless.sniff(jpeg) == ("JPEG", (9, 7))
    try:
        pillowless.sniff(b"hello")
        raise AssertionError
    except OSError:
        pass


def test_edit_and_move_chart(app, tmp_path):
    path = chart_file(tmp_path / "c.xlsx")
    g = open_in(app, path)
    ch = g.sheet.charts[0]
    assert charts.data_range(ch) == "A1:B4"
    moved = charts.nudge(app, ch, 2, 1)
    assert moved["cell"] == (ch["cell"][0] + 2, ch["cell"][1] + 1) and moved["id"] == ch["id"]
    bigger = charts.nudge(app, moved, 3, 0, resize=True)
    assert charts.size_of(bigger)[0] == charts.size_of(moved)[0] + 3
    assert app.save()
    anchor = openpyxl.load_workbook(path)["Data"]._charts[0].anchor
    assert (anchor._from.row, anchor._from.col) == bigger["cell"]
    g = open_in(app, path)
    d = charts.EditDialog(app, g.sheet.charts[0])
    d.kind.set("Line")
    d.title_var.set("Monthly")
    new = d.ok()
    assert new["kind"] == "line" and new["title"] == "Monthly" and "id" not in new
    assert app.save()
    c = openpyxl.load_workbook(path)["Data"]._charts[0]
    assert type(c).__name__ == "LineChart"


def test_insert_chart_dialog(app):
    app.set_cells([(0, 0, "Month"), (0, 1, "Sales"), (1, 0, "Jan"), (1, 1, "10"), (2, 0, "Feb"), (2, 1, "30")])
    app.table.select(0, 0, 2, 1)
    d = charts.insert_dialog(app)
    d.kind.set(next(k for k, v in charts.KINDS if v == "line"))
    d.title_var.set("Trend")
    ch = d.ok()                                  # read the choices, then close (it used to read after closing)
    assert not d.winfo_exists()
    assert ch["kind"] == "line" and ch["title"] == "Trend" and app.current.charts[-1] is ch


def test_pick_and_keys(app, tmp_path):
    g = open_in(app, chart_file(tmp_path / "c.xlsx"))
    app.update()
    g.redraw()
    (x1, y1, x2, y2), ch = g.chart_boxes[0]
    g.cv.event_generate("<ButtonPress-1>", x=(x1 + x2) // 2, y=(y1 + y2) // 2, when="now")
    g.cv.event_generate("<ButtonRelease-1>", x=(x1 + x2) // 2, y=(y1 + y2) // 2, when="now")
    assert g.selected_chart is ch
    key(app, "<Right>")
    assert g.sheet.charts[0]["cell"][1] == ch["cell"][1] + 1
    key(app, "<Delete>")
    assert g.sheet.charts == []
    app.undo()
    assert len(g.sheet.charts) == 1


def drag(g, x, y, x2, y2):
    """Press, move in a few steps (B1-Motion), release - like the mouse does."""
    g.cv.event_generate("<ButtonPress-1>", x=x, y=y, when="now")
    for k in (1, 2, 3):
        g.cv.event_generate("<B1-Motion>", x=x + (x2 - x) * k // 3, y=y + (y2 - y) * k // 3, when="now")
    g.cv.event_generate("<ButtonRelease-1>", x=x2, y=y2, when="now")


def test_drag_chart_with_mouse(app, tmp_path):
    g = open_in(app, chart_file(tmp_path / "c.xlsx"))
    app.update()
    g.redraw()
    (x1, y1, x2, y2), ch = g.chart_boxes[0]
    before = ch["cell"]
    g.cv.event_generate("<ButtonPress-1>", x=(x1 + x2) // 2, y=(y1 + y2) // 2, when="now")
    g.cv.event_generate("<B1-Motion>", x=(x1 + x2) // 2 + 40, y=(y1 + y2) // 2, when="now")
    assert g.cv.find_withtag("dragpreview")                  # an outline shows where it lands
    g.cv.event_generate("<ButtonRelease-1>", x=(x1 + x2) // 2 + 200, y=(y1 + y2) // 2 + 100, when="now")
    moved = g.sheet.charts[0]
    assert moved["cell"][0] > before[0] and moved["cell"][1] > before[1]
    app.update()
    g.redraw()
    (x1, y1, x2, y2), ch = g.chart_boxes[0]
    w0 = charts.size_of(ch)[0]
    drag(g, x2 - 2, y2 - 2, x2 + 150, y2 + 20)              # the corner resizes it
    assert charts.size_of(g.sheet.charts[0])[0] > w0
    app.undo()
    app.undo()
    assert g.sheet.charts[0]["cell"] == before


def png(w, h):
    """A plain grey PNG (stdlib only)."""
    import struct
    import zlib

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\0" + b"\x80\x80\x80" * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def test_drag_picture_with_mouse(app, tmp_path):
    pic = tmp_path / "big.png"
    pic.write_bytes(png(60, 40))
    g = open_in(app, workbooks.plain(tmp_path / "p.xlsx"))
    g.goto(3, 2)
    images.insert(app, str(pic))
    app.update()
    g.redraw()
    (x1, y1, x2, y2), im = g.image_items[0]
    drag(g, x1 + 5, y1 + 5, x1 + 250, y1 + 120)              # body: moves
    moved = g.sheet.extra["images"][0]
    assert moved["cell"][0] > 3 and moved["cell"][1] > 2
    app.update()
    g.redraw()
    (x1, y1, x2, y2), im = g.image_items[0]
    drag(g, x2 - 1, y2 - 1, x2 + 60, y2 + 40)                # corner: resizes
    assert g.sheet.extra["images"][0]["width"] > moved["width"]


def big_picture(app, tmp_path, path=None):
    pic = tmp_path / "big.png"
    pic.write_bytes(png(60, 40))
    g = open_in(app, path or workbooks.plain(tmp_path / "p.xlsx"))
    g.goto(3, 2)
    images.insert(app, str(pic))
    app.update()
    g.redraw()
    return g


def test_cursor_over_picture_and_handle(app, tmp_path):
    g = big_picture(app, tmp_path)
    (x1, y1, x2, y2), im = g.image_items[0]
    g.cv.event_generate("<Motion>", x=x2 + 3, y=y2 + 3, when="now")      # the handle sticks out of the corner
    assert g.cv["cursor"] == "size_nw_se"
    g.cv.event_generate("<Motion>", x=(x1 + x2) // 2, y=(y1 + y2) // 2, when="now")
    assert g.cv["cursor"] == "fleur"
    g.cv.event_generate("<Motion>", x=x2 + 60, y=y2 + 60, when="now")
    assert g.cv["cursor"] == ""


def test_grab_handle_outside_picture_resizes_not_selects(app, tmp_path):
    g = big_picture(app, tmp_path)
    (x1, y1, x2, y2), im = g.image_items[0]
    cur = g.cur
    drag(g, x2 + 4, y2 + 4, x2 + 64, y2 + 44)                # pressed on the part of the handle outside
    new = g.sheet.extra["images"][0]
    assert new["width"] > im["width"] and new["cell"] == im["cell"] and g.cur == cur
    assert abs(new["width"] / new["height"] - 1.5) < 0.05     # proportions kept


def test_resized_picture_drawn_and_reopened_at_new_size(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = big_picture(app, tmp_path, path)
    im = g.sheet.extra["images"][0]
    images.replace(app, im, {**im, "width": 90, "height": 60, "moved": True})
    app.update()
    g.redraw()
    (x1, y1, x2, y2), _ = g.image_items[0]
    assert abs((x2 - x1) - g.px(90)) <= 2 and abs((y2 - y1) - g.px(60)) <= 2     # drawn at that size
    assert app.save()
    g = open_in(app, path)
    again = g.sheet.extra["images"][0]
    assert (again["width"], again["height"]) == (90, 60)
    images.replace(app, again, {**again, "width": 45, "height": 30, "moved": True})
    assert app.save()                                          # resized again after a save
    assert (open_in(app, path).sheet.extra["images"][0]["width"]) == 45


def test_size_from_two_cell_anchor(tmp_path):
    """Excel saves pictures anchored to two cells: the size is the space between them."""
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

    import extras
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["B"].width = 20
    a = TwoCellAnchor(_from=AnchorMarker(col=1, row=1), to=AnchorMarker(col=3, row=4))
    assert extras.anchor_size(ws, a) == (round(20 * 7 + 5 + 8.43 * 7 + 5), 60)


def test_ratio():
    assert images.ratio(90, 60) == (3, 2) and images.ratio(30, 60) == (1, 2) and images.ratio(60, 60) == (1, 1)


def test_insert_move_resize_delete_picture(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    pic = tmp_path / "dot.png"
    pic.write_bytes(PNG)
    g = open_in(app, path)
    g.goto(3, 2)
    im = images.insert(app, str(pic))
    assert im["cell"] == (3, 2) and (im["width"], im["height"]) == (2, 2)
    app.update()
    g.redraw()
    assert g.image_items
    im = images.nudge(app, im, 1, 1)
    im = images.nudge(app, im, 1, 0, resize=True)
    assert im["cell"] == (4, 3) and im["width"] >= 2
    assert app.save()
    wb = openpyxl.load_workbook(path)
    assert len(wb.active._images) == 1
    g = open_in(app, path)
    assert len(g.sheet.extra["images"]) == 1
    images.delete(app, g.sheet.extra["images"][0])
    assert app.save()
    assert len(openpyxl.load_workbook(path).active._images) == 0


def test_pictures_kept_without_pillow(app, tmp_path):
    from openpyxl.drawing.image import Image
    pic = tmp_path / "dot.png"
    pic.write_bytes(PNG)
    wb = openpyxl.Workbook()
    wb.active.add_image(Image(str(pic)), "B2")
    path = tmp_path / "i.xlsx"
    wb.save(path)
    pillowless.install(force=True)                   # as in the exe: no Pillow
    try:
        open_in(app, path)
        app.set_cells([(0, 0, "x")])
        assert app.save()
        with zipfile.ZipFile(path) as z:
            media = [n for n in z.namelist() if n.startswith("xl/media/")]
            assert media and z.read(media[0]) == PNG     # the very same bytes
    finally:
        pillowless.uninstall()
