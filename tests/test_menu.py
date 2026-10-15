"""Right-click menu, inserting/deleting rows, columns and cells, paste special."""
import openpyxl

import menus
import workbooks
from conftest import open_in


def build_menu(app):
    """The right-click menu, built but not shown."""
    shown = {}
    orig = menus.tk.Menu.tk_popup
    menus.tk.Menu.tk_popup = lambda self, x, y: shown.setdefault("m", self)
    try:
        menus.show(app, 0, 0)
    finally:
        menus.tk.Menu.tk_popup = orig
    return shown["m"]


def menu_states(app):
    """{label: 'normal' / 'disabled'} of the right-click menu."""
    m = build_menu(app)
    return {m.entrycget(i, "label"): m.entrycget(i, "state")
            for i in range(m.index("end") + 1) if m.type(i) != "separator"}


def menu_labels(app):
    return list(menu_states(app))


def test_right_click_selection_and_labels(app, tmp_path, monkeypatch):
    g = open_in(app, workbooks.layout(tmp_path / "m.xlsx"))
    monkeypatch.setattr(app, "context_menu", lambda x, y: None)
    g.goto(0, 0)
    g.cv.event_generate("<Button-3>", x=g.col_x(1) + 5, y=g.row_y(2) + 5, when="now")
    assert g.cur == (2, 1)
    g.select(1, 0, 2, 0)
    g.cv.event_generate("<Button-3>", x=g.col_x(0) + 5, y=g.row_y(1) + 5, when="now")
    assert g.sel() == (1, 0, 2, 0)
    labels = menu_labels(app)
    assert {"Insert 2 rows above", "Delete rows 2 - 3", "Paste special", "Insert 1 column left"} <= set(labels)


def test_insert_delete_rows_cols(app, tmp_path):
    path = workbooks.layout(tmp_path / "m.xlsx")
    g = open_in(app, path)
    other = app.sheets["Other"]

    app.restructure("row", 2, 1)
    assert g.text(2, 0) == "" and g.text(3, 0) == "B"
    assert g.text(3, 3) == "=B4*C4" and g.text(5, 3) == "=SUM(D2:D5)" and g.value(5, 3) == "140"
    assert other.rows[0][0] == "=Main!D6" and other.rows[1][0] == "=SUM(Main!B2:B5)"
    assert (7, 0, 7, 1) in g.sheet.merges and g.sheet.heights.get(3) == 30

    app.restructure("col", 1, 1)
    assert g.text(3, 2) == "2" and g.text(3, 4) == "=C4*D4"
    assert 3 in g.sheet.widths and (7, 0, 7, 2) in g.sheet.merges
    app.restructure("col", 1, -1)
    assert g.text(3, 1) == "2" and g.text(3, 3) == "=B4*C4"
    app.undo()
    assert g.text(3, 2) == "2"
    app.redo()

    app.restructure("row", 3, -1)    # delete row "B" (it holds the border and the 30pt height)
    assert g.text(3, 0) == "C" and g.text(4, 3) == "=SUM(D2:D4)" and g.value(4, 3) == "100"
    assert app.dirty()

    assert app.save() and not app.dirty()
    wb = openpyxl.load_workbook(path)
    w = wb["Main"]
    assert [w.cell(r, 1).value for r in range(1, 6)] == ["Item", "A", None, "C", None]
    assert w["D4"].value == "=B4*C4" and w["D5"].value == "=SUM(D2:D4)"
    assert not w["B3"].border.top.style and not w["B4"].border.top.style
    assert "A7:B7" in str(w.merged_cells) and w.column_dimensions["C"].width == 20
    assert not w.row_dimensions[3].height and not w.row_dimensions[4].height
    assert wb["Other"]["A1"].value == "=Main!D5"

    app.undo()   # undo after save brings row "B" back
    assert g.text(3, 0) == "B" and app.dirty()
    assert app.save()
    w = openpyxl.load_workbook(path)["Main"]
    assert w["A4"].value == "B" and w["D6"].value == "=SUM(D2:D5)"
    assert w["B4"].border.top.style == "thin"   # borders are cell formatting: undo brings them back too


def test_border_moves_with_inserted_row(app, tmp_path):
    path = workbooks.layout(tmp_path / "m.xlsx")
    open_in(app, path)
    app.restructure("row", 0, 2)
    assert app.save()
    w = openpyxl.load_workbook(path)["Main"]
    assert w["B5"].border.top.style == "thin" and w.row_dimensions[5].height == 30
    assert "A9:B9" in str(w.merged_cells) and w["A3"].value == "Item"


def test_shift_cells_and_paste_special(app, tmp_path):
    g = open_in(app, workbooks.layout(tmp_path / "m.xlsx"))
    app.set_cells([(10, 0, "x1"), (11, 0, "x2"), (10, 1, "y1")])
    g.select(10, 0, 10, 0)
    app.shift_cells("down")
    assert (g.text(10, 0), g.text(11, 0), g.text(12, 0), g.text(10, 1)) == ("", "x1", "x2", "y1")
    app.shift_cells("up")
    assert (g.text(10, 0), g.text(11, 0)) == ("x1", "x2")
    g.select(10, 0, 10, 0)
    app.shift_cells("right")
    assert (g.text(10, 1), g.text(10, 2)) == ("x1", "y1")
    app.undo()

    g.goto(1, 3)
    g.copy()
    g.goto(12, 3)
    g.paste("values")
    assert g.text(12, 3) == "10" and not g.sheet.styles.get((12, 3))
    g.goto(0, 0)
    g.copy()
    g.goto(12, 4)
    g.paste("format")
    assert g.text(12, 4) == "" and (g.sheet.styles.get((12, 4)) or {}).get("bold")
