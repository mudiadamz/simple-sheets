"""Defined names (named ranges) in formulas, name manager, saving."""
import openpyxl
from openpyxl.workbook.defined_name import DefinedName

import sheet_io as io
from conftest import open_in


def named_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 5):
        ws.append([f"i{r}", r * 10])
    ws["D1"] = "=SUM(prices)"
    wb.defined_names["prices"] = DefinedName("prices", attr_text="Data!$B$1:$B$4")
    wb.defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area", localSheetId=0, attr_text="Data!$A$1:$B$4")
    wb.save(path)
    return path


def test_load_and_use_names(app, tmp_path):
    path = named_file(tmp_path / "n.xlsx")
    assert io.load_names(str(path)) == {"prices": "Data!$B$1:$B$4"}
    g = open_in(app, path)
    assert g.value(0, 3) == "100"
    app.set_cells([(5, 0, "=AVERAGE(Prices)*2"), (6, 0, "=nosuchname")])
    assert g.value(5, 0) == "50" and g.value(6, 0) == "#NAME?"


def test_name_manager_add_delete_undo(app):
    g = app.table
    app.set_cells([(r, 0, str(r + 1)) for r in range(3)] + [(0, 2, "=SUM(nums)")])
    assert g.value(0, 2) == "#NAME?"
    g.select(0, 0, 2, 0)
    m = app.manage_names()
    assert m.refers.get() == "Sheet1!$A$1:$A$3"           # current selection
    m.name.insert(0, "nums")
    assert m.add()
    assert app.names == {"nums": "Sheet1!$A$1:$A$3"} and g.value(0, 2) == "6" and app.dirty()
    m.name.delete(0, "end")
    m.name.insert(0, "A1")                                 # looks like a cell: refused
    assert not m.add()
    m.name.delete(0, "end")
    m.name.insert(0, "nums")
    m.delete()
    assert app.names == {} and g.value(0, 2) == "#NAME?"
    app.undo()
    assert app.names == {"nums": "Sheet1!$A$1:$A$3"}
    app.undo()
    assert app.names == {} and app.hpos == 1             # back to just the typed cells
    m.destroy()


def test_names_follow_rows_and_renames(app, tmp_path):
    g = open_in(app, named_file(tmp_path / "n.xlsx"))
    app.restructure("row", 0, 2)
    assert app.names["prices"] == "Data!$B$3:$B$6" and g.value(2, 3) == "100"
    app.rename_sheet("Data", "My Data")
    assert app.names["prices"] == "'My Data'!$B$3:$B$6"
    app.undo()
    assert app.names["prices"] == "Data!$B$3:$B$6"
    app.undo()
    assert app.names["prices"] == "Data!$B$1:$B$4"


def test_name_box_goes_to_name(app, tmp_path):
    g = open_in(app, named_file(tmp_path / "n.xlsx"))
    app.goto_ref("prices")
    assert g.sel() == (0, 1, 3, 1)


def test_saved(app, tmp_path):
    path = named_file(tmp_path / "n.xlsx")
    open_in(app, path)
    app.set_names({**app.names, "first": "Data!$A$1"})
    assert app.save(), app.dialogs
    wb = openpyxl.load_workbook(path)
    assert wb.defined_names["first"].attr_text == "Data!$A$1"
    assert wb.defined_names["prices"].attr_text == "Data!$B$1:$B$4"
    assert wb["Data"].print_area                          # built-in names untouched
    assert not app.dirty()
