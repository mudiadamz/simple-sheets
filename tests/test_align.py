"""Horizontal / vertical alignment dropdowns."""
import openpyxl

import sheet_io as io
import workbooks
from conftest import open_in


def image_of(w):
    """Name of a widget's image (cget returns a tuple for ttk widgets)."""
    img = w.cget("image")
    return str(img[0] if isinstance(img, tuple) and img else img)


def text_anchor(g):
    items = g.cv.find_withtag("txt")
    return g.cv.itemcget(items[0], "anchor")


def test_align_buttons_are_grouped(app):
    fb = app.fmtbar
    images = {str(fb.icons[a]) for a in ("left", "center", "right", "justify")}
    on_bar = [w for w in fb.winfo_children() if "image" in w.keys() and image_of(w) in images]
    assert len(on_bar) == 1 and on_bar[0] is fb.align_menu.button   # one button, not four


def test_hover_opens_and_pick_applies(app):
    g, fb = app.table, app.fmtbar
    dd = fb.align_menu
    app.set_cells([(0, 0, "x")])
    g.goto(0, 0)
    dd.button.event_generate("<Enter>", when="now")
    app.update()
    assert dd.popup is not None and dd.popup.winfo_ismapped()
    choices = [w for w in dd.popup.winfo_children()[0].winfo_children()]
    assert len(choices) == 4
    choices[1].invoke()                                   # center
    app.update()
    assert g.style(0, 0)["align"] == "center" and dd.popup is None
    assert image_of(dd.button) == str(fb.icons["center"])   # shows the current choice
    g.goto(1, 0)
    assert image_of(dd.button) == str(fb.icons["left"])     # default for a plain cell


def test_popup_closes_when_mouse_leaves(app, monkeypatch):
    dd = app.fmtbar.valign_menu
    dd.open()
    app.update()
    monkeypatch.setattr(dd.button, "winfo_pointerxy", lambda: (dd.button.winfo_rootx() + 2,
                                                               dd.button.winfo_rooty() + 2))
    dd.close_if_away()
    assert dd.popup is not None                            # still over the button
    monkeypatch.setattr(dd.button, "winfo_pointerxy", lambda: (-5000, -5000))
    dd.close_if_away()
    assert dd.popup is None
    dd.toggle()                                            # a click also opens / closes it
    assert dd.popup is not None
    dd.toggle()
    assert dd.popup is None


def test_vertical_align(app):
    g, fb = app.table, app.fmtbar
    app.set_cells([(0, 0, "text"), (0, 1, "12")])
    g.goto(0, 0)
    g.select(0, 0, 0, 1)
    assert text_anchor(g) == "sw"                          # default: bottom (like Excel)
    fb.set_valign("top")
    assert g.style(0, 0)["valign"] == "top" and g.style(0, 1)["valign"] == "top"
    anchors = sorted(g.cv.itemcget(i, "anchor") for i in g.cv.find_withtag("txt"))
    assert anchors == ["ne", "nw"]                         # number right-aligned, both at the top
    fb.set_valign("center")
    anchors = sorted(g.cv.itemcget(i, "anchor") for i in g.cv.find_withtag("txt"))
    assert anchors == ["e", "w"]
    g.goto(0, 0)
    assert fb.valign.get() == "center" and "Middle" in app.fmt_info["text"]
    fb.set_valign("bottom")                               # back to the default: not stored
    assert "valign" not in (g.style(0, 0) or {})
    app.undo()
    assert g.style(0, 0)["valign"] == "center"


def test_center_both_ways(app):
    g, fb = app.table, app.fmtbar
    app.set_cells([(0, 0, "x")])
    g.goto(0, 0)
    fb.set_align("center")
    fb.set_valign("center")
    assert text_anchor(g) == "center"


def test_valign_saved_and_loaded(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.goto(1, 0)
    app.fmtbar.set_valign("top")
    g.goto(1, 1)
    app.fmtbar.set_valign("center")
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Data"]
    assert ws["A2"].alignment.vertical == "top" and ws["B2"].alignment.vertical == "center"
    s = io.load(str(path))["Data"]
    assert s.styles[(1, 0)]["valign"] == "top" and s.styles[(1, 1)]["valign"] == "center"


def test_valign_icons_follow_theme(app):
    icon = app.fmtbar.icons["v-top"]
    app.set_theme(True)
    assert app.fmtbar.icons["v-top"] is icon                # redrawn in place
    app.set_theme(False)
