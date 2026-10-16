"""Line breaks in cells (Ctrl+Enter) and wrapped text."""
import openpyxl
import pytest

import sheet_io as io
import workbooks
from conftest import key, open_in

LONG = "The quick brown fox jumps over the lazy dog and keeps running far away"


def test_ctrl_enter_adds_line_and_turns_on_wrap(app):
    g = app.table
    g.goto(0, 0)
    g.edit("first")
    app.update()
    e = g.editor
    h1 = e.winfo_height()
    key(app, "<Control-Return>", widget=e)
    e.insert("end", "second")
    key(app, "<Alt-Return>", widget=e)
    e.insert("end", "third")
    key(app, "<KeyRelease>", widget=e)
    assert g.editor is e and e.get() == "first\nsecond\nthird"
    assert e.winfo_height() > h1                             # editor grows per line
    key(app, "<Return>", widget=e)                           # plain Enter finishes
    assert g.text(0, 0) == "first\nsecond\nthird" and g.cur == (1, 0)
    assert g.style(0, 0) == {"wrap": True}                   # like Excel's Alt+Enter
    assert g.rh(0) > g.rowh                                  # row grew to show 3 lines
    app.undo()                                               # one step: text and wrap
    assert g.text(0, 0) == "" and not g.style(0, 0)


@pytest.mark.parametrize("text, lines", [
    ("short", 1),
    ("a\nb\nc", 3),
    (LONG, None),               # several lines, each fits
    ("x" * 200, None),          # one long word gets cut
])
def test_wrap_lines_fit(app, text, lines):
    g = app.table
    font = g.font(None)
    width = g.px(80)
    out = g.wrap_lines(text, font, width)
    assert all(font.measure(line) <= width for line in out)
    assert "".join(out).replace(" ", "") == text.replace("\n", "").replace(" ", "")
    if lines:
        assert len(out) == lines
    else:
        assert len(out) > 1


def test_wrap_toggle_from_toolbar(app):
    g, fb = app.table, app.fmtbar
    app.set_cells([(0, 0, LONG)])
    g.goto(0, 0)
    one_line = g.rh(0)
    fb.wrap.set(True)
    fb.toggle_wrap()
    assert g.style(0, 0) == {"wrap": True} and g.rh(0) > one_line
    text_items = [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")]
    assert any("\n" in t for t in text_items)                 # drawn on several lines
    g.goto(1, 0)
    assert not fb.wrap.get()
    g.goto(0, 0)
    assert fb.wrap.get() and "Wrap" in __import__("formatting").describe(app.current.styles.get(app.table.cur))
    fb.wrap.set(False)
    fb.toggle_wrap()
    assert g.rh(0) == one_line


def test_narrower_column_more_lines(app):
    g = app.table
    app.set_cells([(0, 0, LONG)])
    g.goto(0, 0)
    app.format_cells(lambda st: {**st, "wrap": True})
    h_wide = g.rh(0)
    g.sheet.px[0] = 40
    g.rowpx = {}
    assert g.rh(0) > h_wide


def test_editing_wrapped_cell_wraps_in_editor(app):
    g = app.table
    app.set_cells([(0, 0, LONG)])
    g.goto(0, 0)
    app.format_cells(lambda st: {**st, "wrap": True})
    g.edit()
    app.update()
    assert g.editor.cget("wrap") == "word" and g.editor.winfo_width() == g.colw(0)
    g.commit(cancel=True)


def test_formula_bar_shows_line_breaks(app):
    g = app.table
    app.set_cells([(0, 0, "a\nb")])
    g.goto(0, 0)
    assert app.value.get() == "a↵b"
    app.value.focus_force()
    app.value.delete(0, "end")
    app.value.insert(0, "x↵y↵z")
    app.value_commit()
    assert g.text(0, 0) == "x\ny\nz"


def test_wrap_saved_and_loaded(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.goto(3, 0)
    g.edit("line 1")
    app.update()
    key(app, "<Control-Return>", widget=g.editor)
    g.editor.insert("end", "line 2")
    g.commit()
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Data"]
    assert ws["A4"].value == "line 1\nline 2" and ws["A4"].alignment.wrap_text
    s = io.load(str(path))["Data"]
    assert s.rows[3][0] == "line 1\nline 2" and s.styles[(3, 0)]["wrap"]
