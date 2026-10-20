"""A big .xlsx (over the full-load size) still gets its column widths, row heights, merges, hidden rows /
columns and frozen panes - read by streaming the sheet XML, not by loading the whole workbook."""
import openpyxl

import sheet_io as io


def layout_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 40):
        ws.append([f"r{r}", r, r * 2, "x"])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["C"].hidden = True
    ws.row_dimensions[3].height = 40
    ws.row_dimensions[5].hidden = True
    ws.merge_cells("B10:D11")
    ws.freeze_panes = "B2"
    other = wb.create_sheet("Other")
    other["A1"] = "plain"
    other.merge_cells("A3:B3")
    wb.save(path)
    return path


def test_light_layout_reads_the_xml(tmp_path):
    lay = io.light_layout(str(layout_book(tmp_path / "big.xlsx")))
    d = lay["Data"]
    assert d["widths"][0] == 30 and 2 in d["hidden_cols"]
    assert d["heights"][2] == 40 and d["hidden_rows"] == {4}
    assert d["merges"] == [(9, 1, 10, 3)] and d["freeze"] == (1, 1)
    assert lay["Other"]["merges"] == [(2, 0, 2, 1)] and lay["Other"]["freeze"] is None


def test_big_file_saved_after_an_edit(app, tmp_path, monkeypatch):
    from conftest import open_in
    p = layout_book(tmp_path / "big.xlsx")
    monkeypatch.setattr(io, "FULL_LOAD_BYTES", 0)
    g = open_in(app, p)
    assert g.sheet.merges == [(9, 1, 10, 3)] and 4 in g.sheet.hidden_rows
    app.apply_mixed([("v", 0, 0, "edited")])
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(p)["Data"]
    assert ws["A1"].value == "edited" and [str(m) for m in ws.merged_cells.ranges] == ["B10:D11"]
    assert ws.column_dimensions["A"].width == 30 and ws.column_dimensions["C"].hidden
    assert ws.row_dimensions[5].hidden and ws.row_dimensions[3].height == 40 and ws.freeze_panes == "B2"


def test_big_file_mode_keeps_the_layout(tmp_path, monkeypatch):
    p = layout_book(tmp_path / "big.xlsx")
    full = io.load(str(p))["Data"]
    monkeypatch.setattr(io, "FULL_LOAD_BYTES", 0)                 # force the big-file (read-only) loader
    big = io.load(str(p))["Data"]
    assert big.rows == full.rows                                   # (values as before)
    assert big.widths[0] == full.widths[0] == 30
    assert big.merges == full.merges == [(9, 1, 10, 3)]
    assert big.hidden_rows == full.hidden_rows == {4} and 2 in big.hidden_cols
    assert big.heights.get(2) == 40 and tuple(big.freeze) == (1, 1)
