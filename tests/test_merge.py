"""Merging / unmerging cells."""
import openpyxl

import sheet_io as io
import workbooks
from conftest import open_in
from test_menu import menu_states


def test_merge_keeps_first_value_and_undo(app):
    g = app.table
    app.set_cells([(0, 0, ""), (0, 1, "b"), (1, 0, "c")])
    g.select(0, 0, 1, 2)
    assert app.merge()
    assert g.sheet.merges == [(0, 0, 1, 2)] and g.span(1, 1) == (0, 0, 1, 2)
    assert g.text(0, 0) == "b" and g.text(0, 1) == "" and g.text(1, 0) == ""
    assert any(d[0] == "askokcancel" for d in app.dialogs)      # warned about losing values
    assert g.sel() == (0, 0, 0, 0) and app.fmtbar.merged.get()
    app.undo()                                                    # one step back
    assert g.sheet.merges == [] and (g.text(0, 1), g.text(1, 0)) == ("b", "c")
    assert app.hpos == 1                                          # back to just the typed values


def test_merge_cancel_and_single_cell(app, monkeypatch):
    import simple_sheets
    g = app.table
    app.set_cells([(0, 0, "a"), (0, 1, "b")])
    monkeypatch.setattr(simple_sheets.messagebox, "askokcancel", lambda *a, **k: False)
    g.select(0, 0, 0, 1)
    assert not app.merge() and g.sheet.merges == []
    g.goto(3, 3)
    assert not app.merge()


def test_merge_swallows_overlapping_and_unmerge(app):
    g = app.table
    g.select(0, 1, 1, 2)
    app.merge()
    g.select(1, 0, 1, 1)            # overlaps the B1:C2 block
    app.merge()
    assert g.sheet.merges == [(0, 0, 1, 2)]
    g.goto(1, 1)
    assert app.unmerge() and g.sheet.merges == []


def test_merge_and_center(app):
    g = app.table
    app.set_cells([(0, 0, "Title")])
    g.select(0, 0, 0, 3)
    app.merge(center=True)
    assert g.sheet.styles[(0, 0)]["align"] == "center" and g.sheet.merges == [(0, 0, 0, 3)]


def test_toolbar_toggle(app):
    g = app.table
    fb = app.fmtbar
    g.select(2, 2, 3, 3)
    fb.merged.set(True)
    fb.toggle_merge()
    assert g.sheet.merges == [(2, 2, 3, 3)]
    g.goto(2, 2)
    assert fb.merged.get()
    fb.merged.set(False)
    fb.toggle_merge()
    assert g.sheet.merges == [] and not fb.merged.get()


def test_menu_items(app):
    g = app.table
    g.goto(0, 0)
    states = menu_states(app)
    assert states["Merge cells"] == "disabled" and states["Unmerge cells"] == "disabled"
    g.select(0, 0, 1, 1)
    states = menu_states(app)
    assert states["Merge cells"] == "normal" and states["Merge & center"] == "normal"
    app.merge()
    assert menu_states(app)["Unmerge cells"] == "normal"


def test_save_merge_and_unmerge(app, tmp_path):
    path = workbooks.layout(tmp_path / "m.xlsx")        # has A7:B7 merged
    g = open_in(app, path)
    g.goto(6, 0)
    app.unmerge()
    g.select(9, 0, 9, 3)
    app.merge()
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Main"]
    assert str(ws.merged_cells) == "A10:D10"
    assert ws["B3"].border.top.style == "thin"          # rest of the sheet untouched
    assert io.load(str(path))["Main"].merges == [(9, 0, 9, 3)]
    assert not app.dirty()
    app.undo()                                          # undo the merge after save
    assert app.dirty() and app.save()
    assert str(openpyxl.load_workbook(path)["Main"].merged_cells) == ""


def test_merge_on_new_workbook_saved(app, tmp_path, monkeypatch):
    import simple_sheets
    dst = tmp_path / "new.xlsx"
    monkeypatch.setattr(simple_sheets.filedialog, "asksaveasfilename", lambda **k: str(dst))
    g = app.table
    app.set_cells([(0, 0, "Report")])
    g.select(0, 0, 0, 2)
    app.merge(center=True)
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(dst).active
    assert str(ws.merged_cells) == "A1:C1" and ws["A1"].alignment.horizontal == "center"
