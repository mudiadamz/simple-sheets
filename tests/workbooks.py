"""Small sample workbooks built with openpyxl for the tests."""
import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color


def styled(path):
    """Merged title with a theme fill, number/date formats, a custom column width."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Quarterly report"
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=-0.25))
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["Item", "Qty", "Price", "Date"])
    ws.append(["Apple", 1200, 0.5, dt.date(2024, 3, 5)])
    ws["B3"].number_format = "#,##0"
    ws["C3"].number_format = '"$"0.00'
    ws["D3"].number_format = "mmm d, yyyy"
    ws["A3"].font = Font(italic=True, color="FF0000")
    ws.column_dimensions["A"].width = 25
    wb.save(path)
    return path


def plain(path):
    """Two sheets of plain values."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "qty", "price"])
    ws.append(["apple", 3, 1.5])
    ws["D2"] = "=B2*C2"
    ws["A1"].font = Font(bold=True)
    wb.create_sheet("Other").append(["x"])
    wb.save(path)
    return path


def sales(path):
    """A data block A1:D6 plus a separate block G2:G3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    for row in [["Region", "Q1", "Q2", "Q3"], ["North", 120, 135, 150], ["South", 98, 110, 105],
                ["East", 143, 150, 162], ["West", 87, 95, 101], ["Central", 110, 118, 125]]:
        ws.append(row)
    ws["G2"] = "Notes"
    ws["G3"] = "draft"
    wb.save(path)
    return path


def budget(path):
    """Formulas with currency formats; openpyxl saves no cached results."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Qty", "Price", "Total"])
    for i, (item, q, p) in enumerate([("Paper", 10, 2.5), ("Pens", 25, 0.8), ("Ink", 3, 19.99)], start=2):
        ws.append([item, q, p, f"=B{i}*C{i}"])
        ws[f"D{i}"].number_format = '"$"#,##0.00'
    ws["C6"] = "Sum"
    ws["D6"] = "=SUM(D2:D4)"
    ws["D6"].number_format = '"$"#,##0.00'
    ws["D6"].font = Font(bold=True)
    wb.save(path)
    return path


def hidden(path):
    """Rows 3-4 and column C hidden; columns E:G share one (visible) dimension group."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "H"
    for r in range(1, 7):
        ws.append([f"{col}{r}" for col in "ABCDEFG"])
    ws.row_dimensions[3].hidden = True
    ws.row_dimensions[4].hidden = True
    ws.column_dimensions["C"].hidden = True
    ws.column_dimensions.group("E", "G", hidden=False)
    ws.column_dimensions["E"].width = 15
    wb.save(path)
    return path


def hidden_sheet(path):
    """Sheets A, Secret (hidden), B, Deep (veryHidden). B uses Secret in a formula."""
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.active["A1"] = "a"
    s = wb.create_sheet("Secret")
    s["A1"] = 42
    s.sheet_state = "hidden"
    b = wb.create_sheet("B")
    b["A1"] = "=Secret!A1*2"
    d = wb.create_sheet("Deep")
    d.sheet_state = "veryHidden"
    wb.save(path)
    return path


def links(path):
    """Column B holds hyperlinks whose text is the address (like a list of URLs)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links"
    for r in range(1, 6):
        ws.cell(row=r, column=1, value=f"name {r}")
        url = f"https://example.com/page{r}"
        cell = ws.cell(row=r, column=2, value=url)
        cell.hyperlink = url
        cell.font = Font(color="FF0563C1", underline="single")   # styled like real links (matters:
        # openpyxl only keeps the link of an emptied cell when the cell still has a style)
    wb.save(path)
    return path


def layout(path):
    """Formulas, a border, a merge, a row height, a column width and a second sheet referring back."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main"
    ws.append(["Item", "Qty", "Price", "Total"])
    for i, (item, q, p) in enumerate([("A", 1, 10), ("B", 2, 20), ("C", 3, 30)], start=2):
        ws.append([item, q, p, f"=B{i}*C{i}"])
    ws["D5"] = "=SUM(D2:D4)"
    thin = Side(style="thin", color="000000")
    ws["B3"].border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A7:B7")
    ws["A7"] = "merged"
    ws.row_dimensions[3].height = 30
    ws.column_dimensions["C"].width = 20
    ws2 = wb.create_sheet("Other")
    ws2["A1"] = "=Main!D5"
    ws2["A2"] = "=SUM(Main!B2:B4)"
    wb.save(path)
    return path
