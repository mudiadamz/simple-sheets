"""Formatting toolbar, table styles and smart Ctrl+A."""
import openpyxl

import sheet_io as io
import workbooks
from conftest import key, open_in


def style(g, r, c):
    return g.sheet.styles.get((r, c)) or {}


def test_toolbar_variables_belong_to_this_window(app):
    """Tk variables without a master attach to the first Tk ever made (flaky toggles)."""
    fb = app.fmtbar
    assert all(v._tk is app.tk for v in fb.vars.values()) and fb.align._tk is app.tk
    assert fb.merged._tk is app.tk
    assert all(img.tk is app.tk for img in fb.icons.values())          # images too
    assert all(img.tk is app.tk for img in app.save_icons.values())


def test_format_bar(app, tmp_path):
    path = workbooks.styled(tmp_path / "s.xlsx")
    g, fb = open_in(app, path), app.fmtbar

    g.select(1, 0, 1, 3)
    key(app, "<Control-b>")
    assert all(style(g, 1, c).get("bold") for c in range(4)) and fb.vars["bold"].get()
    key(app, "<Control-b>")
    assert not any(style(g, 1, c).get("bold") for c in range(4))

    app.set_cells([(5, 0, "Big title")])
    g.goto(5, 0)
    fb.preset.set("Title")
    fb.set_preset()
    assert style(g, 5, 0) == {"size": 20.0, "bold": True} and fb.preset.get() == "Title"
    assert g.rh(5) > g.rowh
    fb.size.set("14")
    fb.set_size()
    assert style(g, 5, 0).get("size") == 14.0 and fb.preset.get() == ""

    g.select(2, 0, 2, 3)
    fb.vars["italic"].set(True)
    fb.toggle("italic")
    assert all(style(g, 2, c).get("italic") for c in range(4))
    fb.normal()
    assert not any(style(g, 2, c).get("italic") for c in range(4))
    assert style(g, 2, 1).get("numfmt") == "#,##0"
    fb.align.set("center")
    fb.set_align("center")
    app.format_cells(lambda s: {**s, "fg": "#c00000", "bg": "#fff2cc"})
    assert style(g, 2, 1).get("align") == "center" and style(g, 2, 2).get("bg") == "#fff2cc"
    g.goto(2, 3)
    fb.set_align("justify")
    app.undo()
    assert style(g, 2, 3).get("align") == "center"
    app.redo()
    assert style(g, 2, 3).get("align") == "justify"

    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws["A6"].font.sz == 14 and ws["A6"].font.b
    assert ws["C3"].font.color.rgb == "FFC00000" and ws["C3"].fill.fgColor.rgb == "FFFFF2CC"
    assert ws["B3"].alignment.horizontal == "center" and ws["B3"].number_format == "#,##0"
    assert "A1:D1" in str(ws.merged_cells)
    again = io.load(str(path))["Report"]
    assert again.styles[(5, 0)]["size"] == 14.0 and again.styles[(2, 3)]["align"] == "justify"

    g.select(2, 0, 2, 3)
    fb.clear()
    assert style(g, 2, 1) == {"numfmt": "#,##0"} and style(g, 2, 0) == {}


def test_smart_select_all(app, tmp_path):
    g = open_in(app, workbooks.sales(tmp_path / "t.xlsx"))
    g.goto(3, 2)
    key(app, "<Control-a>")
    assert g.sel() == (0, 0, 5, 3) and app.ref["text"] == "A1:D6" and g.cur == (3, 2)
    key(app, "<Control-a>")
    assert app.ref["text"] == "All"
    g.goto(2, 6)
    key(app, "<Control-a>")
    assert g.sel() == (1, 6, 2, 6)
    g.goto(15, 10)
    key(app, "<Control-a>")
    assert app.ref["text"] == "All"
    g.goto(6, 4)   # empty cell touching the block diagonally
    key(app, "<Control-a>")
    assert g.sel() == (0, 0, 6, 4)


def test_table_styles(app, tmp_path):
    path = workbooks.sales(tmp_path / "t.xlsx")
    g, fb = open_in(app, path), app.fmtbar
    g.goto(2, 1)
    fb.apply_table("4472C4", False)
    assert g.sel() == (0, 0, 5, 3)
    assert style(g, 0, 2) == {"bg": "#4472c4", "fg": "#ffffff", "bold": True}
    assert "bg" in style(g, 1, 0) and "bg" not in style(g, 2, 0) and "bg" in style(g, 3, 3)
    fb.apply_table("70AD47", True)
    assert style(g, 0, 0)["fg"] == "#000000" and style(g, 1, 0)["bg"] != "#4472c4"
    fb.remove_table()
    assert all(not style(g, r, c) for r in range(6) for c in range(4))
    for _ in range(3):
        app.undo()
    assert not app.dirty()
    fb.apply_table("ED7D31", False)
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws["A1"].fill.fgColor.rgb == "FFED7D31" and ws["A1"].font.b


def test_bigger_smaller_text_buttons(app):
    g, fb = app.table, app.fmtbar
    app.set_cells([(0, 0, "a"), (0, 1, "b")])
    g.select(0, 0, 0, 1)
    fb.step_size(1)                                               # 11 -> 12
    assert g.style(0, 0)["size"] == 12 and g.style(0, 1)["size"] == 12
    fb.step_size(1)
    assert g.style(0, 0)["size"] == 14                             # next standard size
    for _ in range(3):
        fb.step_size(-1)
    assert g.style(0, 0)["size"] == 10
    g.select(1, 0, 1, 0)
    app.format_cells(lambda st: {**st, "size": 8.0})
    fb.step_size(-1)
    assert g.style(1, 0)["size"] == 7                              # below the list: one point at a time
    app.undo()
    assert g.style(1, 0)["size"] == 8
    texts = [w["text"] for w in fb.winfo_children() if w.winfo_class() == "TButton"]
    assert "A▴" in texts and "A▾" in texts



def test_toolbar_more_button_on_a_narrow_window(app):
    """A MacBook Air / half-screen window: one row, the groups that don't fit behind ». » shows them as more
    rows, « hides them again; the choice is remembered."""
    from simplecore import settings
    bar = app.fmtbar
    app.update()
    need = sum(sum(w.winfo_reqwidth() + p[0] + p[1] for w, p in g) for g in bar.groups) + 14 * (len(bar.groups) - 1)
    bar.expanded = False
    bar.flow(need + 20)
    assert bar.flow_rows == 1 and bar.hidden_groups == 0 and not bar.more_btn.place_info()
    one_row = bar.cget("height")
    narrow = need // 2
    bar.flow(narrow)                                      # collapsed: still one row, the rest hidden
    assert bar.flow_rows == 1 and bar.hidden_groups > 0 and bar.cget("height") == one_row
    assert bar.more_btn.place_info() and bar.more_btn["text"] == "»"
    shown = [g for g in bar.groups if g[0][0].place_info()]
    assert len(shown) == len(bar.groups) - bar.hidden_groups
    more_x = int(bar.more_btn.place_info()["x"])
    for g in shown:                                       # nothing under the » button
        for w, _ in g:
            assert int(w.place_info()["x"]) + w.winfo_reqwidth() <= more_x
    bar.configure(width=narrow)
    bar.toggle_more()                                     # »: every group shown, on more rows
    bar.flow(narrow)
    assert settings.get("toolbar_expanded") is True
    assert bar.flow_rows >= 2 and bar.hidden_groups == 0 and bar.cget("height") > one_row
    assert bar.more_btn["text"] == "«"
    for g in bar.groups:                                  # a group is never split over two rows
        ys = {int(w.place_info()["y"]) + w.winfo_reqheight() // 2 for w, _ in g}
        assert max(ys) - min(ys) <= 2, g
    placed = [s for s in bar.seps if s.place_info()]
    assert len(placed) == len(bar.groups) - bar.flow_rows  # no separator dangling at a row's end
    bar.toggle_more()
    bar.flow(narrow)
    assert bar.flow_rows == 1 and settings.get("toolbar_expanded") is False


def test_toolbar_waits_for_its_width(app):
    """Nothing is laid out before the window has a size (no one-row-then-wrapped flash at the first open)."""
    bar = app.fmtbar
    bar._flow_key = None
    bar.flow(1)
    assert bar._flow_key is None
