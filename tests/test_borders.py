"""Cell borders: load, draw, apply from the toolbar popup, save."""
import openpyxl
import pytest

import borders
import sheet_io as io
import workbooks
from conftest import open_in


def border(g, r, c):
    return (g.sheet.styles.get((r, c)) or {}).get("border") or {}


def lines(g):
    return [g.cv.coords(i) for i in g.cv.find_withtag("border")]


@pytest.mark.parametrize("kind, cell, want", [
    ("all", (1, 1), {"top", "bottom", "left", "right"}),
    ("outer", (0, 0), {"top", "left"}),
    ("outer", (1, 1), set()),                     # middle cell of a 3x3 block
    ("outer", (2, 2), {"bottom", "right"}),
    ("inner", (0, 0), {"bottom", "right"}),
    ("inner", (1, 1), {"top", "bottom", "left", "right"}),
    ("horizontal", (1, 0), {"top", "bottom"}),
    ("vertical", (0, 1), {"left", "right"}),
    ("top", (0, 2), {"top"}),
    ("top", (1, 2), set()),
    ("left", (2, 0), {"left"}),
])
def test_sides_for(kind, cell, want):
    assert borders.sides_for(kind, (0, 0, 2, 2), *cell) == want


def test_load_and_draw_file_borders(app, tmp_path):
    g = open_in(app, workbooks.layout(tmp_path / "m.xlsx"))     # B3 has a thin border all round
    assert border(g, 2, 1) == {s: ["thin", "#000000"] for s in ("left", "right", "top", "bottom")}
    assert len(lines(g)) == 4


def test_apply_outer_then_clear_with_undo(app):
    g = app.table
    g.select(1, 1, 3, 3)
    app.fmtbar.borders.apply("outer")
    assert border(g, 1, 1) == {"top": ["thin", "#000000"], "left": ["thin", "#000000"]}
    assert border(g, 2, 2) == {}
    assert border(g, 3, 3) == {"bottom": ["thin", "#000000"], "right": ["thin", "#000000"]}
    assert len(lines(g)) == 12                                 # 3 cells on each of 4 edges
    g.select(1, 1, 3, 3)
    app.fmtbar.borders.apply("none")
    assert not any(border(g, r, c) for r in range(1, 4) for c in range(1, 4)) and not lines(g)
    app.undo()
    assert border(g, 3, 3)
    app.undo()
    assert not border(g, 3, 3) and app.hpos == 0


def test_style_and_colour(app, monkeypatch):
    g = app.table
    bb = app.fmtbar.borders
    bb.style.set("thick")
    monkeypatch.setattr(borders.colorchooser, "askcolor", lambda *a, **k: ((0, 0, 0), "#c00000"))
    bb.pick_color()
    g.goto(0, 0)
    bb.apply("bottom")
    assert border(g, 0, 0) == {"bottom": ["thick", "#c00000"]}
    widths = [g.cv.itemcget(i, "width") for i in g.cv.find_withtag("border")]
    assert widths == ["3.0"]
    bb.style.set("double")
    bb.apply("top")
    assert len(lines(g)) == 3                                 # double = two lines


def test_popup_opens_and_picks(app):
    g = app.table
    bb = app.fmtbar.borders
    g.select(0, 0, 1, 1)
    bb.button.event_generate("<Enter>", when="now")
    app.update()
    assert bb.popup is not None
    grid = bb.popup.winfo_children()[0].winfo_children()[0]
    buttons = grid.winfo_children()
    assert len(buttons) == len(borders.PLACEMENTS)
    buttons[0].invoke()                                       # "all borders"
    app.update()
    assert bb.popup is None
    for cell in ((0, 0), (0, 1), (1, 0), (1, 1)):
        assert set(border(g, *cell)) == {"top", "bottom", "left", "right"}


def test_save_and_reopen(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    bb = app.fmtbar.borders
    bb.style.set("dashed")
    bb.color = "#1f4e79"
    g.select(0, 0, 1, 2)
    bb.apply("all")
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Data"]
    side = ws["B2"].border.right
    assert side.style == "dashed" and side.color.rgb == "FF1F4E79"
    assert io.load(str(path))["Data"].styles[(1, 1)]["border"]["right"] == ["dashed", "#1f4e79"]
    assert ws["A1"].font.b                                    # other formatting kept


def test_clear_formatting_removes_borders(app):
    g = app.table
    g.goto(0, 0)
    app.fmtbar.borders.apply("all")
    app.fmtbar.clear()
    assert border(g, 0, 0) == {}


def test_paste_format_and_fill_copy_borders(app):
    g = app.table
    app.set_cells([(0, 0, "1")])
    g.goto(0, 0)
    app.fmtbar.borders.apply("all")
    app.fill((0, 0, 0, 0), "down", 2)
    assert border(g, 2, 0) == border(g, 0, 0)
    g.goto(0, 0)
    g.copy()
    g.goto(0, 3)
    g.paste("format")
    assert border(g, 0, 3) == border(g, 0, 0)


def test_dark_mode_black_border_visible(app):
    g = app.table
    g.goto(0, 0)
    app.fmtbar.borders.apply("all")
    app.set_theme(True)
    from simplecore import theme
    colours = {g.cv.itemcget(i, "fill") for i in g.cv.find_withtag("border")}
    assert colours == {theme.DARK["cell_fg"]}                 # black lines drawn light in dark mode
    app.set_theme(False)


def test_formatting_an_empty_block_formats_all_of_it(app):
    """Regression: on empty cells only the first row of the selection got formatted."""
    g = app.table
    g.select(1, 1, 3, 3)
    app.format_cells(lambda st: {**st, "bold": True})
    assert all((g.sheet.styles.get((r, c)) or {}).get("bold") for r in range(1, 4) for c in range(1, 4))
