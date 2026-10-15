"""Sheet / workbook settings: gridlines, headings, formulas view, tab colours, duplicate / copy sheet,
protection, document properties, calculation mode — all saved and undoable."""
import openpyxl

import sheet_props
import workbook_tools as wt
import workbooks
from conftest import open_in


def test_gridlines_headings_saved(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    wt.toggle_sheet_setting(app, "grid")
    wt.toggle_sheet_setting(app, "headings")
    assert g.sheet.props["grid"] is False and g.sheet.props["headings"] is False
    assert g.hdr_w() == 0 and g.hdr_h() == 0 and app.dirty()
    app.update()
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws.sheet_view.showGridLines is False and ws.sheet_view.showRowColHeaders is False
    g2 = open_in(app, path)
    assert g2.sheet.props.get("grid") is False and not app.dirty()
    app.undo()   # nothing to undo after a fresh open
    wt.toggle_sheet_setting(app, "grid")
    assert "grid" not in g2.sheet.props
    app.undo()
    assert g2.sheet.props["grid"] is False


def test_show_formulas(app):
    app.set_cells([(0, 0, "2"), (0, 1, "=A1*3")])
    assert app.table.value(0, 1) == "6"
    wt.toggle_formulas(app)
    assert app.table.value(0, 1) == "=A1*3"
    wt.toggle_formulas(app)
    assert app.table.value(0, 1) == "6"


def test_full_screen_toggle(app):
    wt.toggle_full_screen(app)
    assert app.full_var.get()
    wt.toggle_full_screen(app)
    assert not app.full_var.get()


def test_tab_colour(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    open_in(app, path)
    name = app.sheet
    wt.set_tab_color(app, name, "#e06666")
    assert app.tabs.swatches and app.sheets[name].props["tab"] == "#e06666"
    assert app.save()
    assert openpyxl.load_workbook(path)[name].sheet_properties.tabColor.rgb == "FFE06666"
    open_in(app, path)
    assert app.sheets[name].props["tab"] == "#e06666"


def test_duplicate_and_copy_sheet(app, tmp_path):
    path = workbooks.styled(tmp_path / "s.xlsx")
    g = open_in(app, path)
    src = app.sheet
    new = wt.duplicate_sheet(app, src)
    assert new == f"{src} (2)" and app.sheet == new
    assert app.sheets[new].rows == app.sheets[src].rows and app.sheets[new].styles == app.sheets[src].styles
    app.set_cells([(0, 0, "changed")])
    assert app.sheets[src].rows[0][0] != "changed"         # a real copy
    assert app.save()
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == [src, new] and wb[new]["A1"].value == "changed"
    app.undo()
    app.undo()
    assert new not in app.sheets
    dst = wt.copy_to_new_workbook(app, src, str(tmp_path / "one.xlsx"))
    wb2 = openpyxl.load_workbook(dst)
    assert wb2.sheetnames == [src] and wb2[src]["A1"].value == g.text(0, 0)


def test_protection(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.goto(1, 1)
    wt.toggle_locked(app)                                     # B2 stays editable
    assert g.style(1, 1)["unlocked"]
    assert wt.protect_sheet(app, "secret")
    n = len(app.dialogs)
    app.set_cells([(0, 0, "nope")])
    assert g.text(0, 0) != "nope" and app.dialogs[n][1] == "Protected sheet"
    g.goto(0, 0)
    g.edit("x")
    assert g.editor is None                                   # locked cell: no editing
    app.set_cells([(1, 1, "ok")])
    assert g.text(1, 1) == "ok"
    app.restructure("row", 0, 1)
    assert g.text(1, 1) == "ok"                               # no row inserted
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws.protection.sheet and ws["B2"].protection.locked is False and ws["A1"].protection.locked
    g = open_in(app, path)
    assert not wt.unprotect_sheet(app, "wrong") and g.sheet.props.get("protect")
    assert wt.unprotect_sheet(app, "secret") and not g.sheet.props.get("protect")
    app.set_cells([(0, 0, "free")])
    assert g.text(0, 0) == "free"


def test_protected_sheet_cannot_be_deleted_renamed_hidden(app):
    app.add_sheet()
    name = app.sheet
    assert wt.protect_sheet(app, "")
    n = len(app.dialogs)
    assert app.delete_sheet(name) is False and name in app.sheets
    assert app.dialogs[n][1] == "Protected sheet"
    assert app.rename_sheet(name, "Other") is False and name in app.sheets
    assert app.hide_sheet(name) is False and not app.sheets[name].hidden
    app.tabs.start_rename(name)
    assert app.tabs.editor is None                           # no rename box on a protected tab
    from test_hide_sheet import tab_menu
    states = tab_menu(app, name)[0]
    assert states["Delete"] == states["Rename"] == states["Hide"] == "disabled"
    assert wt.unprotect_sheet(app, "")
    assert tab_menu(app, name)[0]["Delete"] == "normal"
    assert app.delete_sheet(name) is not False and name not in app.sheets


def test_modern_password_hash():
    import base64
    import hashlib
    salt = b"0123456789abcdef"
    h = hashlib.sha512(salt + "pw".encode("utf-16-le")).digest()
    for i in range(10):
        h = hashlib.sha512(h + i.to_bytes(4, "little")).digest()
    prot = {"algo": "SHA-512", "hash": base64.b64encode(h).decode(), "salt": base64.b64encode(salt).decode(), "spin": 10}
    assert sheet_props.check_password(prot, "pw") and not sheet_props.check_password(prot, "PW")


def test_document_properties(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    open_in(app, path)
    d = wt.PropertiesDialog(app)
    d.vars["title"].set("Budget 2026")
    d.vars["creator"].set("Ana")
    d.ok()
    assert app.book_props["title"] == "Budget 2026" and app.dirty()
    assert app.save() and not app.dirty()
    wb = openpyxl.load_workbook(path)
    assert wb.properties.title == "Budget 2026" and wb.properties.creator == "Ana"
    open_in(app, path)
    assert app.book_props["title"] == "Budget 2026"


def test_manual_and_iterative_calculation(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    app.set_cells([(10, 0, "1"), (10, 1, "=A11*2")])
    assert g.value(10, 1) == "2"
    d = wt.CalculationDialog(app)
    d.mode.set("manual")
    d.ok()
    app.set_cells([(10, 0, "5")])
    assert g.value(10, 1) == "2"                  # manual: not recalculated yet
    wt.recalculate(app)
    assert g.value(10, 1) == "10"
    d = wt.CalculationDialog(app)
    d.mode.set("auto")
    d.iterate.set(True)
    d.ok()
    app.set_cells([(11, 0, "=B12/2+10"), (11, 1, "=A12")])
    assert abs(float(g.value(11, 0)) - 20) < 0.01
    assert app.save()
    calc = openpyxl.load_workbook(path).calculation
    assert calc.iterate and calc.calcMode == "auto"
