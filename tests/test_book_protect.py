"""Workbook-structure protection: read from the file and enforced (no adding, deleting, renaming, moving, hiding
or copying sheets), Protect / Unprotect workbook with a password, saved back so Excel sees it too."""
import openpyxl
from openpyxl.workbook.protection import WorkbookProtection

import sheet_props
import workbook_tools as wt
from conftest import open_in


def locked_book(path, password="secret"):
    wb = openpyxl.Workbook()
    wb.active.title = "One"
    wb.active["A1"] = 1
    wb.create_sheet("Two")
    wb.security = WorkbookProtection(lockStructure=True)
    if password:
        wb.security.set_workbook_password(password)
    wb.save(path)
    return path


def test_locked_file_blocks_sheet_changes(app, tmp_path):
    open_in(app, locked_book(tmp_path / "locked.xlsx"))
    assert app.book_props["structure"]["password"] and app.structure_locked(warn=False)
    names = list(app.sheets)
    app.add_sheet()
    assert app.rename_sheet("One", "First") is False
    assert app.delete_sheet("Two") is False and app.hide_sheet("Two") is False
    app.move_sheet("Two", 0)
    assert wt.duplicate_sheet(app, "One") is None
    assert list(app.sheets) == names and not app.dirty()                      # nothing changed


def test_unprotect_with_password_then_save(app, tmp_path):
    open_in(app, locked_book(tmp_path / "locked.xlsx"))
    assert not wt.unprotect_book(app, "wrong") and app.structure_locked(warn=False)
    assert wt.unprotect_book(app, "secret") and not app.structure_locked(warn=False)
    app.add_sheet()
    assert len(app.sheets) == 3
    out = tmp_path / "open.xlsx"
    assert app.write(str(out)), app.dialogs
    wb = openpyxl.load_workbook(out)
    assert not (wb.security and wb.security.lockStructure) and len(wb.sheetnames) == 3


def test_protect_book_saved_and_read_back(app, tmp_path):
    src = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 5
    wb.save(src)
    open_in(app, src)
    assert wt.protect_book(app, "pw") and app.structure_locked(warn=False)
    app.add_sheet()
    assert len(app.sheets) == 1
    app.undo()                                                                 # (undo takes the lock off)
    assert not app.structure_locked(warn=False)
    app.redo()
    out = tmp_path / "locked.xlsx"
    assert app.write(str(out)), app.dialogs
    wb = openpyxl.load_workbook(out)
    assert wb.security.lockStructure and wb.security.workbookPassword
    prot = sheet_props.read_book(str(out))["structure"]
    assert sheet_props.check_password(prot, "pw") and not sheet_props.check_password(prot, "nope")


def test_no_password_lock(app, tmp_path):
    open_in(app, locked_book(tmp_path / "nopw.xlsx", password=None))
    assert app.book_props["structure"] == {} and app.structure_locked(warn=False)
    assert wt.unprotect_book(app) and not app.structure_locked(warn=False)       # no password asked
