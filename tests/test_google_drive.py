"""Google sign-in (OAuth with PKCE on a local address) and Google Drive open / save, against a fake Google."""
import base64
import hashlib
import io
import json
import threading
import time
import urllib.parse
import urllib.request

import openpyxl
import pytest
from simplecore import google_drive as gd
from simplecore import settings

ID_TOKEN = "x." + base64.urlsafe_b64encode(json.dumps({"email": "adam@example.com"}).encode()).decode().rstrip("=") + ".y"
FILES = [{"id": "sheet1", "name": "Budget 2026", "mimeType": gd.NATIVE, "modifiedTime": "2026-10-15T09:30:00Z"},
         {"id": "xl1", "name": "Stock.xlsx", "mimeType": gd.EXPORT, "modifiedTime": "2026-10-14T08:00:00Z"}]


def xlsx_bytes(value):
    wb = openpyxl.Workbook()
    wb.active["A1"] = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class FakeGoogle:
    def __init__(self):
        self.calls, self.verifier_ok, self.uploads = [], None, []
        self.challenge = None

    def __call__(self, method, url, token=None, data=None, headers=None, timeout=60):
        self.calls.append((method, url, token, data, headers))
        if url == gd.TOKEN_URL:
            form = dict(urllib.parse.parse_qsl(data.decode()))
            if form["grant_type"] == "authorization_code":
                want = base64.urlsafe_b64encode(hashlib.sha256(form["code_verifier"].encode()).digest()).rstrip(b"=")
                self.verifier_ok = want.decode() == self.challenge and form["code"] == "the-code"
                body = {"access_token": "at-1", "refresh_token": "rt-1", "expires_in": 3600, "id_token": ID_TOKEN}
            else:
                body = {"access_token": "at-2", "expires_in": 3600}
            return 200, json.dumps(body).encode()
        if url.startswith(gd.REVOKE_URL):
            return 200, b""
        assert token, "Drive calls carry the access token"
        if url.startswith(gd.DRIVE + "/files?"):
            q = dict(urllib.parse.parse_qsl(urllib.parse.urlsplit(url).query))["q"]
            hits = [f for f in FILES if "name contains" not in q or q.split("'")[-2].lower() in f["name"].lower()]
            return 200, json.dumps({"files": hits}).encode()
        if "/export?" in url:
            return 200, xlsx_bytes("from Google Sheets")
        if "alt=media" in url:
            return 200, xlsx_bytes("from an xlsx")
        if method == "PATCH":
            self.uploads.append((url, data, headers))
            return 200, b"{}"
        if method == "POST" and url.startswith(gd.UPLOAD + "/files?uploadType=multipart"):
            self.uploads.append((url, data, headers))
            return 200, json.dumps({"id": "new1", "name": "Copy", "mimeType": gd.NATIVE}).encode()
        raise AssertionError(url)


@pytest.fixture
def google(monkeypatch):
    fake = FakeGoogle()
    monkeypatch.setattr(gd, "transport", fake)
    gd.set_client("123.apps.googleusercontent.com", "shh-secret")
    return fake


def browser_that_signs_in(fake, answer=None):
    """Stands in for the browser: Google's page 'redirects' back to the app's local address."""
    def open_browser(url):
        q = dict(urllib.parse.parse_qsl(urllib.parse.urlsplit(url).query))
        assert q["client_id"] == "123.apps.googleusercontent.com" and q["code_challenge_method"] == "S256"
        assert q["redirect_uri"].startswith("http://127.0.0.1:") and "drive" in q["scope"]
        fake.challenge = q["code_challenge"]
        back = answer or {"code": "the-code", "state": q["state"]}

        def go():                    # (slow CI machines: no proxy for the local address, a few tries)
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
            for _ in range(5):
                time.sleep(0.1)
                try:
                    with opener.open(q["redirect_uri"] + "/?" + urllib.parse.urlencode(back), timeout=20) as r:
                        fake.page = r.read().decode()
                    return
                except OSError:
                    continue
        threading.Thread(target=go, daemon=True).start()
    return open_browser


def test_sign_in_with_pkce_and_keep_it_encrypted(google):
    assert not gd.signed_in()
    who = gd.sign_in(open_browser=browser_that_signs_in(google), timeout=60)
    assert who == "adam@example.com" and gd.signed_in() and google.verifier_ok
    end = time.time() + 5                     # the 'browser' reads the page on its own thread (slow CI machines)
    while not hasattr(google, "page") and time.time() < end:
        time.sleep(0.02)
    assert "Signed in" in google.page
    raw = settings.get("google_refresh")
    assert "rt-1" not in raw and "shh-secret" not in settings.get("google_client_secret")   # encrypted
    assert gd.token() == "at-1"                                          # still fresh
    gd._access["until"] = 0
    assert gd.token() == "at-2"                                          # renewed with the refresh token
    gd.sign_out()
    assert not gd.signed_in() and gd.email() == ""
    with pytest.raises(gd.GoogleError, match="Sign in"):
        gd.token()


def test_sign_in_refuses_a_wrong_state(google):
    with pytest.raises(gd.GoogleError, match="didn't match"):
        gd.sign_in(open_browser=browser_that_signs_in(google, {"code": "x", "state": "forged"}), timeout=60)
    assert not gd.signed_in()


def test_sign_in_needs_a_client_and_can_be_cancelled(google):
    gd.set_client("", "")
    with pytest.raises(gd.GoogleError, match="Client ID"):
        gd.sign_in(open_browser=lambda url: None)
    gd.set_client("123.apps.googleusercontent.com", "s")
    with pytest.raises(gd.GoogleError, match="cancelled"):
        gd.sign_in(open_browser=lambda url: None, cancelled=lambda: True)


def signed(google):
    gd.sign_in(open_browser=browser_that_signs_in(google), timeout=60)


def test_list_open_and_save_back(app, google):
    signed(google)
    assert [f["name"] for f in gd.list_files()] == ["Budget 2026", "Stock.xlsx"]
    assert [f["name"] for f in gd.list_files("stock")] == ["Stock.xlsx"]
    sheet = FILES[0]
    path = gd.download(sheet)
    assert path.endswith("Budget 2026.xlsx")
    gd.open_drive_file(app, sheet, path)
    assert app.drive_file["id"] == "sheet1" and app.table.text(0, 0) == "from Google Sheets"
    app.set_cells([(1, 0, "changed")])
    assert app.save()
    end = time.monotonic() + 5
    while not google.uploads and time.monotonic() < end:                 # goes up in the background
        app.update()
        time.sleep(0.02)
    app.update()
    url, data, headers = google.uploads[-1]
    assert "/files/sheet1?uploadType=media" in url and headers["Content-Type"] == gd.EXPORT
    assert openpyxl.load_workbook(io.BytesIO(data)).active["A2"].value == "changed"
    for _ in range(10):
        app.update()
        time.sleep(0.02)
    assert "Saved Budget 2026 to Google Drive" in app.status["text"]


def test_other_files_dont_go_to_drive(app, google, tmp_path):
    signed(google)
    f = FILES[1]
    gd.open_drive_file(app, f, gd.download(f))
    other = tmp_path / "local.xlsx"
    openpyxl.Workbook().save(other)
    app.open_file(str(other))
    assert app.drive_file is None
    app.set_cells([(0, 0, "x")])
    app.save()
    app.update()
    assert not google.uploads


def test_save_as_new_google_sheet(app, google, tmp_path):
    signed(google)
    path = tmp_path / "Report.xlsx"
    openpyxl.Workbook().save(path)
    made = gd.create(str(path), "Report", as_native=True)
    assert made["id"] == "new1"
    _url, body, headers = google.uploads[-1]
    assert headers["Content-Type"].startswith("multipart/related; boundary=")
    assert b'"mimeType": "' + gd.NATIVE.encode() + b'"' in body and b'"name": "Report"' in body


def test_account_dialog(app, google, monkeypatch):
    monkeypatch.setattr(gd, "sign_in", lambda cancelled=None: (time.sleep(0.05), "adam@example.com")[1])
    d = gd.AccountDialog(app)
    d.cid.set("999.apps.googleusercontent.com")
    d.toggle()                                                        # Sign in (in the background)
    assert d.waiting and d.btn["text"] == "Cancel"
    end = time.monotonic() + 3
    while d.waiting and time.monotonic() < end:
        app.update()
        time.sleep(0.02)
    assert gd.client()[0] == "999.apps.googleusercontent.com"         # the client was saved first
    d.close()


def test_menu_has_google_drive(app):
    import tkinter as tk
    menus = [w for w in app.winfo_children() if isinstance(w, tk.Menu)]
    found = []

    def walk(m):
        for i in range((m.index("end") or -1) + 1):
            if m.type(i) == "cascade":
                found.append(m.entrycget(i, "label"))
                walk(m.nametowidget(m.entrycget(i, "menu")))
            elif m.type(i) == "command":
                found.append(m.entrycget(i, "label"))
    for w in app.winfo_children():
        for m in [c for c in w.winfo_children() if isinstance(c, tk.Menu)] + menus:
            walk(m)
    if "Google Drive" not in found:                                    # menus hang under the ☰ button
        import app_menu
        root = tk.Frame(app)
        app_menu.build(app, root)
        for mb in root.winfo_children():
            if mb.winfo_class() == "TMenubutton":
                walk(mb.nametowidget(mb["menu"]))
        root.destroy()
    assert {"Google Drive", "Open from Google Drive…", "Save to Google Drive…", "Google account…"} <= set(found)
