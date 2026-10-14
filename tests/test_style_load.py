"""Every cell's own style must come back when a file is opened (regression: styles were cached by
id() of temporary openpyxl proxies, so after a save some cells reopened with another cell's / no style)."""
import openpyxl
import pytest
from openpyxl.styles import Alignment, Font, PatternFill

import sheet_io as io


def many_styles(path, n=400):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    expected = {}
    for i in range(n):
        r, c = divmod(i, 20)
        cell = ws.cell(row=r + 1, column=c + 1, value=f"v{i}")
        colour = f"{(i * 37) % 256:02X}{(i * 91) % 256:02X}{(i * 53) % 256:02X}"
        cell.fill = PatternFill("solid", fgColor="FF" + colour)
        cell.font = Font(bold=i % 2 == 0, italic=i % 3 == 0, sz=10 + i % 7)
        cell.alignment = Alignment(horizontal=("left", "center", "right")[i % 3])
        expected[(r, c)] = ("#" + colour.lower(), i % 2 == 0, i % 3 == 0, 10 + i % 7,
                            ("left", "center", "right")[i % 3])
    wb.save(path)
    return expected


def check(sheet, expected):
    wrong = []
    for (r, c), (bg, bold, italic, size, align) in expected.items():
        st = sheet.styles.get((r, c)) or {}
        got = (st.get("bg"), bool(st.get("bold")), bool(st.get("italic")), st.get("size", 11.0), st.get("align"))
        if got != (bg, bold, italic, float(size), align):
            wrong.append(((r, c), got))
    assert not wrong, f"{len(wrong)} cells loaded the wrong style, e.g. {wrong[:3]}"


def test_every_cell_keeps_its_style(tmp_path):
    p = tmp_path / "many.xlsx"
    expected = many_styles(p)
    check(io.load(str(p))["S"], expected)


def test_every_cell_keeps_its_style_big_file_mode(tmp_path, monkeypatch):
    p = tmp_path / "many.xlsx"
    expected = many_styles(p)
    monkeypatch.setattr(io, "FULL_LOAD_BYTES", 0)        # force the read-only (large file) loader
    check(io.load(str(p))["S"], expected)


@pytest.mark.parametrize("rounds", [3])
def test_styles_survive_repeated_save_and_reopen(app, tmp_path, rounds):
    from conftest import open_in
    p = tmp_path / "many.xlsx"
    expected = many_styles(p, n=200)
    for k in range(rounds):
        g = open_in(app, p)
        app.set_cells([(30 + k, 0, f"edit {k}")])     # any change, then save
        assert app.save(), app.dialogs
    check(io.load(str(p))["S"], expected)
