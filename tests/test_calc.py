"""Formulas inside the app: recalculation, paste shifting, save as xlsx/csv."""
import openpyxl

import workbooks
from conftest import open_in


def test_calculation_flow(app, tmp_path):
    path = workbooks.budget(tmp_path / "b.xlsx")
    g = open_in(app, path)
    assert g.value(1, 3) == "25" and g.value(5, 3) == "104.97"
    g.goto(5, 3)
    assert app.value.get() == "=SUM(D2:D4)"
    app.set_cells([(1, 1, "20")])
    assert g.value(1, 3) == "50" and g.value(5, 3) == "129.97"

    # copy a formula one row down: references move, format comes along, one undo step
    app.set_cells([(4, 0, "Tape"), (4, 1, "4"), (4, 2, "1.25")])
    g.goto(3, 3)
    g.copy()
    g.goto(4, 3)
    g.paste()
    assert g.text(4, 3) == "=B5*C5" and g.value(4, 3) == "5"
    assert (g.sheet.styles.get((4, 3)) or {}).get("numfmt") == '"$"#,##0.00'
    app.undo()
    assert g.text(4, 3) == "" and not g.sheet.styles.get((4, 3))
    app.redo()
    assert g.value(5, 3) == "129.97"   # like Excel, the SUM range is not auto-extended
    app.set_cells([(5, 3, "=SUM(D2:D5)")])
    assert g.value(5, 3) == "134.97"

    # one formula pasted into a selection fills it
    app.set_cells([(1, 4, "=D2*2")])
    g.goto(1, 4)
    g.copy()
    g.select(2, 4, 4, 4)
    g.paste()
    assert [g.text(r, 4) for r in range(2, 5)] == ["=D3*2", "=D4*2", "=D5*2"]

    app.set_cells([(6, 3, "=D2/0")])
    assert g.value(6, 3) == "#DIV/0!"

    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws["D6"].value == "=SUM(D2:D5)" and ws["D5"].value == "=B5*C5" and ws["B2"].value == 20
    csv_path = tmp_path / "b.csv"
    assert app.write(str(csv_path))
    lines = csv_path.read_text(encoding="utf-8").splitlines()
    assert lines[1] == "Paper,20,2.5,50,100" and lines[5] == ",,Sum,134.97"
