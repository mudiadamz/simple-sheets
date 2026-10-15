"""Excel tables (banding, ▼ filter, growing) and printing with the file's page setup."""
import re

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

import pdf_export
import tables
from conftest import open_in


def table_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Item", "Qty"], ["a", 1], ["b", 2], ["a", 3]]:
        ws.append(row)
    t = Table(displayName="Stock", ref="A1:B4")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws["D1"] = "=SUM(Stock[Qty])"
    wb.save(path)
    return path


def test_banding_and_header(app, tmp_path):
    g = open_in(app, table_file(tmp_path / "t.xlsx"))
    head = tables.look(g.sheet, 0, 0)
    assert head["bold"] and head["bg"] == "#4472c4" and head["fg"] == "#ffffff"
    assert tables.look(g.sheet, 1, 0)["bg"] and tables.look(g.sheet, 2, 0) is None   # stripes
    assert tables.look(g.sheet, 5, 5) is None
    app.update()
    g.redraw()
    assert len(g.table_buttons) == 2


def test_table_grows_and_saves(app, tmp_path):
    path = table_file(tmp_path / "t.xlsx")
    g = open_in(app, path)
    app.set_cells([(4, 0, "c"), (4, 1, "10")])
    assert g.sheet.extra["tables"][0]["range"] == (0, 0, 4, 1)
    assert g.value(0, 3) == "16"                                # Stock[Qty] includes the new row
    app.undo()
    assert g.sheet.extra["tables"][0]["range"] == (0, 0, 3, 1)
    app.redo()
    assert app.save()
    assert openpyxl.load_workbook(path)["Data"].tables["Stock"].ref == "A1:B5"


def test_table_filter(app, tmp_path):
    path = table_file(tmp_path / "t.xlsx")
    g = open_in(app, path)
    t = g.sheet.extra["tables"][0]
    tables.start_filter(app, t)
    import data_tools
    data_tools.set_criteria(app, 0, ["b"])
    assert g.sheet.filtered_rows == {1, 3} and g.sheet.filter["table"] == "Stock"
    assert app.save()
    ws = openpyxl.load_workbook(path)["Data"]
    assert ws.auto_filter.ref is None                          # not a sheet filter: the table's own
    fc = ws.tables["Stock"].autoFilter.filterColumn
    assert fc and fc[0].filters.filter == ["b"]


def page_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 121):
        ws.append([f"r{r}", r, r * 2])
    ws["A1"], ws["B1"], ws["C1"] = "Name", "Value", "Double"
    ws.print_title_rows = "1:1"
    ws.print_area = "A1:B120"
    ws.page_setup.orientation = "landscape"
    ws.print_options.gridLines = True
    ws.page_margins.left = 0.25
    wb.save(path)
    return path


def test_pdf_uses_page_setup(app, tmp_path):
    open_in(app, page_file(tmp_path / "p.xlsx"))
    page = app.current.props["page"]
    assert page["orientation"] == "landscape" and page["title_rows"] == (0, 0) and page["area"] == (0, 0, 119, 1)
    dst = str(tmp_path / "p.pdf")
    assert app.export_pdf(dst)
    data = open(dst, "rb").read()
    pages = data.count(b"/Type /Page ")
    assert pages >= 3 and b"/MediaBox [0 0 842 595]" in data
    texts = re.findall(rb"\(((?:[^()\\]|\\.)*)\) Tj", data)
    assert texts.count(b"Name") == pages                       # title row on every page
    assert b"Double" not in texts                              # column C is outside the print area
    assert b"0.85 0.85 0.85 RG" in data                       # gridlines asked for
    assert re.search(rb"q [\d.]+ 0 0 [\d.]+ 18\.00 ", data)              # 0.25 inch left margin
    assert pdf_export.A4
