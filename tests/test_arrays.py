"""Dynamic arrays (spill), array functions, structured table references, links to other workbooks,
iterative calculation."""
import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from conftest import open_in
from formula import Engine
from sheet_io import Sheet


def book(rows, **more):
    s = Sheet([list(r) for r in rows])
    e = Engine()
    e.set_book({"S": s, **more})
    return s, e


def shown(s, e, r, c):
    t = s.rows[r][c] if r < len(s.rows) and c < len(s.rows[r]) else ""
    if t.startswith("="):
        return e.display(s, r, c)
    return e.spilled(s, r, c) or t


def test_sequence_spills_and_feeds_other_formulas():
    s, e = book([["=SEQUENCE(3,2)", "", "=SUM(A1#)"], ["", "", "=B3*10"], ["", "", ""]])
    assert [[shown(s, e, r, c) for c in range(2)] for r in range(3)] == [["1", "2"], ["3", "4"], ["5", "6"]]
    assert shown(s, e, 0, 2) == "21" and shown(s, e, 1, 2) == "60"


def test_blocked_spill():
    s, e = book([["=SEQUENCE(3)"], [""], ["x"]])
    assert shown(s, e, 0, 0) == "#SPILL!" and shown(s, e, 1, 0) == ""


def test_range_formula_and_elementwise_math_spill():
    s, e = book([["1", "=A1:A3*2"], ["2", ""], ["3", ""]])
    assert [shown(s, e, r, 1) for r in range(3)] == ["2", "4", "6"]


@pytest.mark.parametrize("formula, want", [
    ('=FILTER(A2:B5,B2:B5>15)', [["b", "20"], ["d", "40"]]),
    ('=FILTER(A2:A5,B2:B5>99,"none")', [["none"]]),
    ('=UNIQUE(C2:C5)', [["x"], ["y"]]),
    ('=SORT(B2:B5,1,-1)', [["40"], ["20"], ["15"], ["10"]]),
    ('=SORTBY(A2:A5,B2:B5,-1)', [["d"], ["b"], ["c"], ["a"]]),
    ('=TRANSPOSE(A2:A4)', [["a", "b", "c"]]),
    ('=TAKE(A2:A5,-2)', [["c"], ["d"]]),
    ('=DROP(A2:A5,3)', [["d"]]),
    ('=CHOOSECOLS(A2:C3,3,1)', [["x", "a"], ["y", "b"]]),
    ('=VSTACK(A2:A3,{"z"})', [["a"], ["b"], ["z"]]),
    ('=HSTACK(A2:A3,B2:B3)', [["a", "10"], ["b", "20"]]),
    ('=TOROW(A2:B3)', [["a", "10", "b", "20"]]),
    ('=TEXTSPLIT("a,b;c,d",",",";")', [["a", "b"], ["c", "d"]]),
    ('=MMULT({1,2;3,4},{1;1})', [["3"], ["7"]]),
    ('=SUM(SEQUENCE(4))', [["10"]]),
    ('=ROWS(SEQUENCE(4))', [["4"]]),
])
def test_array_functions(formula, want):
    rows = [["k", "v", "t"], ["a", "10", "x"], ["b", "20", "y"], ["c", "15", "x"], ["d", "40", "y"]]
    rows = [r + ["", "", "", ""] for r in rows] + [[""] * 7 for _ in range(5)]
    rows[0][4] = formula
    s, e = book(rows)
    got = [[shown(s, e, r, 4 + c) for c in range(len(want[0]))] for r in range(len(want))]
    assert got == want


def tables_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Item", "Qty", "Price"], ["a", 2, 1.5], ["b", 3, 2], ["c", 4, 0.5]]:
        ws.append(row)
    ws["D1"] = "Total"
    for r in range(2, 5):
        ws[f"D{r}"] = "=[@Qty]*[@Price]"
    ws.add_table(Table(displayName="Sales", ref="A1:D4"))
    ws["F1"] = "=SUM(Sales[Qty])"
    ws["F2"] = "=ROWS(Sales)"
    ws["F3"] = "=COUNTA(Sales[#All])"
    ws["F4"] = "=SUM(Sales[Total])"
    wb.save(path)
    return path


def test_structured_references(app, tmp_path):
    g = open_in(app, tables_file(tmp_path / "t.xlsx"))
    assert [g.value(r, 3) for r in range(1, 4)] == ["3", "6", "2"]
    assert g.value(0, 5) == "9" and g.value(1, 5) == "3" and g.value(2, 5) == "16" and g.value(3, 5) == "11"
    app.set_cells([(1, 1, "10")])
    assert g.value(0, 5) == "17" and g.value(1, 3) == "15"


def test_external_link_reads_other_file(app, tmp_path):
    other = openpyxl.Workbook()
    other.active.title = "Prices"
    other.active["B2"] = 42
    other.active["B3"] = "=B2*2"
    other.save(tmp_path / "Prices.xlsx")
    main = openpyxl.Workbook()
    main.active["A1"] = "=1"
    main.save(tmp_path / "main.xlsx")
    g = open_in(app, tmp_path / "main.xlsx")
    app.set_cells([(0, 0, "='[Prices.xlsx]Prices'!B2+1"), (1, 0, "=SUM([Prices.xlsx]Prices!B2:B3)"),
                   (2, 0, "='[Missing.xlsx]Prices'!B2")])
    assert g.value(0, 0) == "43" and g.value(1, 0) == "126" and g.value(2, 0) == "#REF!"


def test_external_link_uses_stored_values(tmp_path):
    import zipfile

    import externals
    p = tmp_path / "linked.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("xl/externalLinks/externalLink1.xml",
                   '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><externalBook>'
                   '<sheetNames><sheetName val="Rates"/></sheetNames><sheetDataSet><sheetData sheetId="0">'
                   '<row r="1"><cell r="A1"><v>0.25</v></cell><cell r="B1" t="str"><v>EUR</v></cell></row>'
                   '</sheetData></sheetDataSet></externalBook></externalLink>')
        z.writestr("xl/externalLinks/_rels/externalLink1.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Target="gone/Rates.xlsx"/></Relationships>')

    class App:
        pass
    links = externals.Links(App(), str(p))
    assert links.value("[1]Rates!A1") == 0.25 and links.value("[1]Rates!B1") == "EUR"


def test_iterative_calculation():
    s, e = book([["=B1/2+10", "=A1"]])
    assert e.display(s, 0, 0) == "#CIRC!"
    e.set_iteration({"iterate": True, "iterate_count": 100, "iterate_delta": 0.0001})
    assert abs(float(e.display(s, 0, 0)) - 20) < 0.01


def test_array_constant_and_spill_ref():
    s, e = book([['={1,2;3,4}', "", "=SUM(A1#)", "=ROWS(A1#)"], ["", "", "", ""]])
    assert shown(s, e, 0, 0) == "1" and shown(s, e, 1, 1) == "4"
    assert shown(s, e, 0, 2) == "10" and shown(s, e, 0, 3) == "2"
