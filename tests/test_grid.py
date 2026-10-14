"""Selection, editing, clipboard, undo/redo, zoom and saving through the real window."""
import openpyxl

import workbooks
from conftest import key, open_in


def test_window_opens_maximized(app):
    assert app.opened_state == "zoomed"


def test_no_view_toggles(app):
    """Formatting and formula results are always on: no checkboxes for them."""
    def texts(w):
        out = []
        for child in w.winfo_children():
            try:
                out.append(str(child.cget("text")))
            except Exception:
                pass
            out += texts(child)
        return out
    labels = texts(app)
    assert "Show formatting" not in labels and "Show formulas" not in labels


def test_formatting_and_results_always_shown(app, tmp_path):
    g = open_in(app, workbooks.budget(tmp_path / "b.xlsx"))
    assert g.style(1, 3) == {"numfmt": '"$"#,##0.00'}
    assert g.value(1, 3) == "25"   # result, not the formula


def test_open_shows_formula_and_tabs(app, tmp_path):
    g = open_in(app, workbooks.plain(tmp_path / "p.xlsx"))
    assert g.text(1, 3) == "=B2*C2" and g.value(1, 3) == "4.5"
    assert app.tabs.names() == ["Data", "Other"]
    app.tabs.nb.select(1)
    app.update()
    assert app.sheet == "Other"


def test_edit_undo_redo_save(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.goto(1, 1)
    g.edit("9")
    g.commit(1, 0)
    assert g.text(1, 1) == "9" and g.cur == (2, 1)
    assert app.dirty() and app.title().startswith("Simple Sheets - *")
    app.undo()
    assert g.text(1, 1) == "3" and not app.dirty()
    app.redo()
    assert g.text(1, 1) == "9"
    assert app.save() and not app.dirty()
    ws = openpyxl.load_workbook(path)["Data"]
    assert ws["B2"].value == 9 and ws["D2"].value == "=B2*C2" and ws["A1"].font.b


def test_unsaved_prompt_cancel_keeps_window(app, tmp_path, monkeypatch):
    import simple_sheets
    g = open_in(app, workbooks.plain(tmp_path / "p.xlsx"))
    app.set_cells([(0, 0, "zz")])
    monkeypatch.setattr(simple_sheets.messagebox, "askyesnocancel", lambda *a, **k: None)  # Cancel
    app.on_close()
    assert app.winfo_exists()
    app.undo()
    assert not app.dirty()


def test_keyboard_selection_and_copy(app, tmp_path):
    g = open_in(app, workbooks.styled(tmp_path / "s.xlsx"))
    g.goto(1, 0)
    key(app, "<Shift-Right>")
    key(app, "<Shift-Right>")
    key(app, "<Shift-Down>")
    assert g.sel() == (1, 0, 2, 2) and app.ref["text"] == "A2:C3"
    assert "Sum:" in app.status["text"]
    key(app, "<Control-c>")
    assert app.clipboard_get() == "Item\tQty\tPrice\r\nApple\t1200\t0.5"
    g.goto(1, 0)
    key(app, "<Control-Right>")
    assert g.cur == (1, 3)
    key(app, "<Control-Shift-Left>")
    assert g.sel() == (1, 0, 1, 3)


def test_mouse_selection(app, tmp_path):
    g = open_in(app, workbooks.styled(tmp_path / "s.xlsx"))
    g.goto(0, 0)
    x1, y1 = g.col_x(1) + 5, g.row_y(2) + 5   # B3
    x2, y2 = g.col_x(2) + 5, g.row_y(5) + 5   # C6
    g.cv.event_generate("<Button-1>", x=x1, y=y1, when="now")
    g.cv.event_generate("<B1-Motion>", x=x2, y=y2, when="now")
    g.cv.event_generate("<ButtonRelease-1>", x=x2, y=y2, when="now")
    assert g.sel() == (2, 1, 5, 2)
    g.cv.event_generate("<Shift-Button-1>", x=g.col_x(3) + 5, y=g.row_y(3) + 5, when="now")
    assert g.sel() == (2, 1, 3, 3)
    g.cv.event_generate("<Button-1>", x=g.col_x(1) + 5, y=3, when="now")
    g.cv.event_generate("<ButtonRelease-1>", x=g.col_x(1) + 5, y=3, when="now")
    assert app.ref["text"] == "B:B"


def test_paste_delete_undo_groups(app, tmp_path):
    g = open_in(app, workbooks.styled(tmp_path / "s.xlsx"))
    g.goto(5, 0)
    app.clipboard_clear()
    app.clipboard_append("x\ty\r\nz\tw")
    key(app, "<Control-v>")
    assert [g.text(5, 0), g.text(5, 1), g.text(6, 0), g.text(6, 1)] == ["x", "y", "z", "w"]
    key(app, "<Delete>")
    assert g.text(5, 0) == "" and g.text(6, 1) == ""
    app.undo()
    assert g.text(6, 1) == "w"
    app.undo()
    assert g.text(5, 0) == "" and not app.dirty()


def test_formula_bar_and_format_info(app, tmp_path):
    path = workbooks.styled(tmp_path / "s.xlsx")
    g = open_in(app, path)
    g.goto(2, 2)
    assert "$" in app.fmt_info["text"]
    app.value.delete(0, "end")
    app.value.insert(0, "0.75")
    app.value_commit()
    assert g.text(2, 2) == "0.75"
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws["C3"].value == 0.75 and ws["C3"].number_format == '"$"0.00' and "A1:D1" in str(ws.merged_cells)


def test_edit_survives_switching_to_another_program(app, monkeypatch):
    g = app.table
    g.goto(0, 0)
    g.edit("hel")
    app.update()
    e = g.editor
    monkeypatch.setattr(g, "focus_get", lambda: None)          # window lost focus (Alt+Tab)
    g.focus_left(e)
    assert g.editor is e                                        # still editing
    monkeypatch.setattr(g, "focus_get", lambda: app.value)      # clicked the formula bar
    g.focus_left(e)
    assert g.editor is None and g.text(0, 0) == "hel"


def test_tab_rename_survives_switching_to_another_program(app, monkeypatch):
    t = app.tabs
    t.start_rename("Sheet1")
    app.update()
    e = t.editor
    e.delete(0, "end")
    e.insert(0, "Data")
    monkeypatch.setattr(t, "focus_get", lambda: None)
    t.focus_left(e)
    assert t.editor is e
    monkeypatch.setattr(t, "focus_get", lambda: app.table.cv)
    t.focus_left(e)
    assert list(app.sheets) == ["Data"]


def test_zoom(app, tmp_path):
    g = open_in(app, workbooks.plain(tmp_path / "p.xlsx"))
    w0, h0 = g.colw(0), g.rowh
    g.cv.focus_set()
    key(app, "<Control-plus>")
    key(app, "<Control-equal>")
    assert app.zoom_label["text"] == "125%" and g.colw(0) > w0 and g.rowh > h0
    key(app, "<Control-MouseWheel>", delta=-120)
    assert app.zoom_label["text"] == "110%"
    for _ in range(20):
        app.zoom_step(-1)
    assert app.zoom_label["text"] == "50%"
    for _ in range(20):
        app.zoom_step(1)
    assert app.zoom_label["text"] == "300%"
    key(app, "<Control-0>")
    assert app.zoom_label["text"] == "100%" and g.colw(0) == w0 and g.editor is None
