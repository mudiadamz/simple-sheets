"""App-level P3: HTML export, templates, .ods saving, split window, help, settings, translations,
update check, keyboard-only commands, warning before parts of a file would be lost."""
import io as pyio  # noqa: F401
import json
import zipfile

import openpyxl
from simplecore import settings, updates

import help_ui
import html_export
import i18n
import settings_ui
import sheet_io as io
import workbooks
from conftest import open_in


def test_html_export(app, tmp_path):
    g = open_in(app, workbooks.styled(tmp_path / "s.xlsx"))
    app.set_cells([(0, 7, "a < b & c"), (1, 7, "=2*3")])
    dst = html_export.export(app, str(tmp_path / "out.html"))
    page = open(dst, encoding="utf-8").read()
    assert page.startswith("<!doctype html>") and "a &lt; b &amp; c" in page and ">6<" in page
    assert "font-weight:bold" in page and "background:" in page
    assert g.sheet.merges and ("colspan" in page or "rowspan" in page)


def template_file(path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Invoice"
    wb.active["A1"].font = openpyxl.styles.Font(bold=True)
    wb.template = True
    wb.save(path)
    return path


def test_new_from_template(app, tmp_path):
    import templates
    tpl = template_file(tmp_path / "invoice.xltx")
    assert templates.new_from_template(app, str(tpl))
    g = app.table
    assert app.path is None and g.text(0, 0) == "Invoice" and g.style(0, 0)["bold"]
    app.set_cells([(1, 0, "No. 7")])
    assert app.write(str(tmp_path / "inv7.xlsx"))
    wb = openpyxl.load_workbook(tmp_path / "inv7.xlsx")
    assert wb.active["A2"].value == "No. 7" and not wb.template
    assert openpyxl.load_workbook(tpl).active["A2"].value is None          # template untouched
    with zipfile.ZipFile(tmp_path / "inv7.xlsx") as z:
        assert "template" not in z.read("[Content_Types].xml").decode()
    assert app.write(str(tmp_path / "copy.xltx"))
    assert openpyxl.load_workbook(tmp_path / "copy.xltx").template


def test_save_ods_roundtrip(app, tmp_path):
    g = open_in(app, workbooks.styled(tmp_path / "s.xlsx"))
    app.set_cells([(9, 0, "=1+2"), (9, 1, "=SUM(B2:B3)"), (10, 0, "2024-03-05"), (10, 1, "TRUE")])
    dst = str(tmp_path / "out.ods")
    assert app.write(dst)
    book = io.load(dst)
    s = book[app.sheet]
    assert s.rows[9][0] == "=1+2" and s.rows[9][1] == "=SUM(B2:B3)" and s.rows[10][:2] == ["2024-03-05", "TRUE"]
    assert s.styles.get((0, 0), {}).get("bold") == g.style(0, 0).get("bold")
    assert s.merges == g.sheet.merges
    with zipfile.ZipFile(dst) as z:
        assert z.namelist()[0] == "mimetype" and z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
    g2 = open_in(app, dst)
    assert g2.value(9, 0) == "3"
    app.set_cells([(11, 0, "x")])
    assert app.save()                                  # .ods saves in place
    assert io.load(dst)[app.sheet].rows[11][0] == "x"


def test_of_formula():
    import ods_write
    assert ods_write.of_formula("=SUM(A1:B2,'My Sheet'!C3)") == "of:=SUM([.A1:.B2];[$'My Sheet'.C3])"
    assert ods_write.of_formula('=IF(A1>1,"a,b",2)') == 'of:=IF([.A1]>1;"a,b";2)'
    assert ods_write.data_style("#,##0.00", "N1").count('number:decimal-places="2"') == 1
    assert "percentage-style" in ods_write.data_style("0%", "N2")
    assert "date-style" in ods_write.data_style("dd/mm/yyyy", "N3")


def test_split_window(app):
    import split_view
    app.set_cells([(0, 0, "a")])
    assert split_view.toggle(app)
    sg = app.split_grid
    app.update()
    assert sg.sheet is app.table.sheet
    sg.on_press(type("E", (), {"x": sg.col_x(1) + 5, "y": sg.row_y(2) + 5})())
    assert app.table.cur == (2, 1)
    assert sg.edit("x") == "break" and app.table.text(2, 1) == ""
    assert not split_view.toggle(app) and app.split_grid is None


def test_help_windows(app):
    w = help_ui.show(app, "shortcuts")
    assert "Flash fill" in w.text.get("1.0", "end")
    w.destroy()
    w = help_ui.show(app, "guide")
    assert "Simple Sheets" in w.text.get("1.0", "end")
    w.destroy()
    help_ui.show(app, "about")
    assert app.dialogs[-1][0] == "showinfo"


def test_settings_dialog(app, monkeypatch):
    d = settings_ui.SettingsDialog(app)
    d.rows.set("120000")
    d.minutes.set("5")
    d.lang.set("Bahasa Indonesia")
    assert d.ok()
    assert settings.get("max_rows") == 120000 and io.MAX_ROWS == 120000
    import recovery
    assert recovery.INTERVAL_MS == 300_000 and settings.get("lang") == "id"
    settings.set("max_rows", 50_000)
    settings_ui.apply_startup()
    assert io.MAX_ROWS == 50_000


def test_translations(app, monkeypatch):
    import tkinter as tk
    i18n.set_lang("id")
    try:
        assert i18n.tr("Save") == "Simpan" and i18n.tr("Insert 2 rows above") == "Sisipkan 2 baris di atas"
        m = tk.Menu(app, tearoff=False)
        sub = tk.Menu(m, tearoff=False)
        sub.add_command(label="Copy")
        m.add_command(label="Open…")
        m.add_cascade(label="Data", menu=sub)
        i18n.translate_menu(m)
        assert m.entrycget(0, "label") == "Buka…" and sub.entrycget(0, "label") == "Salin"
    finally:
        i18n.set_lang("en")
    assert i18n.tr("Save") == "Save"


def test_update_check(app, tmp_path):
    feed = tmp_path / "latest.json"
    feed.write_text(json.dumps({"tag_name": "v99.0.0", "html_url": "https://example.invalid/r"}))
    assert updates.check(app, feed.as_uri()) is True and app.dialogs[-1][0] == "askokcancel"
    feed.write_text(json.dumps({"version": "0.1"}))
    assert updates.check(app, feed.as_uri()) is False and app.dialogs[-1][0] == "showinfo"
    assert updates.check(app, (tmp_path / "missing.json").as_uri()) is None
    assert updates.check(app) is None                     # no address set
    assert updates.parse_version("v1.10.2") > updates.parse_version("1.9")


def test_keyboard_only(app):
    app.add_sheet()
    first = app.visible_sheets()[0]
    app.next_sheet(-1)
    assert app.sheet == first
    app.next_sheet(1)
    assert app.sheet != first


def test_warns_before_dropping_form_controls(app, tmp_path, monkeypatch):
    path = workbooks.plain(tmp_path / "p.xlsx")
    data = {}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            data[n] = z.read(n)
    data["xl/ctrlProps/ctrlProp1.xml"] = b'<formControlPr objectType="CheckBox"/>'   # a form control
    with zipfile.ZipFile(path, "w") as z:
        for n, b in data.items():
            z.writestr(n, b)
    open_in(app, path)
    assert app.lost_parts == ["form controls (buttons, check boxes…)"]
    app.set_cells([(0, 0, "x")])
    import simple_sheets
    monkeypatch.setattr(simple_sheets.messagebox, "askokcancel", lambda *a, **k: False)
    assert not app.save()
