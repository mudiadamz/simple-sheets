"""Strikethrough button, indent, text rotation, gradient fills."""
import openpyxl
from openpyxl.styles import Alignment, GradientFill

import sheet_io as io
import workbooks
from conftest import open_in


def test_strike_indent_rotate_saved(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.goto(0, 0)
    fb = app.fmtbar
    fb.vars["strike"].set(True)
    fb.toggle("strike")
    fb.indent(1)
    fb.indent(1)
    fb.rotate(-45)
    st = g.style(0, 0)
    assert st["strike"] and st["indent"] == 2 and st["rotate"] == -45
    app.update()
    g.redraw()
    texts = [i for i in g.cv.find_withtag("txt") if g.cv.itemcget(i, "text") == g.value(0, 0)]
    assert texts and float(g.cv.itemcget(texts[0], "angle")) in (-45.0, 315.0)
    fb.indent(-5)
    assert "indent" not in g.style(0, 0)
    fb.indent(3)
    assert app.save()
    cell = openpyxl.load_workbook(path).active["A1"]
    assert cell.font.strike and cell.alignment.indent == 3 and cell.alignment.text_rotation == 135
    s = io.load(str(path))["Data"]
    assert s.styles[(0, 0)]["rotate"] == -45 and s.styles[(0, 0)]["indent"] == 3


def test_gradient_shown_and_kept(app, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "grad"
    ws["A1"].fill = GradientFill(stop=("FF0000", "0000FF"))
    ws["A1"].alignment = Alignment(horizontal="left")
    p = tmp_path / "g.xlsx"
    wb.save(p)
    g = open_in(app, p)
    assert g.style(0, 0)["bg"] == "#800080" and g.style(0, 0)["gradient"]
    app.format_cells(lambda st: {**st, "bold": True})   # restyling keeps the gradient itself
    assert app.save()
    cell = openpyxl.load_workbook(p).active["A1"]
    assert cell.fill.stop and cell.fill.stop[0].color.rgb.endswith("FF0000") and cell.font.b
