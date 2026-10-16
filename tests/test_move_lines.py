"""Drag selected column / row headers to move whole columns / rows."""
import openpyxl

import formula
import structure
import workbooks
from conftest import open_in
from grid import ALL


def setup(app):
    g = app.table
    app.set_cells([(0, c, h) for c, h in enumerate(["A", "B", "C", "D"])] +
                  [(1, c, str(v)) for c, v in enumerate([1, 2, 3, 4])] +
                  [(2, 4, "=B2*10"), (3, 4, "=SUM(A2:D2)")])
    g.select(0, 1, 0, 1)
    app.format_cells(lambda st: {**st, "bold": True, "bg": "#ffff00"})
    app.set_width(1, 20)
    return g


def test_mover():
    p = structure.mover(1, 2, 5)          # lines 1, 2 go before 5: 0 3 4 1 2 5
    assert [p(i) for i in range(6)] == [0, 3, 4, 1, 2, 5]
    p = structure.mover(3, 1, 0)          # line 3 to the top: 3 0 1 2
    assert [p(i) for i in range(4)] == [1, 2, 3, 0]


def test_permute_formula():
    p = structure.mover(1, 1, 4)          # column B to after D
    assert formula.permute("=B2*10+A2", "S", "S", "col", p) == "=D2*10+A2"
    assert formula.permute("=SUM(A2:D2)", "S", "S", "col", p) == "=SUM(A2:D2)"   # the range keeps its ends
    assert formula.permute("=Other!B2", "S", "S", "col", p) == "=Other!B2"       # another sheet: untouched
    assert formula.permute("=$B$1", "S", "S", "col", p) == "=$D$1"


def test_move_column_with_everything(app):
    g = setup(app)
    first = app.move_lines("col", 1, 1, 4)           # B goes after D
    assert first == 3
    assert [g.text(0, c) for c in range(4)] == ["A", "C", "D", "B"]
    assert [g.text(1, c) for c in range(4)] == ["1", "3", "4", "2"]
    assert g.sheet.styles[(0, 3)]["bold"] and not (g.sheet.styles.get((0, 1)) or {}).get("bold")
    assert g.sheet.widths.get(3) == 20 and 1 not in g.sheet.widths
    assert g.text(2, 4) == "=D2*10" and g.value(2, 4) == "20"   # the formula follows its cell
    assert g.value(3, 4) == "10"
    assert g.selected_lines() == ("col", 3, 3)            # the moved column stays selected
    app.undo()
    assert [g.text(0, c) for c in range(4)] == ["A", "B", "C", "D"] and g.text(2, 4) == "=B2*10"
    assert g.sheet.widths.get(1) == 20
    app.redo()
    assert g.text(0, 3) == "B"


def test_move_rows_up(app):
    g = app.table
    app.set_cells([(r, 0, f"r{r}") for r in range(6)] + [(0, 2, "=A5")])
    g.sheet.heights[4] = 30.0
    first = app.move_lines("row", 4, 2, 1)           # rows 5-6 go between row 1 and row 2
    assert first == 1
    assert [g.text(r, 0) for r in range(6)] == ["r0", "r4", "r5", "r1", "r2", "r3"]
    assert g.sheet.heights.get(1) == 30.0 and g.text(0, 2) == "=A2"
    assert g.selected_lines() == ("row", 1, 2)


def test_no_move_onto_itself(app):
    setup(app)
    assert app.move_lines("col", 1, 1, 1) is None and app.move_lines("col", 1, 1, 2) is None


def test_merged_block_not_split(app):
    g = setup(app)
    g.select(0, 1, 0, 2)
    app.merge()
    n = len(app.dialogs)
    assert app.move_lines("col", 1, 1, 4) is None                 # B alone: part of the B:C merge
    assert app.dialogs[n][0] == "showwarning" and g.text(0, 1) == "B"
    assert app.move_lines("col", 1, 2, 4) == 2                    # B:C together: fine
    assert g.sheet.merges == [(0, 2, 0, 3)]


def test_defined_names_follow(app):
    setup(app)
    app.names["Price"] = "Sheet1!$B$2" if app.sheet == "Sheet1" else f"{app.sheet}!$B$2"
    app.move_lines("col", 1, 1, 4)
    assert app.names["Price"].endswith("$D$2")


def test_saved_after_move(app, tmp_path):
    path = workbooks.plain(tmp_path / "m.xlsx")
    g = open_in(app, path)
    before = [g.text(0, c) for c in range(3)]
    app.set_cells([(5, 0, "=A1")])
    g.select(0, 0, 0, 0)
    app.format_cells(lambda st: {**st, "italic": True, "bg": "#00ff00"})
    app.move_lines("col", 0, 1, 3)                                 # A goes after C
    assert [g.text(0, c) for c in range(3)] == before[1:] + before[:1]
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert [ws.cell(1, c).value for c in (1, 2, 3)] == [v or None for v in before[1:] + before[:1]]
    assert ws["C6"].value == "=C1"
    assert ws["C1"].font.i and ws["C1"].fill.fgColor.rgb.endswith("00FF00") and not ws["A1"].font.i


def header_xy(g, axis, i):
    if axis == "col":
        return g.col_x(i) + g.colw(i) // 2, g.hdr_h() // 2
    return g.hdr_w() // 2, g.row_y(i) + g.rh(i) // 2


def mouse(g, x, y, kind):
    g.cv.event_generate(kind, x=x, y=y, when="now")


def test_drag_column_header(app):
    g = setup(app)
    app.update()
    x, y = header_xy(g, "col", 1)
    mouse(g, x, y, "<Button-1>")                                   # select column B
    mouse(g, x, y, "<ButtonRelease-1>")
    assert g.selected_lines() == ("col", 1, 1)
    mouse(g, x, y, "<Button-2>")                                   # (not a double-click: a moment later)
    mouse(g, x, y, "<Button-1>")                                   # press it again and drag to the right of D
    tx = g.col_x(3) + g.colw(3) - 3
    mouse(g, tx, y, "<B1-Motion>")
    assert g.cv.find_withtag("moveline")                           # a guide line shows where it lands
    mouse(g, tx, y, "<ButtonRelease-1>")
    assert [g.text(0, c) for c in range(4)] == ["A", "C", "D", "B"]


def test_drag_row_header_and_plain_click(app):
    g = app.table
    app.set_cells([(r, 0, f"r{r}") for r in range(5)])
    app.update()
    g.select_line("row", 3, 4)                                     # rows 4-5 selected
    x, y = header_xy(g, "row", 3)
    mouse(g, x, y, "<Button-1>")
    ty = g.row_y(0) + 2
    mouse(g, x, ty, "<B1-Motion>")
    mouse(g, x, ty, "<ButtonRelease-1>")
    assert [g.text(r, 0) for r in range(5)] == ["r3", "r4", "r0", "r1", "r2"]
    assert g.selected_lines() == ("row", 0, 1)
    x, y = header_xy(g, "row", 1)                                  # a click (no drag) inside the selection
    mouse(g, x, y, "<Button-2>")
    mouse(g, x, y, "<Button-1>")
    mouse(g, x, y, "<ButtonRelease-1>")
    assert g.selected_lines() == ("row", 1, 1)                     # just selects that row, as before
    assert g.sel() == (1, 0, 1, ALL)


def test_move_cursor_on_selected_header(app):
    g = setup(app)
    app.update()
    g.select_line("col", 1, 2)
    x, y = header_xy(g, "col", 2)
    mouse(g, x, y, "<Motion>")
    assert g.cv["cursor"] == "fleur"
    x, y = header_xy(g, "col", 3)                                  # not selected: no move pointer
    mouse(g, x, y, "<Motion>")
    assert g.cv["cursor"] != "fleur"
