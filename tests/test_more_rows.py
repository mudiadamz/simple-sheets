"""Long sheets: more rows are read while scrolling, without reopening (unsaved edits stay)."""
import openpyxl
from openpyxl.styles import Font

import sheet_io as io


def long_file(path, n=130):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, n + 1):
        ws.append([f"row {r}", r, f"=B{r}*2"])
    ws["A100"].font = Font(bold=True)
    wb.save(path)
    return path


def test_rows_read_while_scrolling(app, tmp_path, monkeypatch):
    path = long_file(tmp_path / "long.xlsx")
    monkeypatch.setattr(io, "MAX_ROWS", 40)
    app.open_file(str(path))
    g = app.table
    assert len(g.sheet.rows) == 40 and g.sheet.loaded_limit == 40
    app.set_cells([(0, 0, "edited before")])
    assert app.load_more_rows(chunk=40) == 40
    assert len(g.sheet.rows) == 80 and g.text(0, 0) == "edited before"
    g.goto(79, 0)                                      # scrolling to the end reads the next part
    for _ in range(5):
        app.update()
    assert len(g.sheet.rows) >= 100
    while app.load_more_rows():
        pass
    assert len(g.sheet.rows) == 130 and g.sheet.loaded_limit is None and not app.truncated()
    assert g.style(99, 0) == {"bold": True} and g.value(129, 2) == "260"
    assert app.save()
    wb = openpyxl.load_workbook(path)
    assert wb.active["A1"].value == "edited before" and wb.active.max_row == 130
