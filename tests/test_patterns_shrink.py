"""Pattern fills (stripes, dots, grids) shown and kept; shrink to fit drawn smaller, turned on / off, kept."""
import os

import openpyxl

import grid
import workbooks
from conftest import open_in


def test_pattern_fill_read_drawn_and_kept(app, tmp_path):
    src = workbooks.patterns(tmp_path / "p.xlsx")
    g = open_in(app, src)
    st = g.sheet.styles[(1, 0)]
    assert st["pattern"] == "darkDown" and st["pattern_fg"] == "#ff0000" and st["bg"] == "#ffff00"
    assert g.sheet.styles[(2, 0)]["pattern"] == "lightGray" and "bg" not in g.sheet.styles[(2, 0)]
    g.redraw()
    app.update()
    stipples = {g.cv.itemcget(i, "stipple") for i in g.cv.find_withtag("pattern")}
    assert "gray25" in stipples and any(s.startswith("@") and s.endswith("darkDown.xbm") for s in stipples)
    assert os.path.exists(grid.pattern_stipple("darkDown")[1:])
    assert grid.pattern_stipple("solid") is None and grid.pattern_stipple("mediumGray") == "gray50"
    g.goto(0, 0)
    app.apply_mixed([("v", 0, 0, "edited")])                          # save after an edit: the patterns stay
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(src).active
    assert ws["A2"].fill.patternType == "darkDown" and ws["A2"].fill.fgColor.rgb == "FFFF0000"
    assert ws["A2"].fill.bgColor.rgb == "FFFFFF00" and ws["A3"].fill.patternType == "lightGray"


def test_picking_a_fill_replaces_the_pattern(app, tmp_path, monkeypatch):
    g = open_in(app, workbooks.patterns(tmp_path / "p.xlsx"))
    g.goto(1, 0)
    from tkinter import colorchooser
    monkeypatch.setattr(colorchooser, "askcolor", lambda *a, **k: ((0, 0, 255), "#0000ff"))
    app.fmtbar.pick("bg")
    st = g.sheet.styles[(1, 0)]
    assert st["bg"] == "#0000ff" and "pattern" not in st and "pattern_fg" not in st


def test_shrink_to_fit(app, tmp_path):
    src = workbooks.patterns(tmp_path / "p.xlsx")
    g = open_in(app, src)
    assert g.sheet.styles[(1, 1)]["shrink"] is True
    g.redraw()
    app.update()

    def font_of(text):
        item = next(i for i in g.cv.find_withtag("txt") if g.cv.itemcget(i, "text").startswith(text))
        return g.cv.itemcget(item, "font"), g.cv.itemcget(item, "text")
    font, shown = font_of("A long")
    normal = g.font(None)
    import tkinter.font as tkfont
    assert tkfont.Font(font=font).actual("size") < normal.actual("size")       # drawn smaller...
    assert shown == "A long text that fits drawn smaller"                       # (all of it)
    g.goto(1, 1)
    assert app.fmtbar.toggle_shrink() is False and "shrink" not in (g.sheet.styles.get((1, 1)) or {})
    g.redraw()
    app.update()
    assert tkfont.Font(font=font_of("A long")[0]).actual("size") == normal.actual("size")   # ...not any more
    assert app.fmtbar.toggle_shrink() is True
    g.goto(2, 2)
    assert app.fmtbar.toggle_shrink() is True                                   # on an empty cell too
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(src).active
    assert ws["B2"].alignment.shrink_to_fit and ws["C3"].alignment.shrink_to_fit
    m = app.fmtbar.rotate_menu
    assert "Shrink to fit (on / off)" in [m.entrycget(i, "label") for i in range(m.index("end") + 1)
                                          if m.type(i) == "command"]
