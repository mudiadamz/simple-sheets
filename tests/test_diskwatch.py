"""Warnings when the open file changes on disk or is locked by another program (Excel)."""
import ctypes
import os
import sys

import openpyxl
import pytest

import diskwatch
import workbooks
from conftest import open_in


def touch_other_program(path):
    """Another program rewrites the file."""
    wb = openpyxl.load_workbook(path)
    wb.active["A1"] = "changed elsewhere"
    wb.save(path)
    st = os.stat(path)
    os.utime(path, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))


def test_reload_offered_when_changed(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    app.check_disk()
    assert not app.dialogs                                      # unchanged: nothing asked
    touch_other_program(path)
    app.check_disk()
    assert app.dialogs[-1][0] == "askokcancel"                  # stub answers OK -> reloaded
    assert g.text(0, 0) == "changed elsewhere"
    n = len(app.dialogs)
    app.check_disk()
    assert len(app.dialogs) == n                                # asked once per change


def test_save_over_changed_file_asks(app, tmp_path, monkeypatch):
    path = workbooks.plain(tmp_path / "p.xlsx")
    open_in(app, path)
    app.set_cells([(5, 0, "mine")])
    touch_other_program(path)
    import simple_sheets
    monkeypatch.setattr(simple_sheets.messagebox, "askokcancel", lambda *a, **k: False)
    assert not app.save()                                       # said Cancel: file untouched
    assert openpyxl.load_workbook(path).active["A1"].value == "changed elsewhere"
    monkeypatch.setattr(simple_sheets.messagebox, "askokcancel", lambda *a, **k: True)
    assert app.save()
    assert openpyxl.load_workbook(path).active["A6"].value == "mine"
    assert app.disk_stamp == diskwatch.stamp(str(path))


@pytest.mark.skipif(sys.platform != "win32", reason="Windows file sharing")
def test_locked_file(app, tmp_path):
    path = str(workbooks.plain(tmp_path / "p.xlsx"))
    k32 = ctypes.windll.kernel32
    k32.CreateFileW.restype = ctypes.c_void_p
    # like Excel: open for read/write, let others only read
    h = k32.CreateFileW(path, 0xC0000000, 1, None, 3, 0x80, None)
    try:
        assert diskwatch.is_locked(path)
        open_in(app, path)
        assert app.locked and "in use by another program" in app.title()
        app.set_cells([(5, 0, "x")])
        assert not app.save()
        assert app.dialogs[-1][:2] == ("showerror", "File in use")
    finally:
        k32.CloseHandle(ctypes.c_void_p(h))
    assert not diskwatch.is_locked(path)
    app.check_disk()                                            # Excel closed it: title back to normal
    assert not app.locked and "in use" not in app.title()
    assert app.save()
