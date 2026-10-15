"""Data validation (drop-down lists, rejecting input) and conditional formatting shown from the file."""
import openpyxl
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation

from conftest import key, open_in


def rules_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Qty", "Status", "Score", "Tag"])
    for name, qty, status, score, tag in [("a", 5, "Open", 10, "x"), ("b", 50, "Done", 20, "y"),
                                          ("c", 15, "Open", 30, "x"), ("d", 2, "Hold", 40, "zz")]:
        ws.append([name, qty, status, score, tag])
    ws["G1"], ws["G2"], ws["G3"] = "Open", "Done", "Hold"
    lst = DataValidation(type="list", formula1="$G$1:$G$3", showErrorMessage=True, error="Pick from the list")
    ws.add_data_validation(lst)
    lst.add("C2:C10")
    yn = DataValidation(type="list", formula1='"Yes,No"', showErrorMessage=True, errorStyle="warning")
    ws.add_data_validation(yn)
    yn.add("F2:F10")
    num = DataValidation(type="whole", operator="between", formula1="1", formula2="100", showErrorMessage=True)
    ws.add_data_validation(num)
    num.add("B2:B10")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ws.conditional_formatting.add("B2:B5", CellIsRule(operator="greaterThan", formula=["10"], fill=red,
                                                      font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add("A2:A5", FormulaRule(formula=['$C2="Hold"'], fill=red))
    ws.conditional_formatting.add("D2:D5", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                          end_type="max", end_color="00FF00"))
    ws.conditional_formatting.add("H2:H5", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    for r, v in enumerate([1, 2, 3, 4], start=2):
        ws.cell(r, 8, v)
    dup = Rule(type="duplicateValues", dxf=DifferentialStyle(fill=red))
    ws.conditional_formatting.add("E2:E5", dup)
    txt = Rule(type="containsText", operator="containsText", text="z", dxf=DifferentialStyle(font=Font(italic=True)),
               formula=['NOT(ISERROR(SEARCH("z",E2)))'])
    ws.conditional_formatting.add("E2:E5", txt)
    top = Rule(type="top10", rank=1, dxf=DifferentialStyle(font=Font(underline="single")))
    ws.conditional_formatting.add("D2:D5", top)
    avg = Rule(type="aboveAverage", aboveAverage=False, dxf=DifferentialStyle(font=Font(strike=True)))
    ws.conditional_formatting.add("H2:H5", avg)
    wb.save(path)
    return path


def test_conditional_styles(app, tmp_path):
    g = open_in(app, rules_file(tmp_path / "r.xlsx"))
    cf = lambda r, c: app.rules.conditional(g.sheet, r, c) or {}   # noqa: E731
    assert cf(2, 1) == {"bg": "#ffc7ce", "bold": True, "fg": "#9c0006"}     # 50 > 10
    assert cf(1, 1) == {}                                                     # 5
    assert cf(4, 0).get("bg") == "#ffc7ce" and cf(1, 0) == {}                 # row with Status Hold
    assert cf(1, 3)["bg"] == "#ffffff" and cf(4, 3)["bg"] == "#00ff00"        # colour scale ends
    assert cf(2, 3)["bg"] not in ("#ffffff", "#00ff00")
    assert cf(4, 3).get("underline") and not cf(3, 3).get("underline")       # top 1
    assert cf(1, 7)["bar"][0] < cf(4, 7)["bar"][0] == 1.0                    # data bars grow
    assert cf(1, 7).get("strike") and not cf(4, 7).get("strike")             # below average
    assert cf(1, 4).get("bg") == "#ffc7ce" and not cf(4, 4).get("bg")        # duplicates x x
    assert cf(4, 4).get("italic") and not cf(1, 4).get("italic")             # contains "z"
    app.set_cells([(1, 1, "99")])                                             # recalculated after edits
    assert cf(1, 1).get("bold")
    app.update()
    assert app.table.cv.find_withtag("databar")


def test_follows_inserted_rows(app, tmp_path):
    g = open_in(app, rules_file(tmp_path / "r.xlsx"))
    app.restructure("row", 1, 1)                  # insert a row above row 2
    assert not app.rules.conditional(g.sheet, 1, 1)          # the new row has no rules
    assert app.rules.conditional(g.sheet, 3, 1).get("bold")  # 50 moved down one row


def test_dropdown_list(app, tmp_path):
    g = open_in(app, rules_file(tmp_path / "r.xlsx"))
    assert app.rules.dropdown(g.sheet, 1, 2) == ["Open", "Done", "Hold"]
    assert app.rules.dropdown(g.sheet, 1, 5) == ["Yes", "No"]
    assert app.rules.dropdown(g.sheet, 1, 1) is None
    g.goto(1, 2)
    app.update()
    assert g.dv_button and g.cv.find_withtag("dvbtn")
    g.open_dropdown()
    top, lb, opts, pick = g.dropdown
    assert lb.curselection() == (0,)              # current value "Open" highlighted
    lb.selection_clear(0, "end")
    lb.selection_set(2)
    pick()
    assert g.text(1, 2) == "Hold" and g.dropdown is None
    app.undo()
    assert g.text(1, 2) == "Open"
    g.goto(1, 0)
    app.update()
    assert g.dv_button is None


def test_click_arrow_opens_list(app, tmp_path):
    g = open_in(app, rules_file(tmp_path / "r.xlsx"))
    g.goto(2, 2)
    app.update()
    x1, y1, x2, y2 = g.dv_button
    g.cv.event_generate("<Button-1>", x=(x1 + x2) // 2, y=(y1 + y2) // 2, when="now")
    app.update()
    assert g.dropdown is not None
    key(app, "<Escape>", g.dropdown[1])
    assert g.dropdown is None


def test_reject_invalid_input(app, tmp_path):
    g = open_in(app, rules_file(tmp_path / "r.xlsx"))
    g.goto(1, 2)
    g.edit("Maybe")
    g.commit(1, 0)
    assert g.text(1, 2) == "Open"                                 # stop: rejected ...
    assert app.dialogs[-1][:3] == ("showerror", "Invalid value", "Pick from the list")
    assert g.editor is not None and g.editor.get() == "Maybe"     # ... and back to editing
    g.editor.delete(0, "end")
    g.editor.insert(0, "done")                                    # list match ignores case
    g.commit()
    assert g.text(1, 2) == "done"
    g.goto(1, 1)
    g.edit("150")
    g.commit()
    assert g.text(1, 1) == "5" and "between 1 and 100" in app.dialogs[-1][2]
    g.commit(cancel=True)
    g.edit("2.5")
    g.commit()
    assert g.text(1, 1) == "5" and "whole number" in app.dialogs[-1][2]
    g.commit(cancel=True)
    g.edit("7")
    g.commit()
    assert g.text(1, 1) == "7"
    g.goto(1, 5)                                                  # warning style: asks, stub says OK
    g.edit("Perhaps")
    g.commit()
    assert g.text(1, 5) == "Perhaps" and app.dialogs[-1][0] == "askokcancel"


def test_formula_bar_validation(app, tmp_path):
    g = open_in(app, rules_file(tmp_path / "r.xlsx"))
    g.goto(1, 1)
    app.update()
    app.value.delete(0, "end")
    app.value.insert(0, "0")
    app.value_commit()
    assert g.text(1, 1) == "5" and app.dialogs[-1][0] == "showerror"


def test_rules_kept_on_save(app, tmp_path):
    path = rules_file(tmp_path / "r.xlsx")
    open_in(app, path)
    app.set_cells([(1, 0, "zz")])
    assert app.save()
    ws = openpyxl.load_workbook(path)["Data"]
    assert len(ws.data_validations.dataValidation) == 3
    assert sum(len(cf.rules) for cf in ws.conditional_formatting) == 8
