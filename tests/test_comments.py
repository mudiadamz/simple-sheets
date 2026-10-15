"""Cell comments / notes."""
import openpyxl
from openpyxl.comments import Comment

import comments
import sheet_io as io
from conftest import open_in
from test_menu import menu_states


def note_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["B2"] = "x"
    ws["B2"].comment = Comment("check this", "Ann")
    wb.save(path)
    return path


def test_load_show_and_hover(app, tmp_path):
    g = open_in(app, note_file(tmp_path / "n.xlsx"))
    assert g.sheet.comments == {(1, 1): ["check this", "Ann"]}
    assert len(g.cv.find_withtag("note")) == 1                    # red corner
    g.notes.motion(g.col_x(1) + 5, g.row_y(1) + 5)
    g.notes.show()
    assert g.notes.tip is not None
    label = g.notes.tip.winfo_children()[0]
    assert label.cget("text") == "Ann:\ncheck this"
    g.notes.motion(g.col_x(4) + 5, g.row_y(4) + 5)              # away: popup goes
    assert g.notes.tip is None


def test_add_edit_delete_with_undo(app):
    g = app.table
    g.goto(0, 0)
    assert "Insert comment" in menu_states(app)
    d = app.edit_comment()
    d.text.insert("1.0", "first note")
    d.ok()
    assert g.sheet.comments[(0, 0)][0] == "first note" and app.dirty()
    assert {"Edit comment", "Delete comment"} <= set(menu_states(app))
    app.set_comment(0, 0, "changed")
    assert g.sheet.comments[(0, 0)][0] == "changed"
    app.set_comment(0, 0, "")
    assert (0, 0) not in g.sheet.comments
    app.undo()
    assert g.sheet.comments[(0, 0)][0] == "changed"
    app.undo()
    app.undo()
    assert not g.sheet.comments and not app.dirty()


def test_saved_and_moved(app, tmp_path):
    path = note_file(tmp_path / "n.xlsx")
    g = open_in(app, path)
    app.set_comment(4, 0, "new one")
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["S"]
    assert ws["A5"].comment.text == "new one" and ws["B2"].comment.text == "check this"
    app.restructure("row", 0, 1)                                # comments move down
    assert set(g.sheet.comments) == {(2, 1), (5, 0)}
    app.set_comment(2, 1, "")                                    # delete one on a rebuilt sheet
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["S"]
    assert ws["B3"].comment is None and ws["A6"].comment.text == "new one"
    assert io.load(str(path))["S"].comments == {(5, 0): ["new one", comments.author()]}


def test_new_workbook_comment_saved(app, tmp_path, monkeypatch):
    import simple_sheets
    dst = tmp_path / "w.xlsx"
    monkeypatch.setattr(simple_sheets.filedialog, "asksaveasfilename", lambda **k: str(dst))
    app.set_cells([(0, 0, "a")])
    app.set_comment(0, 0, "hello")
    assert app.save(), app.dialogs
    assert openpyxl.load_workbook(dst).active["A1"].comment.text == "hello"
