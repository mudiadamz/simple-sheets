"""P1 data-safety fixes: rich text kept, ranges move with inserted/deleted rows and columns,
.xls dates, formula results stored in saved files."""
import datetime as dt
import os
import zipfile

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

import sheet_io as io
from conftest import open_in


def features(path):
    """A workbook using the range-based features that must follow inserted/deleted rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["name", "qty", "price"])
    for r in range(2, 7):
        ws.append([f"item{r}", r, r * 10])
    ws["E1"] = CellRichText(["plain ", TextBlock(InlineFont(b=True, color="FFFF0000"), "bold red")])
    dv = DataValidation(type="list", formula1='"yes,no"')
    ws.add_data_validation(dv)
    dv.add("D2:D6")
    ws.conditional_formatting.add("B2:B6", CellIsRule(operator="greaterThan", formula=["$C$2"],
                                                      fill=PatternFill("solid", fgColor="FFFFC7CE")))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:C6"
    ws.print_title_rows = "1:1"
    wb.defined_names["prices"] = openpyxl.workbook.defined_name.DefinedName("prices", attr_text="S!$C$2:$C$6")
    t = wb.create_sheet("T")
    t.append(["h1", "h2"])
    t.append([1, 2])
    t.add_table(Table(displayName="Tbl", ref="A1:B2"))
    wb.save(path)
    return path


# ---------- rich text ----------

def test_rich_text_survives_normal_save(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    g = open_in(app, path)
    assert g.text(0, 4) == "plain bold red"                     # shown as its plain text
    app.set_cells([(10, 0, "edit")])
    assert app.save(), app.dialogs
    v = openpyxl.load_workbook(path, rich_text=True)["S"]["E1"].value
    assert isinstance(v, CellRichText) and str(v) == "plain bold red"


def test_rich_text_survives_rebuild_but_edit_makes_plain(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    open_in(app, path)
    app.restructure("row", 0, 1)                                # E1 -> E2, sheet rebuilt on save
    app.set_cells([(3, 0, "changed")])
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path, rich_text=True)["S"]
    assert isinstance(ws["E2"].value, CellRichText)
    open_in(app, path)
    app.set_cells([(1, 4, "now plain")])
    assert app.save(), app.dialogs
    assert openpyxl.load_workbook(path, rich_text=True)["S"]["E2"].value == "now plain"


# ---------- ranges follow inserted / deleted rows and columns ----------

def test_ranges_move_with_inserted_rows(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    open_in(app, path)
    app.restructure("row", 0, 2)                                # 2 rows on top: everything moves down 2
    assert app.save(), app.dialogs
    wb = openpyxl.load_workbook(path)
    ws = wb["S"]
    assert str(ws.data_validations.dataValidation[0].sqref) == "D4:D8"
    cf = list(ws.conditional_formatting)
    assert str(cf[0].sqref) == "B4:B8" and cf[0].rules[0].formula == ["$C$4"]
    assert ws.freeze_panes == "B4"
    assert ws.auto_filter.ref == "A3:C8"
    assert ws.print_title_rows == "$3:$3"
    assert wb.defined_names["prices"].attr_text == "S!$C$4:$C$8"


def test_ranges_shrink_with_deleted_rows_and_columns(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    open_in(app, path)
    app.restructure("row", 2, -2)                               # delete rows 3-4
    app.restructure("col", 0, -1)                               # delete column A
    assert app.save(), app.dialogs
    wb = openpyxl.load_workbook(path)
    ws = wb["S"]
    assert str(ws.data_validations.dataValidation[0].sqref) == "C2:C4"
    assert str(list(ws.conditional_formatting)[0].sqref) == "A2:A4"
    assert ws.auto_filter.ref == "A1:B4"
    assert wb.defined_names["prices"].attr_text == "S!$B$2:$B$4"


def test_fully_deleted_range_is_dropped(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    open_in(app, path)
    app.restructure("col", 3, -1)                               # column D held the drop-down list
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["S"]
    assert not ws.data_validations.dataValidation


def test_excel_table_grows_with_inserted_column(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    open_in(app, path)
    app.show("T")
    app.restructure("col", 1, 1)                                # new column inside the table
    app.set_cells([(0, 1, "new")])
    assert app.save(), app.dialogs
    t = openpyxl.load_workbook(path)["T"].tables["Tbl"]
    assert t.ref == "A1:C2" and [c.name for c in t.tableColumns] == ["h1", "new", "h2"]


# ---------- .xls ----------

def test_xls_dates_booleans_numbers(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    path = str(tmp_path / "old.xls")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Old")
    date_style = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    ws.write(0, 0, dt.datetime(2024, 3, 5), date_style)
    ws.write(0, 1, dt.datetime(2024, 3, 5, 14, 30), xlwt.easyxf(num_format_str="yyyy-mm-dd hh:mm"))
    ws.write(0, 2, True)
    ws.write(0, 3, 12.5)
    ws.write(0, 4, "text")
    wb.save(path)
    s = io.load(path)["Old"]
    assert s.rows[0] == ["2024-03-05", "2024-03-05 14:30:00", "TRUE", "12.5", "text"]


# ---------- formula results stored in the file ----------

def cached(path, sheet_xml="xl/worksheets/sheet1.xml"):
    """Results stored in the file, as a program that doesn't calculate would read them."""
    return openpyxl.load_workbook(path, data_only=True)


def test_formula_results_written(app, tmp_path):
    from workbooks import budget
    path = budget(tmp_path / "b.xlsx")                          # openpyxl-made: no results stored
    open_in(app, path)
    assert cached(path).active["D6"].value is None
    app.set_cells([(6, 0, '=IF(D6>100,"big","small")'), (7, 0, "=D6>100"), (8, 0, "=1/0")])
    assert app.save(), app.dialogs
    ws = cached(path).active
    assert ws["D2"].value == 25 and ws["D6"].value == pytest.approx(104.97)
    assert ws["A7"].value == "big" and ws["A8"].value is True and ws["A9"].value == "#DIV/0!"
    formulas = openpyxl.load_workbook(path).active
    assert formulas["D6"].value == "=SUM(D2:D4)"                 # the formulas themselves are kept


def test_results_written_for_new_workbook_and_other_sheets(app, tmp_path, monkeypatch):
    import simple_sheets
    dst = tmp_path / "new.xlsx"
    monkeypatch.setattr(simple_sheets.filedialog, "asksaveasfilename", lambda **k: str(dst))
    app.set_cells([(0, 0, "5"), (0, 1, "=A1*2")])
    app.add_sheet()
    app.set_cells([(0, 0, "=Sheet1!B1+1")])
    assert app.save(), app.dialogs
    wb = cached(dst)
    assert wb["Sheet1"]["B1"].value == 10 and wb["Sheet2"]["A1"].value == 11


def test_long_chain_results(app, tmp_path, monkeypatch):
    import simple_sheets
    dst = tmp_path / "chain.xlsx"
    monkeypatch.setattr(simple_sheets.filedialog, "asksaveasfilename", lambda **k: str(dst))
    app.set_cells([(0, 0, "1")] + [(r, 0, f"=A{r}+1") for r in range(1, 3000)])
    assert app.save(), app.dialogs
    assert cached(dst).active["A3000"].value == 3000


def test_saved_file_is_a_valid_zip_with_all_parts(app, tmp_path):
    path = features(tmp_path / "f.xlsx")
    before = set(zipfile.ZipFile(path).namelist())
    open_in(app, path)
    app.set_cells([(10, 0, "=B2*2")])
    assert app.save(), app.dialogs
    z = zipfile.ZipFile(path)
    assert z.testzip() is None and before <= set(z.namelist())
    assert not os.path.exists(str(path) + ".tmp")
