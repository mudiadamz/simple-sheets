"""Cells with hyperlinks: emptying them must stick after save + reopen."""
import openpyxl

import sheet_io as io
import workbooks
from conftest import key, open_in


def test_delete_linked_cells_stays_deleted(app, tmp_path):
    path = workbooks.links(tmp_path / "l.xlsx")
    g = open_in(app, path)
    g.select(1, 1, 3, 1)                                   # B2:B4
    key(app, "<Delete>")
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Links"]
    assert [ws.cell(r, 2).value for r in range(1, 6)] == \
        ["https://example.com/page1", None, None, None, "https://example.com/page5"]
    assert ws["B2"].hyperlink is None and ws["B5"].hyperlink.target == "https://example.com/page5"
    g = open_in(app, path)                                 # reopen in the app
    assert [g.text(r, 1) for r in range(5)] == ["https://example.com/page1", "", "", "",
                                                 "https://example.com/page5"]


def test_right_click_open_link(app, tmp_path, monkeypatch):
    import menus
    from test_menu import menu_labels
    opened = []
    monkeypatch.setattr(menus.webbrowser, "open", lambda url: opened.append(url))
    path = workbooks.links(tmp_path / "l.xlsx")
    g = open_in(app, path)
    assert g.sheet.links[(2, 1)] == "https://example.com/page3"
    g.goto(2, 1)
    labels = menu_labels(app)
    assert labels[0] == "Open link"
    from test_menu import build_menu
    build_menu(app).invoke(0)
    assert opened == ["https://example.com/page3"]
    g.goto(2, 0)                                           # plain text: no link item
    assert "Open link" not in menu_labels(app)


def test_open_link_uses_link_not_text(app, tmp_path):
    import menus
    g = app.table
    app.set_cells([(0, 0, "Docs")])
    g.sheet.links[(0, 0)] = "https://docs.example.com"
    assert menus.link_at(g, 0, 0) == "https://docs.example.com"
    app.set_cells([(0, 0, "")])
    assert menus.link_at(g, 0, 0) is None                  # emptied cell: link is gone


def test_web_address_text_is_a_link(app):
    import menus
    g = app.table
    app.set_cells([(0, 0, "https://aws.amazon.com/x"), (1, 0, "www.example.com"),
                   (2, 0, "mailto:a@b.com"), (3, 0, "see https://x.com"), (4, 0, "#Sheet2!A1")])
    assert menus.link_at(g, 0, 0) == "https://aws.amazon.com/x"
    assert menus.link_at(g, 1, 0) == "http://www.example.com"
    assert menus.link_at(g, 2, 0) == "mailto:a@b.com"
    assert menus.link_at(g, 3, 0) is None                  # text containing a link isn't one
    assert menus.link_at(g, 4, 0) is None


def test_links_move_with_inserted_rows(app, tmp_path):
    path = workbooks.links(tmp_path / "l.xlsx")
    g = open_in(app, path)
    app.restructure("row", 0, 2)
    assert g.sheet.links[(4, 1)] == "https://example.com/page3"
    app.undo()
    assert g.sheet.links[(2, 1)] == "https://example.com/page3"


def test_untouched_links_are_kept(app, tmp_path):
    path = workbooks.links(tmp_path / "l.xlsx")
    open_in(app, path)
    app.set_cells([(0, 0, "changed")])
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Links"]
    assert ws["B3"].hyperlink.target == "https://example.com/page3"


def test_emptied_link_cell_after_row_insert(app, tmp_path):
    path = workbooks.links(tmp_path / "l.xlsx")
    g = open_in(app, path)
    app.restructure("row", 0, 1)                           # sheet is rebuilt on save
    g.goto(2, 1)                                           # old B2, now B3
    key(app, "<Delete>")
    assert app.save(), app.dialogs
    s = io.load(str(path))["Links"]
    assert s.rows[2][1] == "" and s.rows[3][1] == "https://example.com/page3"
    ws = openpyxl.load_workbook(path)["Links"]
    assert ws["B3"].hyperlink is None and ws["B4"].hyperlink.target == "https://example.com/page3"
