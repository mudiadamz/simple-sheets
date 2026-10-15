"""Several workbooks in one window (tabs), and loading all rows of a long sheet."""
import openpyxl

import book_tabs
import settings
import sheet_io as io
import workbooks


def test_tabs_keep_each_workbook(app, tmp_path):
    a = workbooks.plain(tmp_path / "a.xlsx")
    b = workbooks.sales(tmp_path / "b.xlsx")
    assert app.open_file(str(a))
    app.set_cells([(9, 0, "edit in a")])
    assert book_tabs.new_tab(app, str(b))
    assert app.path == str(b) and app.book_bar.visible and book_tabs.count(app) == 2
    assert app.table.text(9, 0) == "" and not app.dirty()
    app.set_cells([(9, 1, "edit in b")])
    book_tabs.switch(app, 0)
    assert app.path == str(a) and app.table.text(9, 0) == "edit in a" and app.dirty()
    app.undo()                                       # each workbook has its own undo
    assert app.table.text(9, 0) == ""
    book_tabs.switch(app, 1)
    assert app.table.text(9, 1) == "edit in b"
    assert app.save()
    assert openpyxl.load_workbook(b).active["B10"].value == "edit in b"
    assert book_tabs.close(app, 1)                   # stub answers "don't save" for nothing unsaved
    assert book_tabs.count(app) == 1 and app.path == str(a) and not app.book_bar.visible


def test_close_all_asks_each(app, tmp_path, monkeypatch):
    book_tabs.new_tab(app)
    app.set_cells([(0, 0, "x")])
    asked = []
    import simple_sheets
    monkeypatch.setattr(simple_sheets.messagebox, "askyesnocancel", lambda *a, **k: asked.append(a) or None)
    assert not book_tabs.close_all(app)              # Cancel keeps the window open
    assert asked and book_tabs.count(app) == 2


def test_open_in_tabs_setting(app, tmp_path):
    a = workbooks.plain(tmp_path / "a.xlsx")
    b = workbooks.sales(tmp_path / "b.xlsx")
    app.open_file(str(a))
    settings.set("open_in_tabs", True)
    app.open_recent(str(b))
    assert book_tabs.count(app) == 2 and app.path == str(b)


def test_load_all_rows(app, tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    for r in range(1, 41):
        wb.active.append([r])
    p = tmp_path / "long.xlsx"
    wb.save(p)
    monkeypatch.setattr(io, "MAX_ROWS", 25)
    app.open_file(str(p))
    assert len(app.table.sheet.rows) == 25 and app.truncated() and "load all" in app.status["text"]
    assert app.load_all_rows()
    assert len(app.table.sheet.rows) == 40 and not app.truncated()
    app.set_cells([(40, 0, "41")])
    assert app.save()
    assert openpyxl.load_workbook(p).active.max_row == 41
