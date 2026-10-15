"""Data tools (text to columns, remove duplicates, flash fill, pivot, goal seek), grouping, charts,
pictures, and creating conditional formatting / validation rules."""
import io as pyio

import openpyxl
from openpyxl.chart import BarChart, Reference

import charts
import data_more
import outline
import rules_dialog
import workbooks
from conftest import open_in


def test_text_to_columns(app):
    g = app.table
    app.set_cells([(0, 0, "a,b,c"), (1, 0, "d,e")])
    assert data_more.split_column(app, (0, 0, 1, 0), ",") == 5
    assert [g.text(0, c) for c in range(3)] == ["a", "b", "c"] and g.text(1, 1) == "e"
    app.undo()
    assert g.text(0, 0) == "a,b,c" and g.text(0, 1) == ""


def test_remove_duplicates(app):
    g = app.table
    app.set_cells([(r, c, v) for r, row in enumerate([["Name", "N"], ["a", "1"], ["b", "2"], ["A", "1"], ["b", "2"]])
                   for c, v in enumerate(row)])
    assert data_more.dedupe(app, (0, 0, 4, 1), header=True) == 2
    assert [g.text(r, 0) for r in range(5)] == ["Name", "a", "b", "", ""]
    g.goto(1, 0)
    assert data_more.remove_duplicates(app) == 0 and app.dialogs[-1][0] == "showinfo"


def test_flash_fill(app):
    g = app.table
    names = ["Ada Lovelace", "Alan Turing", "Grace Hopper", "Linus Torvalds"]
    app.set_cells([(0, 0, "Full name"), (0, 1, "First")] + [(r + 1, 0, n) for r, n in enumerate(names)] +
                  [(1, 1, "Ada")])
    g.goto(1, 1)
    assert data_more.flash_fill(app) == 3
    assert [g.text(r, 1) for r in range(2, 5)] == ["Alan", "Grace", "Linus"]
    app.set_cells([(1, 2, "LOVELACE")])
    g.goto(1, 2)
    data_more.flash_fill(app)
    assert g.text(4, 2) == "TORVALDS"
    app.set_cells([(1, 3, "A.L."), (2, 3, "A.T.")])
    assert data_more.learn([("Ada Lovelace", "AL")]) is not None
    assert data_more.learn([("x", "something unrelated")]) is None


def test_pivot():
    rows = [["East", "Pens", "10"], ["West", "Pens", "5"], ["East", "Ink", "7"], ["East", "Pens", "3"]]
    t = data_more.make_pivot(rows, ["Region", "Item", "Qty"], "Region", "Qty", "Sum")
    assert t == [["Region", "Sum of Qty"], ["East", "20"], ["West", "5"], ["Total", "25"]]
    t = data_more.make_pivot(rows, ["Region", "Item", "Qty"], "Region", "Qty", "Count", "Item")
    assert t[0] == ["Region", "Pens", "Ink"] and t[1] == ["East", "2", "1"] and t[-1] == ["Total", "3", "1"]


def test_pivot_dialog_adds_sheet(app):
    app.set_cells([(0, 0, "Region"), (0, 1, "Qty"), (1, 0, "E"), (1, 1, "2"), (2, 0, "E"), (2, 1, "3")])
    app.table.goto(0, 0)
    d = data_more.pivot(app)
    d.boxes["values"].set("Qty")
    d.ok()
    assert app.sheet == "Pivot" and app.table.text(1, 1) == "5"
    app.undo()
    assert "Pivot" not in app.sheets


def test_goal_seek(app):
    g = app.table
    app.set_cells([(0, 0, "2"), (0, 1, "=A1*A1+3")])
    x = data_more.goal_seek_solve(app, (0, 1), 28, (0, 0))
    assert abs(x - 5) < 1e-6 and g.text(0, 0) == "2"        # the cell itself is left alone


def test_grouping(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.select(1, 0, 3, 0)
    outline.group_selection(app, True)
    assert g.sheet.outline_rows == {1: 1, 2: 1, 3: 1}
    app.update()
    g.redraw()
    assert g.outline_buttons and g.outline_buttons[0][1:] == ("row", 1, 3)
    outline.toggle(app, "row", 1, 3)
    assert {1, 2, 3} <= g.sheet.hidden_rows
    outline.toggle(app, "row", 1, 3)
    assert not ({1, 2, 3} & g.sheet.hidden_rows)
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws.row_dimensions[2].outlineLevel == 1 and ws.row_dimensions[4].outlineLevel == 1
    g = open_in(app, path)
    assert g.sheet.outline_rows == {1: 1, 2: 1, 3: 1}
    g.select(1, 0, 3, 0)
    outline.group_selection(app, False)
    assert g.sheet.outline_rows == {}


def chart_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Month", "Sales"], ["Jan", 10], ["Feb", 30], ["Mar", 20]]:
        ws.append(row)
    ch = BarChart()
    ch.title = "Sales"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(ch, "D2")
    wb.save(path)
    return path


def test_charts_view_insert_delete_save(app, tmp_path):
    path = chart_file(tmp_path / "c.xlsx")
    g = open_in(app, path)
    assert len(g.sheet.charts) == 1 and g.sheet.charts[0]["title"] == "Sales"
    assert charts.numbers(charts.resolve(app, g.sheet.charts[0]["series"][0]["values"], "Data")) == [10, 30, 20]
    app.update()
    g.redraw()
    assert g.cv.find_withtag("chart") and g.chart_boxes
    g.select(0, 0, 3, 1)
    ch = charts.insert(app, "line", "Trend")
    assert len(g.sheet.charts) == 2 and ch["series"][0]["categories"].endswith("$A$2:$A$4") and app.dirty()
    assert app.save()
    wb = openpyxl.load_workbook(path)
    assert len(wb["Data"]._charts) == 2
    g = open_in(app, path)
    charts.delete(app, g.sheet.charts[0])
    assert app.save()
    assert len(openpyxl.load_workbook(path)["Data"]._charts) == 1
    app.undo()
    assert len(g.sheet.charts) == 2


def test_pictures_drawn(app, tmp_path):
    import tkinter as tk

    from openpyxl.drawing.image import Image
    png = tk.PhotoImage(master=app, width=20, height=10)
    png.put("#ff0000", to=(0, 0, 20, 10))
    p = tmp_path / "red.png"
    png.write(str(p), format="png")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "pic"
    wb.active.add_image(Image(str(p)), "B2")
    wb.save(tmp_path / "i.xlsx")
    g = open_in(app, tmp_path / "i.xlsx")
    app.update()
    g.redraw()
    assert g.cv.find_withtag("image") and g.image_boxes
    app.set_cells([(0, 0, "x")])
    assert app.save()
    assert len(openpyxl.load_workbook(tmp_path / "i.xlsx").active._images) == 1   # kept on save


def test_create_conditional_rule(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    app.set_cells([(10, 0, "5"), (11, 0, "50")])
    g.select(10, 0, 11, 0)
    d = rules_dialog.ConditionalDialog(app)
    assert d.range.get() == "A11:A12"
    assert d.add("cellIs", "greaterThan", "10", "", {"bg": "#ffc7ce"})
    assert app.rules.conditional(g.sheet, 11, 0) == {"bg": "#ffc7ce"} and not app.rules.conditional(g.sheet, 10, 0)
    assert d.add("scale2", ranges=[(10, 0, 11, 0)])
    d.destroy()
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    kinds = sorted(r.type for cf in ws.conditional_formatting for r in cf.rules)
    assert kinds == ["cellIs", "colorScale"]
    g = open_in(app, path)
    assert app.rules.conditional(g.sheet, 11, 0).get("bg")
    d = rules_dialog.ConditionalDialog(app)
    d.tree.selection_set(d.tree.get_children())
    d.delete()
    d.destroy()
    assert app.save()
    assert not list(openpyxl.load_workbook(path).active.conditional_formatting)


def test_create_validation(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    g.select(10, 0, 12, 0)
    d = rules_dialog.ValidationDialog(app)
    d.kind.set("List of items")
    d.v1.set("Yes,No")
    assert d.ok()
    assert app.rules.dropdown(g.sheet, 11, 0) == ["Yes", "No"]
    g.goto(11, 0)
    g.edit("Maybe")
    g.commit()
    assert g.text(11, 0) == "" and g.editor is not None
    g.commit(cancel=True)
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path).active
    dv = ws.data_validations.dataValidation
    assert len(dv) == 1 and dv[0].formula1 == '"Yes,No"' and str(dv[0].sqref) == "A11:A13"
    g = open_in(app, path)
    g.select(11, 0, 11, 0)
    rules_dialog.ValidationDialog(app).clear()                   # cut one cell out of the rule
    assert app.rules.dropdown(g.sheet, 11, 0) is None and app.rules.dropdown(g.sheet, 10, 0) == ["Yes", "No"]
    assert app.save()
    dv = openpyxl.load_workbook(path).active.data_validations.dataValidation
    assert sorted(str(x.sqref) for x in dv) == ["A11 A13"] or sorted(str(r) for x in dv for r in x.sqref.ranges) == ["A11", "A13"]


def test_subtract():
    assert rules_dialog.subtract((0, 0, 4, 4), (1, 1, 2, 2)) == [(0, 0, 0, 4), (3, 0, 4, 4), (1, 0, 2, 0), (1, 3, 2, 4)]
    assert rules_dialog.subtract((0, 0, 1, 1), (5, 5, 6, 6)) == [(0, 0, 1, 1)]
    assert pyio is not None


def test_edit_and_reorder_rules(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    app.set_cells([(10, 0, "5"), (11, 0, "50")])
    d = rules_dialog.ConditionalDialog(app)
    d.add("cellIs", "greaterThan", "10", "", {"bg": "#ffc7ce"}, ranges=[(10, 0, 11, 0)])
    d.add("cellIs", "greaterThan", "1", "", {"bg": "#c6efce"}, ranges=[(10, 0, 11, 0)])
    assert app.rules.conditional(g.sheet, 11, 0)["bg"] == "#c6efce"        # the newest rule is on top
    d.tree.selection_set("1")
    assert d.move(-1)                                                      # older rule moved up: it wins now
    assert app.rules.conditional(g.sheet, 11, 0)["bg"] == "#ffc7ce"
    d.tree.selection_set("0")
    assert d.edit_selected() and d.v1.get() == "10" and d.add_btn.cget("text") == "Save changes"
    d.v1.set("100")
    assert d.add()
    assert len(g.sheet.cond) == 2 and g.sheet.cond[0]["f"] == ["100"]       # edited in place
    assert app.rules.conditional(g.sheet, 11, 0)["bg"] == "#c6efce"
    d.destroy()
    assert app.save()
    rules_found = [r for cf in openpyxl.load_workbook(path).active.conditional_formatting for r in cf.rules]
    assert sorted(r.formula[0] for r in rules_found) == ["1", "100"]
