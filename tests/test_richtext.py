"""Rich text: shown run by run, edited with Ctrl+B/I/U on selected characters, saved."""
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

import richtext
from conftest import key, open_in


def rich_file(path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = CellRichText(["Total: ", TextBlock(InlineFont(b=True, color="FFFF0000"), "42 kg")])
    wb.active["A2"] = "plain"
    wb.save(path)
    return path


def test_load_and_draw(app, tmp_path):
    g = open_in(app, rich_file(tmp_path / "r.xlsx"))
    assert g.text(0, 0) == "Total: 42 kg"
    assert g.sheet.rich[(0, 0)] == [["Total: ", {}], ["42 kg", {"bold": True, "fg": "#ff0000"}]]
    app.update()
    g.redraw()
    texts = [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")]
    assert "Total: " in texts and "42 kg" in texts


def test_edit_keeps_and_changes_runs(app, tmp_path):
    path = rich_file(tmp_path / "r.xlsx")
    g = open_in(app, path)
    g.goto(0, 0)
    g.edit()
    e = g.editor
    assert "bold" in e.tag_names("1.8")                     # the bold run came into the editor
    e.select_range(0, 5)                                    # "Total"
    key(app, "<Control-i>", e)
    g.commit()
    runs = g.sheet.rich[(0, 0)]
    assert runs[0] == ["Total", {"italic": True}] and runs[-1] == ["42 kg", {"bold": True, "fg": "#ff0000"}]
    assert app.save()
    cell = openpyxl.load_workbook(path, rich_text=True).active["A1"].value
    assert isinstance(cell, CellRichText) and str(cell) == "Total: 42 kg"
    assert cell[0].font.i and cell[-1].font.b
    app.undo()
    assert g.sheet.rich[(0, 0)][0] == ["Total: ", {}]


def test_plain_cell_gets_rich(app, tmp_path):
    path = rich_file(tmp_path / "r.xlsx")
    g = open_in(app, path)
    g.goto(1, 0)
    g.edit()
    g.editor.select_range(0, 2)
    key(app, "<Control-b>", g.editor)
    g.commit()
    assert g.sheet.rich[(1, 0)] == [["pl", {"bold": True}], ["ain", {}]]
    assert app.save()
    v = openpyxl.load_workbook(path, rich_text=True).active["A2"].value
    assert isinstance(v, CellRichText) and v[0].font.b


def test_typing_over_drops_runs(app, tmp_path):
    path = rich_file(tmp_path / "r.xlsx")
    g = open_in(app, path)
    app.set_cells([(0, 0, "new text")])
    assert richtext.valid(g.sheet, 0, 0, g.text(0, 0)) is None   # no longer matches: shown plain
    assert app.save()
    assert openpyxl.load_workbook(path).active["A1"].value == "new text"
