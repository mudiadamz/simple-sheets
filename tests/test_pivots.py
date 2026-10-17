"""Real Excel pivot tables: saved as pivot cache + pivot table parts; openpyxl reads them back."""
import zipfile

import openpyxl

import data_more
import pivots
import workbooks
from conftest import open_in


def sales_book(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    rows = [["Region", "Item", "Qty"], ["East", "Pens", "10"], ["West", "Pens", "5"], ["East", "Ink", "7"],
            ["West", "Ink", "1"], ["East", "Pens", "3"]]
    app.set_cells([(r + 20, c, v) for r, row in enumerate(rows) for c, v in enumerate(row)])
    g.goto(20, 0)
    return path, g


def test_pivot_workbook_xml_without_declaration(app, tmp_path, monkeypatch):
    """openpyxl without lxml writes workbook.xml with no <?xml …?> line and xmlns:r already on <workbook>: the pivot
    save must not add xmlns:r a second time (CI once failed on that)."""
    real = pivots.write

    def write_like_ci(path, app_):
        with zipfile.ZipFile(path) as z:
            files = {n: z.read(n) for n in z.namelist()}
        wb = files["xl/workbook.xml"].decode("utf-8")
        wb = wb[wb.index("?>") + 2:].lstrip() if wb.startswith("<?xml") else wb
        if "xmlns:r=" not in wb.split(">", 1)[0]:
            wb = wb.replace("<workbook", '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                                         'relationships"', 1)
        files["xl/workbook.xml"] = wb.encode("utf-8")
        with zipfile.ZipFile(path, "w") as z:
            for n, data in files.items():
                z.writestr(n, data)
        return real(path, app_)
    monkeypatch.setattr(pivots, "write", write_like_ci)
    path, g = sales_book(app, tmp_path)
    d = data_more.pivot(app)
    d.boxes["rows"].set("Region")
    d.boxes["values"].set("Qty")
    d.ok()
    assert app.save()
    with zipfile.ZipFile(path) as z:
        wbxml = z.read("xl/workbook.xml").decode()
    assert wbxml.split(">", 1)[0].count("xmlns:r=") == 1
    openpyxl.load_workbook(path)


def test_pivot_saved_as_excel_pivot(app, tmp_path):
    path, g = sales_book(app, tmp_path)
    d = data_more.pivot(app)
    d.boxes["rows"].set("Region")
    d.boxes["values"].set("Qty")
    d.boxes["agg"].set("Sum")
    name = d.ok()
    assert app.sheet == name and app.table.text(1, 1) == "20" and app.pivots[0]["sheet"] == name
    assert app.save()
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert any(n.startswith("xl/pivotCache/pivotCacheDefinition") for n in names)
        assert any(n.startswith("xl/pivotTables/pivotTable") for n in names)
        wbxml = z.read("xl/workbook.xml").decode()
        assert "<pivotCaches>" in wbxml and 'refreshOnLoad="1"' in z.read(
            next(n for n in names if n.startswith("xl/pivotCache/"))).decode()
    wb = openpyxl.load_workbook(path)                      # openpyxl parses the pivot parts
    pt = wb[name]._pivots[0]
    assert pt.cache.cacheSource.worksheetSource.ref == "A21:C26"
    assert [f.name for f in pt.cache.cacheFields] == ["Region", "Item", "Qty"]
    assert pt.dataFields[0].fld == 2 and pt.rowFields[0].x == 0
    assert app.pivots[0]["saved"]
    app.set_cells([(0, 0, "x")])
    assert app.save()                                      # saved once only (the next save keeps it)
    with zipfile.ZipFile(path) as z:
        assert len([n for n in z.namelist() if n.startswith("xl/pivotTables/pivotTable") and n.endswith(".xml")]) == 1


def test_pivot_with_columns_and_count(app, tmp_path):
    path, g = sales_book(app, tmp_path)
    d = data_more.pivot(app)
    d.boxes["rows"].set("Region")
    d.boxes["cols"].set("Item")
    d.boxes["values"].set("Qty")
    d.boxes["agg"].set("Count")
    name = d.ok()
    assert app.save()
    pt = openpyxl.load_workbook(path)[name]._pivots[0]
    assert pt.colFields[0].x == 1 and pt.dataFields[0].subtotal == "count"


def test_refresh(app, tmp_path):
    path, g = sales_book(app, tmp_path)
    src = app.sheet
    d = data_more.pivot(app)
    d.boxes["rows"].set("Region")
    d.boxes["values"].set("Qty")
    name = d.ok()
    app.show(src)
    app.set_cells([(21, 2, "100")])                        # East Pens 10 -> 100
    app.show(name)
    assert pivots.refresh(app)
    assert app.sheets[name].rows[1][1] == "110"
    app.undo()
    assert app.sheets[name].rows[1][1] == "20"


def test_plain_summary_when_not_real(app, tmp_path):
    path, g = sales_book(app, tmp_path)
    d = data_more.pivot(app)
    d.real.set(False)
    d.boxes["values"].set("Qty")
    d.ok()
    assert app.pivots == [] and app.save()
    with zipfile.ZipFile(path) as z:
        assert not any("pivot" in n for n in z.namelist())
