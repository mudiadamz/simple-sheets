"""Array formulas from Excel ({=UNIQUE(A2:A25)} over E4:E9, as in tp1-excel.xlsx): the results Excel stored in
the covered cells don't block the formula (#SPILL!), files keep them, and changing the formula clears them."""
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

from conftest import open_in


def make(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Soal"
    for i, brand in enumerate(["Honda", "Toyota", "Honda", "Nissan", "Toyota", "Ford"], start=2):
        ws[f"A{i}"] = brand
    ws["E4"] = ArrayFormula("E4:E7", "=_xlfn.UNIQUE(A2:A7)")
    for i, brand in enumerate(["Toyota", "Nissan", "Ford"], start=5):     # what Excel stores under it
        ws[f"E{i}"] = brand
    ws["F4"] = "=COUNTIF($A$2:$A$7,E4)"
    ws["G2"] = "=SUM(F4:F7)"
    ws["F12"] = ArrayFormula("F12:H12", "=TRANSPOSE(_xlfn.UNIQUE(A2:A4))")
    ws["G12"], ws["H12"] = "Toyota", "x"
    wb.save(path)
    return path


def test_no_spill_error(app, tmp_path):
    g = open_in(app, make(tmp_path / "tp1.xlsx"))
    assert g.text(3, 4) == "=UNIQUE(A2:A7)" and g.text(4, 4) == ""          # E5 is the formula's, not a value
    assert [g.value(r, 4) for r in range(3, 7)] == ["Honda", "Toyota", "Nissan", "Ford"]
    assert g.value(3, 5) == "2" and g.value(1, 6) == "2"                     # F4 / G2 work again
    assert [g.value(11, c) for c in range(5, 7)] == ["Honda", "Toyota"]       # a sideways one too
    assert not app.dirty()


def test_saving_keeps_the_array_formula(app, tmp_path):
    path = make(tmp_path / "tp1.xlsx")
    open_in(app, path)
    app.set_cells([(0, 9, "note")])
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert isinstance(ws["E4"].value, ArrayFormula) and ws["E4"].value.ref == "E4:E7"
    assert ws["E5"].value == "Toyota" and ws["J1"].value == "note"           # untouched cells stay as Excel had them


def test_changing_it_clears_the_old_results(app, tmp_path):
    path = make(tmp_path / "tp1.xlsx")
    g = open_in(app, path)
    app.set_cells([(3, 4, "=SORT(UNIQUE(A2:A7))")])
    assert app.save()
    ws = openpyxl.load_workbook(path).active
    assert ws["E5"].value is None and ws["E7"].value is None                 # nothing left to block it
    g = open_in(app, path)
    assert [g.value(r, 4) for r in range(3, 7)] == ["Ford", "Honda", "Nissan", "Toyota"]
