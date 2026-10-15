"""Freeze panes: frozen rows/columns stay on screen while scrolling."""
import openpyxl

import sheet_io as io
from conftest import key, open_in


def fill(app, rows=60, cols=12):
    app.set_cells([(r, c, f"{r},{c}") for r in range(rows) for c in range(cols)])


def test_frozen_rows_stay_while_scrolling(app):
    g = app.table
    fill(app)
    app.set_freeze(2, 0)
    assert g.fr == 2 and g.top >= 2
    for _ in range(5):
        g.on_wheel(type("E", (), {"delta": -120})())
    assert g.top > 10
    rows = g.view[0]
    assert rows[:3] == [0, 1, g.top]                        # rows 1-2 still on top, then the scrolled part
    assert g.hit(g.col_x(0) + 3, g.row_y(0) + 3) == (0, 0)
    assert g.hit(g.col_x(0) + 3, g.row_y(g.top) + 3) == (g.top, 0)
    assert len(g.cv.find_withtag("freeze")) == 1


def test_frozen_columns(app):
    g = app.table
    fill(app, cols=40)
    app.set_freeze(0, 1)
    g.goto(0, 35)                                           # far right: column A stays
    cols = [c for c, _, _ in g.view[1]]
    assert cols[0] == 0 and 35 in cols and g.left > 1
    assert g.hit(g.hdr_w() + 3, g.row_y(0) + 3) == (0, 0)


def test_moving_into_frozen_area_doesnt_scroll(app):
    g = app.table
    fill(app)
    app.set_freeze(3, 0)
    g.goto(50, 0)
    key(app, "<Control-Home>")
    assert g.cur == (0, 0) and g.top >= 3                    # scrolled part goes back to its start
    g.goto(1, 0)
    assert g.top >= 3


def test_menu_freeze_and_undo(app):
    from test_menu import build_menu
    g = app.table
    fill(app, rows=10, cols=5)
    g.goto(2, 1)
    m = build_menu(app)
    sub = m.nametowidget(m.entrycget(m.index("Freeze"), "menu"))
    labels = [sub.entrycget(i, "label") for i in range(sub.index("end") + 1)]
    assert labels == ["Up to row 3", "Up to column B", "Up to cell B3", "Unfreeze"]
    sub.invoke(2)
    assert g.sheet.freeze == (2, 1) and app.dirty()
    app.undo()
    assert g.sheet.freeze == (0, 0)


def test_insert_rows_moves_freeze(app):
    g = app.table
    fill(app, rows=10, cols=3)
    app.set_freeze(3, 0)
    app.restructure("row", 1, 2)                             # inside the frozen part: grows
    assert g.sheet.freeze == (5, 0)
    app.restructure("row", 8, 1)                             # below it: unchanged
    assert g.sheet.freeze == (5, 0)
    app.restructure("row", 0, -2)
    assert g.sheet.freeze == (3, 0)


def test_load_and_save(app, tmp_path):
    path = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active.freeze_panes = "C2"
    wb.save(path)
    assert io.load(str(path))["S"].freeze == (1, 2)
    g = open_in(app, path)
    assert g.fr == 1 and g.fc == 2
    app.set_freeze(3, 0)
    assert app.save(), app.dialogs
    assert openpyxl.load_workbook(path)["S"].freeze_panes == "A4"
    app.set_freeze(0, 0)
    assert app.save(), app.dialogs
    assert openpyxl.load_workbook(path)["S"].freeze_panes is None


def test_freeze_with_new_workbook(app, tmp_path, monkeypatch):
    import simple_sheets
    dst = tmp_path / "n.xlsx"
    monkeypatch.setattr(simple_sheets.filedialog, "asksaveasfilename", lambda **k: str(dst))
    app.set_cells([(0, 0, "head")])
    app.set_freeze(1, 0)
    assert app.save(), app.dialogs
    assert openpyxl.load_workbook(dst).active.freeze_panes == "A2"
