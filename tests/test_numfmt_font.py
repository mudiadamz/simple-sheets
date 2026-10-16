"""Number format picker, decimal buttons, font family."""
import openpyxl
import pytest

import sheet_io as io
from formatting import change_decimals, render


@pytest.mark.parametrize("fmt, delta, sample, want", [
    (None, 1, "1.5", "0.00"),
    (None, -1, "1.25", "0.0"),
    (None, 1, "12", "0.0"),
    ("#,##0.00", 1, "", "#,##0.000"),
    ("#,##0.00", -1, "", "#,##0.0"),
    ("#,##0.0", -1, "", "#,##0"),
    ("0", 1, "", "0.0"),
    ('"$"#,##0.00', -1, "", '"$"#,##0.0'),
    ("0.00%", -1, "", "0.0%"),
    ("#,##0.00;[Red](#,##0.00)", 1, "", "#,##0.000;[Red](#,##0.000)"),
    ("0.00E+00", 1, "", "0.000E+00"),
    ("yyyy-mm-dd", 1, "", "yyyy-mm-dd"),        # dates unchanged
    (None, 1, "text", None),                     # not a number
])
def test_change_decimals(fmt, delta, sample, want):
    assert change_decimals(fmt, delta, sample) == want


def test_picker_and_decimals_buttons(app):
    g, fb = app.table, app.fmtbar
    app.set_cells([(0, 0, "1234.5"), (1, 0, "0.256")])
    g.goto(0, 0)
    labels = [fb.numfmt_menu.entrycget(i, "label") for i in range(fb.numfmt_menu.index("end") + 1)
              if fb.numfmt_menu.type(i) != "separator"]
    fb.numfmt_menu.invoke(labels.index("Currency   $1,234.56"))
    assert g.style(0, 0)["numfmt"] == '"$"#,##0.00' and g.value(0, 0) == "1234.5"
    assert [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")][0] == "$1,234.50"
    fb.decimals(-1)
    assert g.style(0, 0)["numfmt"] == '"$"#,##0.0'
    g.goto(1, 0)
    fb.numfmt_menu.invoke(labels.index("Percent   12%"))
    assert render("0.256", g.style(1, 0)["numfmt"]) == "26%"
    fb.decimals(1)
    assert render("0.256", g.style(1, 0)["numfmt"]) == "25.6%"
    fb.numfmt_menu.invoke(labels.index("General"))
    assert "numfmt" not in (g.style(1, 0) or {})
    app.undo()
    assert g.style(1, 0)["numfmt"] == "0.0%"


def test_font_family(app, tmp_path):
    from conftest import open_in
    from workbooks import plain
    path = plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    fb = app.fmtbar
    assert fb.font.get() == "Calibri"
    g.goto(1, 0)
    fb.font.set("Courier New")
    fb.set_font()
    assert g.style(1, 0)["font"] == "Courier New"
    assert g.font(g.style(1, 0)).actual("family") == "Courier New"     # drawn in that font
    assert "Courier New" in __import__("formatting").describe(app.current.styles.get(app.table.cur))
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Data"]
    assert ws["A2"].font.name == "Courier New"
    assert ws["A1"].font.name is None and ws["A1"].font.b          # untouched cell: as it was
    assert io.load(str(path))["Data"].styles[(1, 0)]["font"] == "Courier New"
