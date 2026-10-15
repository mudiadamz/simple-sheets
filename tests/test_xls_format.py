"""Old .xls files: formatting, widths, heights, merges, hidden rows/columns/sheets, frozen panes."""
import pytest

import sheet_io as io

xlwt = pytest.importorskip("xlwt")


def old_file(path):
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Old")
    head = xlwt.easyxf("font: bold on, height 280, colour red, name Verdana; pattern: pattern solid, fore_colour yellow;"
                       "alignment: horizontal center, vertical top, wrap on;"
                       "borders: bottom medium, bottom_colour blue")
    ws.write(0, 0, "Title", head)
    ws.write(1, 0, 1234.5, xlwt.easyxf(num_format_str="#,##0.00"))
    ws.write(1, 1, "slanted", xlwt.easyxf("font: italic on, underline single, struck_out on"))
    ws.write(2, 0, "plain")
    ws.write_merge(3, 3, 0, 2, "merged")
    ws.col(0).width = 256 * 20
    ws.col(3).hidden = True
    ws.row(1).height_mismatch = True
    ws.row(1).height = 20 * 30
    ws.row(2).hidden = True
    ws.panes_frozen = True
    ws.horz_split_pos = 1
    ws.vert_split_pos = 1
    hidden = wb.add_sheet("Gone")
    hidden.write(0, 0, "x")
    hidden.visibility = 1
    wb.save(str(path))
    return path


def test_xls_formatting(tmp_path):
    book = io.load(str(old_file(tmp_path / "old.xls")))
    s = book["Old"]
    head = s.styles[(0, 0)]
    assert head["bold"] and head["size"] == 14 and head["fg"] == "#ff0000" and head["font"] == "Verdana"
    assert head["bg"] == "#ffff00" and head["align"] == "center" and head["valign"] == "top" and head["wrap"]
    assert head["border"]["bottom"] == ["medium", "#0000ff"]
    assert s.styles[(1, 0)]["numfmt"] == "#,##0.00"
    assert s.styles[(1, 1)] == {"italic": True, "underline": True, "strike": True}
    assert (2, 0) not in s.styles                                  # default look: nothing stored
    assert s.widths[0] == 20 and 3 in s.hidden_cols
    assert s.heights[1] == 30 and 2 in s.hidden_rows
    assert s.merges == [(3, 0, 3, 2)]
    assert s.freeze == (1, 1)
    assert book["Gone"].hidden and not s.hidden


def test_xls_opens_in_app_and_saves_as_xlsx(app, tmp_path):
    import openpyxl

    from conftest import open_in
    g = open_in(app, old_file(tmp_path / "old.xls"))
    assert g.style(0, 0)["bold"] and not app.dirty()
    dst = tmp_path / "new.xlsx"
    assert app.write(str(dst))
    ws = openpyxl.load_workbook(dst)["Old"]
    assert ws["A1"].font.b and ws["A1"].fill.fgColor.rgb.endswith("FFFF00")
    assert "A4:C4" in [str(m) for m in ws.merged_cells.ranges]


def formula_file(path):
    wb = xlwt.Workbook()
    ws, other = wb.add_sheet("Main"), wb.add_sheet("Other")
    for r, v in enumerate([1, 2, 3]):
        ws.write(r, 0, v)
    ws.write(0, 1, 10)
    other.write(0, 0, 7)
    for i, f in enumerate(["SUM(A1:A3)", "A1*2+B1", 'IF(A1>1,"big","small")', "ROUND(B1/3,2)", "Other!A1*3",
                           "AVERAGE($A$1:A3)", 'CONCATENATE("x",A2)', "-A3^2", "10%"]):
        ws.write(i, 3, xlwt.Formula(f))
    wb.save(str(path))
    return path


def test_xls_formulas_read(tmp_path):
    import xls_formulas
    path = formula_file(tmp_path / "f.xls")
    found = xls_formulas.read(str(path))[0]
    assert found[(4, 3)] == "=Other!A1*3" and found[(5, 3)] == "=AVERAGE($A$1:A3)" and found[(8, 3)] == "=10%"
    s = io.load(str(path))["Main"]
    assert [s.rows[i][3] for i in range(3)] == ["=SUM(A1:A3)", "=A1*2+B1", '=IF(A1>1,"big","small")']


def test_xls_formula_kept_only_when_it_matches(tmp_path):
    import xls_formulas
    from sheet_io import Sheet
    s = Sheet([["2", "4"]])
    book = {"Main": s}
    xls_formulas.read = lambda path: [{(0, 1): "=A1*2"}]          # matches the saved 4
    try:
        assert xls_formulas.apply(book, "x") == 1 and s.rows[0][1] == "=A1*2"
        s.rows[0][1] = "5"
        xls_formulas.read = lambda path: [{(0, 1): "=A1*2"}]      # gives 4, Excel saved 5: keep 5
        assert xls_formulas.apply(book, "x") == 0 and s.rows[0][1] == "5"
    finally:
        import importlib
        importlib.reload(xls_formulas)


def test_xls_formulas_in_app(app, tmp_path):
    from conftest import open_in
    g = open_in(app, formula_file(tmp_path / "f.xls"))
    assert g.text(0, 3) == "=SUM(A1:A3)" and g.value(0, 3) == "6" and g.value(4, 3) == "21"
    assert g.value(2, 3) == "small" and g.value(7, 3) == "9"   # Excel: (-3)^2
