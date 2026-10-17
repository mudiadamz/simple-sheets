"""Files: password-protected workbooks, .xlsb, spelling."""
import io as pyio
import struct
import sys
import zipfile

import openpyxl
import pytest
from simplecore import office_crypto

import sheet_io as io
import xlsb
from conftest import open_in

# ---------- password-protected ----------

def encrypted(tmp_path, password="pa55"):
    msoff = pytest.importorskip("msoffcrypto.format.ooxml")    # only used to make the test file
    wb = openpyxl.Workbook()
    wb.active.title = "Secret"
    wb.active["A1"] = "classified"
    wb.active["B1"] = "=2*21"
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    enc = tmp_path / "locked.xlsx"
    with open(plain, "rb") as f, open(enc, "wb") as o:
        msoff.OOXMLFile(f).encrypt(password, o)
    return enc


def test_aes_vectors():
    assert office_crypto.AES(bytes(range(16))).decrypt_block(bytes.fromhex("69c4e0d86a7b0430d8cdb78070b4c55a")) \
        == bytes.fromhex("00112233445566778899aabbccddeeff")
    assert office_crypto.AES(bytes(range(24))).decrypt_block(bytes.fromhex("dda97ca4864cdfe06eaf70a0ec0d7191")) \
        == bytes.fromhex("00112233445566778899aabbccddeeff")


def test_decrypt(tmp_path):
    enc = encrypted(tmp_path)
    assert office_crypto.is_encrypted(str(enc)) and not office_crypto.is_encrypted(str(tmp_path / "plain.xlsx"))
    data = office_crypto.decrypt(str(enc), "pa55")
    assert openpyxl.load_workbook(pyio.BytesIO(data))["Secret"]["A1"].value == "classified"
    with pytest.raises(office_crypto.WrongPassword):
        office_crypto.decrypt(str(enc), "wrong")


def test_open_protected_in_app(app, tmp_path, monkeypatch):
    enc = encrypted(tmp_path)
    import simple_sheets
    answers = iter(["wrong", "pa55"])
    monkeypatch.setattr(simple_sheets.simpledialog, "askstring", lambda *a, **k: next(answers))
    assert app.open_file(str(enc))
    g = app.table
    assert g.text(0, 0) == "classified" and g.value(0, 1) == "42"
    assert app.path == str(enc) and app.password == "pa55"
    app.set_cells([(1, 0, "more")])
    assert app.save()                                           # saved in place, still protected
    assert office_crypto.is_encrypted(str(enc))
    wb = openpyxl.load_workbook(pyio.BytesIO(office_crypto.decrypt(str(enc), "pa55")))
    assert wb["Secret"]["A2"].value == "more" and wb["Secret"]["A1"].value == "classified"
    app.set_cells([(2, 0, "again")])
    assert app.save()                                           # a second save patches the plain copy again
    wb = openpyxl.load_workbook(pyio.BytesIO(office_crypto.decrypt(str(enc), "pa55")))
    assert wb["Secret"]["A3"].value == "again" and wb["Secret"]["A2"].value == "more"
    msoff = pytest.importorskip("msoffcrypto")
    with open(enc, "rb") as f:                                  # another reader accepts it (integrity too)
        of = msoff.OfficeFile(f)
        of.load_key(password="pa55")
        out = pyio.BytesIO()
        of.decrypt(out, verify_integrity=True)
    assert openpyxl.load_workbook(out)["Secret"]["A3"].value == "again"
    import os
    tmp = app.unlocked_files[0]
    assert os.path.exists(tmp)
    app.dirty = lambda: False
    destroyed = []
    monkeypatch.setattr(app, "destroy", lambda: destroyed.append(True))
    app.on_close()
    assert destroyed and not os.path.exists(tmp)                # closing deletes the decrypted copy


def test_set_and_remove_password(app, tmp_path):
    import workbooks
    path = workbooks.plain(tmp_path / "p.xlsx")
    app.open_file(str(path))
    assert app.set_password("abc") and app.dirty()
    assert app.save()
    assert office_crypto.is_encrypted(str(path))
    assert openpyxl.load_workbook(pyio.BytesIO(office_crypto.decrypt(str(path), "abc"))).active["A1"].value
    app.set_cells([(5, 0, "x")])
    assert app.save()
    assert app.set_password("") and app.save()
    assert not office_crypto.is_encrypted(str(path))
    assert openpyxl.load_workbook(path).active["A6"].value == "x"


def test_cfb_round_trip():
    tree = {"Small": b"abc" * 10, "Big": bytes(range(256)) * 40, "Store": {"Inner": b"in"}}
    cf = office_crypto.CompoundFile(office_crypto.write_cfb(tree))
    assert cf.stream("Small") == b"abc" * 10 and cf.stream("Big") == bytes(range(256)) * 40
    assert cf.stream("Inner") == b"in"


def test_cancel_password(app, tmp_path, monkeypatch):
    enc = encrypted(tmp_path)
    import simple_sheets
    monkeypatch.setattr(simple_sheets.simpledialog, "askstring", lambda *a, **k: None)
    assert app.open_file(str(enc)) is False and app.path is None


# ---------- .xlsb ----------

def rec(rtype, payload=b""):
    out = bytearray()
    if rtype >= 0x80:
        out += bytes([(rtype & 0x7F) | 0x80, rtype >> 7])
    else:
        out.append(rtype)
    n = len(payload)
    while True:
        b = n & 0x7F
        n >>= 7
        out.append(b | (0x80 if n else 0))
        if not n:
            break
    return bytes(out) + payload


def ws(s):
    return struct.pack("<I", len(s)) + s.encode("utf-16-le")


def cell(col, style=0):
    return struct.pack("<II", col, style)


def xlsb_file(path):
    workbook = rec(0x83) + rec(0x9C, struct.pack("<II", 0, 1) + ws("rId1") + ws("Data")) + \
        rec(0x9C, struct.pack("<II", 1, 2) + ws("rId2") + ws("Hidden")) + rec(0x84)
    sheet1 = b"".join([
        rec(0x00, struct.pack("<I", 0) + b"\0" * 8), rec(0x07, cell(0) + struct.pack("<I", 0)),
        rec(0x02, cell(1) + struct.pack("<I", (12 << 2) | 0x02)),               # RK integer 12
        rec(0x05, cell(2) + struct.pack("<d", 3.25)), rec(0x04, cell(3) + b"\x01"),
        rec(0x00, struct.pack("<I", 1) + b"\0" * 8), rec(0x06, cell(0) + ws("inline")),
        rec(0x03, cell(1) + b"\x07"), rec(0x05, cell(2, 1) + struct.pack("<d", 45356.0)),
        rec(0x09, cell(3) + struct.pack("<d", 99.5) + b"\0\0"),
        rec(0x02, cell(4) + struct.pack("<I", (1234 << 2) | 0x03)),             # RK 1234 / 100
    ])
    sst = rec(0x13, b"\0" + ws("shared text"))
    styles = rec(0x269) + rec(0x2F, struct.pack("<HH", 0, 0) + b"\0" * 12) + \
        rec(0x2F, struct.pack("<HH", 0, 14) + b"\0" * 12) + rec(0x26A)
    rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Target="worksheets/sheet1.bin"/>'
            '<Relationship Id="rId2" Target="worksheets/sheet2.bin"/></Relationships>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.bin", workbook)
        z.writestr("xl/_rels/workbook.bin.rels", rels)
        z.writestr("xl/worksheets/sheet1.bin", sheet1)
        z.writestr("xl/worksheets/sheet2.bin", rec(0x00, struct.pack("<I", 0) + b"\0" * 8) +
                   rec(0x06, cell(0) + ws("h")))
        z.writestr("xl/sharedStrings.bin", sst)
        z.writestr("xl/styles.bin", styles)
    return path


def test_xlsb_values(tmp_path):
    book = io.load(str(xlsb_file(tmp_path / "b.xlsb")))
    assert list(book) == ["Data", "Hidden"] and book["Hidden"].hidden and not book["Data"].hidden
    assert book["Data"].rows[0] == ["shared text", "12", "3.25", "TRUE"]
    assert book["Data"].rows[1] == ["inline", "#DIV/0!", "2024-03-05", "99.5", "12.34"]


def test_xlsb_in_app_saves_as_xlsx(app, tmp_path):
    g = open_in(app, xlsb_file(tmp_path / "b.xlsb"))
    assert g.text(0, 0) == "shared text"
    app.set_cells([(2, 0, "new")])
    assert app.write(str(tmp_path / "b.xlsx"))
    assert openpyxl.load_workbook(tmp_path / "b.xlsx")["Data"]["A3"].value == "new"


def test_rk_numbers():
    assert xlsb.rk((5 << 2) | 2) == 5.0 and xlsb.rk((-3 << 2 & 0xFFFFFFFF) | 2) == -3.0
    hi = struct.unpack("<Q", struct.pack("<d", 1.5))[0] >> 32
    assert xlsb.rk(hi & 0xFFFFFFFC) == 1.5


# ---------- spelling ----------

class FakeChecker:
    words = {"speling": ["spelling"], "teh": ["the"]}

    def __init__(self):
        self.ignored, self.added = set(), []

    def check(self, text):
        import re
        return [(m.start(), len(m.group()), m.group()) for m in re.finditer(r"\w+", text)
                if m.group() in self.words and m.group().casefold() not in self.ignored]

    def suggest(self, word):
        return self.words.get(word, [])

    def add(self, word):
        self.added.append(word)

    def ignore(self, word):
        self.ignored.add(word.casefold())


def test_spell_dialog_flow(app):
    import spell
    g = app.table
    app.set_cells([(0, 0, "teh speling"), (1, 0, "=1+1"), (2, 0, "teh end"), (3, 0, "fine")])
    g.goto(0, 0)
    d = spell.SpellDialog.open(app, FakeChecker())
    assert d.current[4] == "teh" and d.change_to.get() == "the"
    d.change_all()                                    # teh -> the here and later
    assert g.text(0, 0) == "the speling" and d.current[4] == "speling"
    d.change()
    assert g.text(0, 0) == "the spelling"
    assert g.text(2, 0) == "the end"                  # changed by "Change all" on the way
    assert app.dialogs[-1][:2] == ("showinfo", "Spelling")


@pytest.mark.skipif(sys.platform != "win32", reason="Windows spell checker")
def test_windows_spell_checker():
    import spell
    try:
        c = spell.Checker("en-US")
    except OSError:
        pytest.skip("no English spell checker installed")
    errs = c.check("This sentense has a mistake")
    assert errs == [(5, 8, "sentense")] and "sentence" in c.suggest("sentense")
    c.ignore("sentense")
    assert c.check("sentense") == []
