"""Column widths changed in the app (drag / double-click auto-fit) are saved and undoable."""
from types import SimpleNamespace

import openpyxl

import sheet_io as io
import workbooks
from conftest import open_in


def drag_border(app, c, dx):
    g = app.table
    app.update()
    x = g.col_x(c) + g.colw(c)          # right edge of column c in the header
    y = g.rowh // 2
    g.cv.event_generate("<ButtonPress-1>", x=x, y=y, when="now")
    g.cv.event_generate("<B1-Motion>", x=x + dx, y=y, when="now")
    g.cv.event_generate("<ButtonRelease-1>", x=x + dx, y=y, when="now")
    app.update()


def test_drag_width_is_saved(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    w0 = g.colw(1)
    drag_border(app, 1, g.px(80))
    assert g.colw(1) > w0 and 1 in g.sheet.widths and not g.sheet.px   # a real width, not view-only
    assert app.dirty()
    assert app.save(), app.dialogs
    width = openpyxl.load_workbook(path)["Data"].column_dimensions["B"].width
    assert width == g.sheet.widths[1] and width > 13
    assert io.load(str(path))["Data"].widths[1] == width              # reopened: same width
    assert abs(g.colw(1) - (w0 + g.px(80))) <= 2


def test_width_undo_redo(app):
    g = app.table
    w0 = g.colw(2)
    drag_border(app, 2, -g.px(30))
    narrow = g.colw(2)
    assert narrow < w0
    app.undo()
    assert g.colw(2) == w0 and 2 not in g.sheet.widths and not app.dirty()
    app.redo()
    assert g.colw(2) == narrow


def test_autofit_is_saved(app, tmp_path):
    path = workbooks.plain(tmp_path / "p.xlsx")
    g = open_in(app, path)
    app.set_cells([(5, 0, "a much longer piece of text than the column")])
    app.update()
    x = g.col_x(0) + g.colw(0)
    g.on_double(SimpleNamespace(x=x, y=g.rowh // 2))
    assert g.sheet.widths[0] > 20
    assert app.save(), app.dialogs
    assert openpyxl.load_workbook(path)["Data"].column_dimensions["A"].width == g.sheet.widths[0]


def test_width_inside_a_column_group(app, tmp_path):
    path = workbooks.hidden(tmp_path / "h.xlsx")          # E:G share one dimension (width 15)
    open_in(app, path)
    app.set_width(5, 30)                                    # only F
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["H"]
    assert ws.column_dimensions["F"].width == 30
    again = io.load(str(path))["H"]
    assert again.widths[4] == 15 and again.widths[5] == 30 and again.widths[6] == 15


def test_width_after_inserting_columns(app, tmp_path):
    path = workbooks.layout(tmp_path / "m.xlsx")           # column C width 20
    open_in(app, path)
    app.restructure("col", 0, 1)                            # C moves to D (sheet rebuilt on save)
    app.set_width(1, 25)                                    # new width for (now) B
    assert app.save(), app.dialogs
    ws = openpyxl.load_workbook(path)["Main"]
    assert ws.column_dimensions["D"].width == 20 and ws.column_dimensions["B"].width == 25


def test_reset_width_to_default(app, tmp_path):
    path = workbooks.layout(tmp_path / "m.xlsx")
    open_in(app, path)
    app.set_width(2, None)                                  # column C back to default
    assert app.save(), app.dialogs
    assert 2 not in io.load(str(path))["Main"].widths
