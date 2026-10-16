"""=IMAGE(): a web picture inside its cell; converting JPEG / BMP with Windows; Excel's _xlfn. prefix in files."""
import base64
import struct
import sys
import time
import zlib

import openpyxl
import pytest

import cell_images
import formatting
import sheet_io as io
from conftest import open_in
from formula import Engine
from sheet_io import Sheet


def png(w, h, rgb=(255, 0, 0)):
    """A tiny solid PNG (no Pillow)."""
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def bmp(w, h, bgr=(0, 0, 255)):
    row = bytes(bgr) * w + b"\x00" * ((4 - (w * 3) % 4) % 4)
    pixels = row * h
    return (b"BM" + struct.pack("<IHHI", 54 + len(pixels), 0, 0, 54)
            + struct.pack("<IiiHHIIiiII", 40, w, h, 1, 24, 0, len(pixels), 2835, 2835, 0, 0) + pixels)


@pytest.fixture
def ev():
    data = Sheet([["https://example.com/a.png"]])
    eng = Engine()
    eng.set_book({"S": data})

    def run(formula):
        data.rows.append([formula])
        eng.clear()
        try:
            return eng.display(data, len(data.rows) - 1, 0)
        finally:
            data.rows.pop()
    return run


def test_image_value(ev):
    v = ev('=IMAGE(A1, "Logo", 3, 40, 60)')
    assert v == "Logo" and v.url == "https://example.com/a.png" and (v.sizing, v.height, v.width) == (3, 40, 60)
    assert ev("=IMAGE(A1)") == "" and ev("=IMAGE(A1)").sizing == 0
    assert ev('=LEN(IMAGE(A1, "Logo"))') == "4"                     # other formulas see the alt text
    assert ev('=IMAGE("")') == "#VALUE!" and ev("=IMAGE(A1,,7)") == "#VALUE!" and ev("=IMAGE(A1,,3)") == "#VALUE!"


def test_google_sheets_form(ev):
    """IMAGE(url, mode, height, width) from Google Sheets: 1 fit, 2 stretch, 3 original, 4 custom."""
    v = ev("=IMAGE(A1,4,10,15)")                                        # as in world_gdp.xlsx
    assert v.url == "https://example.com/a.png" and (v.sizing, v.height, v.width) == (3, 10, 15) and v == ""
    assert ev("=IMAGE(A1,2)").sizing == 1 and ev("=IMAGE(A1,3)").sizing == 2 and ev("=IMAGE(A1,1)").sizing == 0
    assert ev("=IMAGE(A1,4)") == "#VALUE!"                              # custom needs a size
    v = ev('=IMAGE(A1,"4",1)')                                          # text "4": Excel's alt text
    assert v == "4" and v.sizing == 1
    assert ev("=IMAGE(A1,4,10,15,20)") == "#VALUE!"                     # Excel form: sizing 10 doesn't exist


def test_sizing_boxes():
    v = cell_images.ImageValue("u")
    assert cell_images.box(v, 100, 50, 200, 200, int) == (50, 50)         # fit, keeping its shape
    v.sizing = 1
    assert cell_images.box(v, 100, 50, 200, 200, int) == (100, 50)        # fill
    v.sizing = 2
    assert cell_images.box(v, 100, 50, 30, 20, int) == (30, 20)           # original
    v.sizing, v.height, v.width = 3, 10, 40
    assert cell_images.box(v, 100, 50, 30, 20, int) == (40, 10)           # custom


def test_picture_drawn_in_the_cell(app, monkeypatch):
    got = []
    monkeypatch.setattr(cell_images, "fetch", lambda url, timeout=15: got.append(url) or png(40, 20))
    g = app.table
    app.set_cells([(1, 1, '=IMAGE("https://example.com/red.png", "Red")')])
    g.redraw()
    shown = [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")]
    assert "Red" in shown or g.cv.find_withtag("cellimage")              # the alt text while it downloads
    end = time.monotonic() + 3
    while not g.cv.find_withtag("cellimage") and time.monotonic() < end:
        app.update()
        time.sleep(0.02)
    items = g.cv.find_withtag("cellimage")
    assert items and got == ["https://example.com/red.png"]               # downloaded once
    x, y = g.cv.coords(items[0])
    assert g.col_x(1) < x < g.col_x(2) and g.row_y(1) < y < g.row_y(2)     # inside B2
    assert "Red" not in [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")]   # picture, not alt text
    g.redraw()
    assert got == ["https://example.com/red.png"]                         # cached


def test_failed_picture_shows_alt_text(app, monkeypatch):
    def boom(url, timeout=15):
        raise OSError("404")
    monkeypatch.setattr(cell_images, "fetch", boom)
    g = app.table
    app.set_cells([(0, 0, '=IMAGE("https://example.com/gone.png", "Company logo")')])
    end = time.monotonic() + 3
    while cell_images.state("https://example.com/gone.png")[0] != "error" and time.monotonic() < end:
        g.redraw()
        app.update()
        time.sleep(0.02)
    g.redraw()
    assert "Company logo" in [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")]


def test_data_url_and_formats():
    data = png(3, 2)
    assert cell_images.data_url("data:image/png;base64," + base64.b64encode(data).decode()) == data
    assert cell_images.data_url("data:text/plain,a%20b") == b"a b"
    assert cell_images.to_png_or_gif(data) == data                        # PNG as it is
    gif = b"GIF89a" + b"\x00" * 20
    assert cell_images.to_png_or_gif(gif) == gif
    assert cell_images.to_png_or_gif(b"not a picture") is None


@pytest.mark.skipif(sys.platform != "win32", reason="Windows Imaging Component")
def test_wic_decodes_pictures():
    import pillowless
    out = cell_images.wic_png(bmp(6, 3, bgr=(255, 0, 0)))               # blue BMP
    assert pillowless.sniff(out) == ("PNG", (6, 3))
    src = cell_images.png_bytes(2, 1, bytearray([255, 0, 0, 255, 0, 255, 0, 128]))
    assert pillowless.sniff(src) == ("PNG", (2, 1))
    with pytest.raises(OSError):                                       # WIC's error codes come as OSError
        cell_images.wic_png(b"definitely not a picture")


def test_reasons_in_plain_words(app, monkeypatch):
    assert "web page" in cell_images.why_not(b"<!DOCTYPE html><html><body>")
    assert "SVG" in cell_images.why_not(b'<?xml version="1.0"?><svg xmlns="x"></svg>')
    assert cell_images.to_png_or_gif(b"<html></html>") is None
    monkeypatch.setattr(cell_images, "fetch", lambda url, timeout=15: b"<!doctype html><html>search results</html>")
    g = app.table
    app.set_cells([(0, 0, '=IMAGE("https://www.google.com/search?q=dog")')])
    end = time.monotonic() + 3
    while cell_images.state("https://www.google.com/search?q=dog")[0] != "error" and time.monotonic() < end:
        g.redraw()
        app.update()
        time.sleep(0.02)
    g.goto(0, 0)
    app.on_select()
    assert app.status["text"].startswith("Picture not shown: the link opens a web page")
    g.redraw()
    assert "#picture?" in [g.cv.itemcget(i, "text") for i in g.cv.find_withtag("txt")]


@pytest.mark.skipif(sys.platform != "win32", reason="Windows GDI+")
def test_windows_converts_bmp_to_png():
    out = cell_images.gdiplus_png(bmp(5, 4))
    assert out[:8] == b"\x89PNG\r\n\x1a\n"
    import pillowless
    assert pillowless.sniff(out)[1] == (5, 4)
    assert cell_images.to_png_or_gif(bmp(2, 2))[:4] == b"\x89PNG"


def test_xlfn_prefixes():
    assert formatting.to_file('=XLOOKUP(A1,B:B,C:C)') == "=_xlfn.XLOOKUP(A1,B:B,C:C)"
    assert formatting.to_file('=IMAGE("https://x/SUM(1).png")') == '=_xlfn.IMAGE("https://x/SUM(1).png")'
    assert formatting.to_file("=SUM(FILTER(A:A,B:B>1))") == "=SUM(_xlfn._xlws.FILTER(A:A,B:B>1))"
    assert formatting.to_file('=CONCAT("XLOOKUP(", A1)') == '=_xlfn.CONCAT("XLOOKUP(", A1)'   # strings untouched
    assert formatting.to_file("=STDEV.S(A:A)+SUM(B:B)") == "=_xlfn.STDEV.S(A:A)+SUM(B:B)"
    assert formatting.from_file("=_xlfn.XLOOKUP(1,_xlfn._xlws.FILTER(A:A,1),C:C)") == "=XLOOKUP(1,FILTER(A:A,1),C:C)"
    assert formatting.from_file('="_xlfn.X(" & 1') == '="_xlfn.X(" & 1'
    for f in ("=XLOOKUP(1,A:A,B:B)", '=IMAGE("u","a",0)', "=SUM(A1:A3)"):
        assert formatting.from_file(formatting.to_file(f)) == f


def test_files_keep_prefixes_out_of_sight(app, tmp_path):
    src = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=_xlfn.XLOOKUP(1,B1:B2,C1:C2)"
    wb.save(src)
    g = open_in(app, src)
    assert g.text(0, 0) == "=XLOOKUP(1,B1:B2,C1:C2)"                       # no _xlfn. in the formula bar
    app.set_cells([(1, 0, '=IMAGE("https://example.com/a.png","A")')])
    assert app.save()
    ws = openpyxl.load_workbook(src).active
    assert ws["A2"].value == '=_xlfn.IMAGE("https://example.com/a.png","A")'   # Excel knows it
    assert io.to_value("=SUM(A1)") == "=SUM(A1)"
