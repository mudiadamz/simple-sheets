"""Cell scripts: clicking a cell runs its script (a button in the sheet); saved in the file; trust first."""
import time

import openpyxl

import cell_scripts
import settings
import sheet_props
from conftest import open_in
from test_menu import menu_states

GO_TODAY = "sheet['B2'] = 2026\nsheet['D2'] = 'October'\nprint('Jumped from', clicked)\n"


CLOCK = [int(time.monotonic() * 1000)]


def click(app, r, c, to=None):
    g = app.table
    CLOCK[0] += 5000                                    # clicks far apart in time: never a double-click
    t = CLOCK[0]
    x, y = g.col_x(c) + 5, g.row_y(r) + 5
    g.cv.event_generate("<Button-1>", x=x, y=y, time=t, when="now")
    if to:                                              # dragged to another cell before letting go
        g.cv.event_generate("<B1-Motion>", x=g.col_x(to[1]) + 5, y=g.row_y(to[0]) + 5, time=t + 1, when="now")
        x, y = g.col_x(to[1]) + 5, g.row_y(to[0]) + 5
    g.cv.event_generate("<ButtonRelease-1>", x=x, y=y, time=t + 2, when="now")
    app.update()


def calendar(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cal"
    ws["A2"], ws["B2"], ws["C2"], ws["D2"], ws["G2"] = "Year:", 2030, "Month:", "March", "📍 TODAY"
    sheet_props.write_scripts(wb, {"Cal": {"G2": GO_TODAY}})
    wb.save(path)
    return path


def test_assign_click_and_undo(app):
    g = app.table
    app.set_cells([(1, 6, "Go")])
    cell_scripts.assign(app, 1, 6, "sheet['A1'] = 'clicked ' + clicked")
    assert cell_scripts.get(app, 1, 6).startswith("sheet['A1']") and (1, 6) in g.script_cells
    assert g.cv.find_withtag("script")                                    # the corner marker
    click(app, 1, 6)
    assert g.text(0, 0) == "clicked G2" and "▶ G2" in app.status["text"]
    app.set_cells([(0, 0, "")])
    click(app, 1, 6, to=(3, 6))                                           # a drag doesn't run it
    assert g.text(0, 0) == ""
    click(app, 2, 6)                                                      # other cells don't
    assert g.text(0, 0) == ""
    app.undo()                                                            # the empty A1, then the script
    app.undo()
    app.undo()
    assert cell_scripts.get(app, 1, 6) is None and not g.script_cells


def test_failing_script_changes_nothing(app, monkeypatch):
    import cell_scripts as cs
    shown = []
    monkeypatch.setattr("tkinter.messagebox.showerror", lambda *a, **k: shown.append(a))
    cs.assign(app, 0, 0, "sheet['B1'] = 'x'\n1/0")
    ok, msg = cs.run(app, 0, 0)
    assert not ok and "ZeroDivisionError" in msg and shown and app.table.text(0, 1) == ""


def test_saved_in_the_file_and_asked_before_running(app, tmp_path, monkeypatch):
    path = calendar(tmp_path / "cal.xlsx")
    asked = []
    monkeypatch.setattr("tkinter.messagebox.askyesno", lambda *a, **k: asked.append(a) or False)
    g = open_in(app, path)
    assert (1, 6) in g.script_cells
    click(app, 1, 6)
    assert len(asked) == 1 and "cell script" in asked[0][1] and g.text(1, 1) == "2030"   # said no: nothing ran
    monkeypatch.setattr("tkinter.messagebox.askyesno", lambda *a, **k: asked.append(a) or True)
    click(app, 1, 6)
    assert len(asked) == 2 and g.text(1, 1) == "2026" and g.text(1, 3) == "October"
    app.set_cells([(1, 1, "1999")])
    click(app, 1, 6)
    assert len(asked) == 2 and g.text(1, 1) == "2026"                    # allowed once: not asked again
    assert cell_scripts.fingerprint(app) in settings.get(cell_scripts.TRUST_KEY)
    # changed code is new code: asked again
    data = {"Cal": {"G2": "sheet['A1'] = 'other'"}}
    wb = openpyxl.load_workbook(path)
    sheet_props.write_scripts(wb, data)
    wb.save(tmp_path / "other.xlsx")
    open_in(app, tmp_path / "other.xlsx")
    click(app, 1, 6)
    assert len(asked) == 3


def test_save_keeps_scripts_and_long_code(app, tmp_path):
    long_code = "x = 1\n" + "# a fairly long comment line for the property parts\n" * 40 + "sheet['A1'] = x\n"
    app.set_cells([(0, 0, "a")])
    cell_scripts.assign(app, 4, 2, long_code)
    dst = tmp_path / "saved.xlsx"
    assert app.write(str(dst))
    read = sheet_props.read_book(str(dst))["scripts"]
    assert read == {app.sheet: {"C5": long_code}}
    wb = openpyxl.load_workbook(dst)
    parts = [p for p in wb.custom_doc_props.props if p.name.startswith(sheet_props.SCRIPTS_PROP)]
    assert len(parts) > 1 and all(len(p.value) <= 255 for p in parts)    # Excel-sized pieces
    open_in(app, dst)
    assert cell_scripts.get(app, 4, 2) == long_code
    cell_scripts.assign(app, 4, 2, None)                                  # removed: gone from the file too
    app.save()
    assert "scripts" not in sheet_props.read_book(str(dst))


def test_rename_sheet_keeps_its_scripts(app):
    cell_scripts.assign(app, 0, 0, "print(1)")
    old = app.sheet
    assert app.rename_sheet(old, "Buttons")
    assert cell_scripts.get(app, 0, 0) == "print(1)\n" and "Buttons" in app.book_props["scripts"]
    app.undo()                                                            # one step: name and scripts
    assert app.sheet == old and cell_scripts.get(app, 0, 0) == "print(1)\n"


def test_scripts_can_make_buttons(app):
    import scripting
    ok, _ = scripting.run(app, 'sheet.on_click("G2", "sheet[\'A1\'] = clicked")')
    assert ok and cell_scripts.get(app, 1, 6) and cell_scripts.trusted(app)
    click(app, 1, 6)
    assert app.table.text(0, 0) == "G2"
    ok, _ = scripting.run(app, 'sheet.on_click("G2", None)')
    assert ok and cell_scripts.get(app, 1, 6) is None


def test_right_click_menu_and_editor(app, monkeypatch):
    app.table.goto(1, 6)
    assert "Run a script on click…" in menu_states(app)
    ed = cell_scripts.open_editor(app)
    assert "clicked" in ed.code()                                         # starts from an example
    ed.text.delete("1.0", "end")
    ed.text.insert("1.0", "sheet['A1'] = 7")
    monkeypatch.setattr("tkinter.messagebox.showinfo", lambda *a, **k: None)
    assert ed.test() and app.table.text(0, 0) == "7"
    assert ed.save() and cell_scripts.get(app, 1, 6) == "sheet['A1'] = 7\n"
    states = menu_states(app)
    assert {"Edit cell script…", "Run cell script", "Remove cell script"} <= set(states)
