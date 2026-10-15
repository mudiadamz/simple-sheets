"""Builds docs/showcase.xlsx: a good-looking sample workbook for screenshots of Simple Sheets.

Sheets: Dashboard (KPI cards, monthly table with colour scale / data bars, charts), Orders (an Excel table),
Roadmap (status colours, progress bars). Uses openpyxl only. Run: python tools/make_showcase.py [out.xlsx]
"""
import datetime as dt
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

INK, MUTED, LINE = "1F2937", "6B7280", "E5E7EB"
NAVY, TEAL, VIOLET, AMBER, ROSE = "111827", "0D9488", "7C3AED", "D97706", "E11D48"
SOFT = {TEAL: "CCFBF1", VIOLET: "EDE9FE", AMBER: "FEF3C7", ROSE: "FFE4E6"}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def fill(hex6):
    return PatternFill("solid", fgColor=hex6)


def thin(color=LINE):
    s = Side(style="thin", color=color)
    return Border(top=s, bottom=s, left=s, right=s)


def dashboard(wb):
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    for col, w in zip("ABCDEFGHIJ", (2, 14, 14, 14, 14, 14, 14, 14, 14, 2)):
        ws.column_dimensions[col].width = w

    # title band
    for r in (1, 2, 3):
        for c in range(1, 11):
            ws.cell(r, c).fill = fill(NAVY)
    ws.merge_cells("B2:I2")
    ws["B2"] = "Brew & Co.  ·  2026 Sales Dashboard"
    ws["B2"].font = Font(name="Segoe UI", size=20, bold=True, color="FFFFFF")
    ws["B2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 36
    ws.merge_cells("B3:I3")
    ws["B3"] = "Coffee roastery  ·  four regions  ·  figures in USD"
    ws["B3"].font = Font(name="Segoe UI", size=10, color="9CA3AF")

    # KPI cards: label, big value, small note
    cards = [("Revenue", "=SUM(C13:C24)", '"$"#,##0', TEAL, "+12% vs 2025"),
             ("Profit", "=SUM(E13:E24)", '"$"#,##0', VIOLET, "best month: Dec"),
             ("Margin", "=E25/C25", "0.0%", AMBER, "target 35%"),
             ("Orders", "=COUNTA(Orders!A2:A41)", "#,##0", ROSE, "40 this quarter")]
    for k, (label, value, fmt, color, note) in enumerate(cards):
        c1 = 2 + k * 2
        a, b = get_column_letter(c1), get_column_letter(c1 + 1)
        for r in (5, 6, 7, 8):
            for c in (c1, c1 + 1):
                ws.cell(r, c).fill = fill(SOFT[color])
        ws.cell(5, c1).fill = fill(color)
        ws.cell(5, c1 + 1).fill = fill(color)
        ws.merge_cells(f"{a}5:{b}5")
        ws[f"{a}5"] = label.upper()
        ws[f"{a}5"].font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        ws[f"{a}5"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"{a}6:{b}7")
        ws[f"{a}6"] = value
        ws[f"{a}6"].number_format = fmt
        ws[f"{a}6"].font = Font(name="Segoe UI", size=22, bold=True, color=color)
        ws[f"{a}6"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"{a}8:{b}8")
        ws[f"{a}8"] = note
        ws[f"{a}8"].font = Font(name="Segoe UI", size=9, color=MUTED)
        ws[f"{a}8"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[6].height = 22
    ws.row_dimensions[7].height = 22

    # monthly table
    ws["B11"] = "Monthly performance"
    ws["B11"].font = Font(name="Segoe UI", size=13, bold=True, color=INK)
    heads = ["Month", "Revenue", "Cost", "Profit", "Margin", "Growth", "Target met"]
    for i, h in enumerate(heads):
        cell = ws.cell(12, 2 + i, h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = fill(INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[12].height = 22
    random.seed(7)
    revenue = 41000
    for m, name in enumerate(MONTHS):
        r = 13 + m
        revenue = round(revenue * random.uniform(0.96, 1.09) + (2500 if name in ("Nov", "Dec") else 0), -2)
        cost = round(revenue * random.uniform(0.58, 0.7), -2)
        ws.cell(r, 2, name)
        ws.cell(r, 3, revenue)
        ws.cell(r, 4, cost)
        ws.cell(r, 5, f"=C{r}-D{r}")
        ws.cell(r, 6, f"=E{r}/C{r}")
        ws.cell(r, 7, "" if m == 0 else f"=C{r}/C{r - 1}-1")
        ws.cell(r, 8, f'=IF(F{r}>=0.35,"✔ yes","—")')
        band = fill("F9FAFB") if m % 2 else fill("FFFFFF")
        for c in range(2, 9):
            cell = ws.cell(r, c)
            cell.fill = band
            cell.border = Border(bottom=Side(style="thin", color=LINE))
            cell.font = Font(name="Segoe UI", size=10, color=INK)
            cell.alignment = Alignment(horizontal="left" if c == 2 else "center" if c == 8 else "right")
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = '"$"#,##0'
        ws.cell(r, 6).number_format = "0.0%"
        ws.cell(r, 7).number_format = "+0.0%;-0.0%;0.0%"
    total = 25
    ws.cell(total, 2, "Total")
    for c, f in ((3, "=SUM(C13:C24)"), (4, "=SUM(D13:D24)"), (5, "=SUM(E13:E24)"), (6, "=E25/C25")):
        ws.cell(total, c, f)
    for c in range(2, 9):
        cell = ws.cell(total, c)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=INK)
        cell.fill = fill("F3F4F6")
        cell.border = Border(top=Side(style="medium", color=INK), bottom=Side(style="double", color=INK))
        cell.alignment = Alignment(horizontal="left" if c == 2 else "right")
    for c in (3, 4, 5):
        ws.cell(total, c).number_format = '"$"#,##0'
    ws.cell(total, 6).number_format = "0.0%"

    ws.conditional_formatting.add("F13:F24", ColorScaleRule(start_type="min", start_color="FDE68A",
                                                            mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                                            end_type="max", end_color="99F6E4"))
    ws.conditional_formatting.add("C13:C24", DataBarRule(start_type="min", end_type="max", color="5EEAD4"))
    ws.conditional_formatting.add("G14:G24", CellIsRule(operator="lessThan", formula=["0"],
                                                        font=Font(color=ROSE, bold=True)))
    ws.conditional_formatting.add("G14:G24", CellIsRule(operator="greaterThan", formula=["0"],
                                                        font=Font(color=TEAL, bold=True)))
    ws.conditional_formatting.add("H13:H24", FormulaRule(formula=['LEFT(H13,1)="✔"'],
                                                         font=Font(color="047857", bold=True),
                                                         fill=PatternFill("solid", fgColor="D1FAE5")))

    # region split (for the pie)
    ws["B28"] = "Revenue by region"
    ws["B28"].font = Font(name="Segoe UI", size=13, bold=True, color=INK)
    for i, (region, share) in enumerate((("North", 0.31), ("South", 0.24), ("East", 0.27), ("West", 0.18))):
        r = 29 + i
        ws.cell(r, 2, region).font = Font(name="Segoe UI", size=10, color=INK)
        ws.cell(r, 3, f"=ROUND(C25*{share},0)").number_format = '"$"#,##0'
        ws.cell(r, 3).font = Font(name="Segoe UI", size=10, color=INK)
        ws.cell(r, 4, f"=C{r}/C25").number_format = "0%"
        ws.cell(r, 4).font = Font(name="Segoe UI", size=10, color=MUTED)
    ws.conditional_formatting.add("D29:D32", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                         end_value=0.4, color="C4B5FD"))

    # charts
    col = BarChart()
    col.type, col.title, col.style = "col", "Revenue vs cost", 10
    col.add_data(Reference(ws, min_col=3, max_col=4, min_row=12, max_row=24), titles_from_data=True)
    col.set_categories(Reference(ws, min_col=2, min_row=13, max_row=24))
    col.height, col.width = 7.5, 14
    ws.add_chart(col, "K5")
    line = LineChart()
    line.title = "Profit"
    line.add_data(Reference(ws, min_col=5, min_row=12, max_row=24), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=2, min_row=13, max_row=24))
    line.height, line.width = 7.5, 14
    ws.add_chart(line, "K21")
    pie = PieChart()
    pie.title = "By region"
    pie.add_data(Reference(ws, min_col=3, min_row=29, max_row=32))
    pie.set_categories(Reference(ws, min_col=2, min_row=29, max_row=32))
    pie.height, pie.width = 7, 9
    ws.add_chart(pie, "F27")
    ws.freeze_panes = "A4"


def orders(wb):
    ws = wb.create_sheet("Orders")
    ws.sheet_properties.tabColor = VIOLET
    heads = ["Order", "Date", "Customer", "Product", "Region", "Qty", "Price", "Total", "Status"]
    ws.append(heads)
    random.seed(11)
    products = [("Ethiopia Yirgacheffe", 18.5), ("Colombia Huila", 15.0), ("Sumatra Mandheling", 16.75),
                ("House Espresso", 13.9), ("Decaf Swiss Water", 14.5), ("Cold Brew Pack", 22.0)]
    customers = ["Ana Lima", "Ben Ortiz", "Chen Wei", "Dara Novak", "Eli Haddad", "Femi Ade", "Gita Rao",
                 "Hugo Brandt", "Iris Kowal", "Jae Park"]
    statuses = ["Delivered", "Delivered", "Delivered", "Shipped", "Processing", "Cancelled"]
    day = dt.date(2026, 7, 1)
    for i in range(40):
        day += dt.timedelta(days=random.choice((1, 2, 2, 3)))
        name, price = random.choice(products)
        r = i + 2
        ws.append([f"BC-{1041 + i}", day, random.choice(customers), name, random.choice(["North", "South", "East",
                                                                                          "West"]),
                   random.randint(1, 12), price, f"=F{r}*G{r}", random.choice(statuses)])
        ws.cell(r, 2).number_format = "dd mmm yyyy"
        ws.cell(r, 7).number_format = '"$"#,##0.00'
        ws.cell(r, 8).number_format = '"$"#,##0.00'
    for col, w in zip("ABCDEFGHI", (10, 13, 14, 22, 9, 6, 9, 11, 12)):
        ws.column_dimensions[col].width = w
    table = Table(displayName="Orders", ref="A1:I41")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
    ws.add_table(table)
    for status, (fg, bg) in {"Delivered": ("047857", "D1FAE5"), "Shipped": ("1D4ED8", "DBEAFE"),
                             "Processing": ("B45309", "FEF3C7"), "Cancelled": ("BE123C", "FFE4E6")}.items():
        ws.conditional_formatting.add("I2:I41", CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                           font=Font(color=fg, bold=True), fill=fill(bg)))
    ws.conditional_formatting.add("H2:H41", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                        color="A78BFA"))
    ws.freeze_panes = "A2"


def roadmap(wb):
    ws = wb.create_sheet("Roadmap")
    ws.sheet_properties.tabColor = AMBER
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "Q3 roadmap"
    ws["A1"].font = Font(name="Segoe UI", size=16, bold=True, color=INK)
    ws.row_dimensions[1].height = 28
    heads = ["Task", "Owner", "Due", "Status", "Progress", "Notes"]
    for i, h in enumerate(heads):
        cell = ws.cell(3, 1 + i, h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=MUTED)
        cell.border = Border(bottom=Side(style="medium", color=INK))
    tasks = [("New espresso blend", "Ana", 0.9, "In progress", "Tasting panel Friday"),
             ("Online shop redesign", "Chen", 1.0, "Done", "Live since Aug 3"),
             ("Wholesale price list", "Hugo", 0.45, "In progress", "Waiting for freight quotes"),
             ("Cold brew cans", "Gita", 0.15, "Planned", "Supplier shortlist ready"),
             ("Barista training videos", "Femi", 0.6, "In progress", "3 of 5 recorded"),
             ("Loyalty card app", "Jae", 0.0, "Blocked", "Needs payment provider"),
             ("Autumn market stall", "Iris", 1.0, "Done", "Booked for Oct 12")]
    for i, (task, owner, pct, status, note) in enumerate(tasks):
        r = 4 + i
        ws.cell(r, 1, task).font = Font(name="Segoe UI", size=10, bold=True, color=INK)
        ws.cell(r, 2, owner).font = Font(name="Segoe UI", size=10, color=INK)
        ws.cell(r, 3, dt.date(2026, 8, 15) + dt.timedelta(days=9 * i)).number_format = "dd mmm"
        ws.cell(r, 4, status).alignment = Alignment(horizontal="center")
        ws.cell(r, 5, pct).number_format = "0%"
        ws.cell(r, 6, note).font = Font(name="Segoe UI", size=10, italic=True, color=MUTED)
        for c in range(1, 7):
            ws.cell(r, c).border = Border(bottom=Side(style="thin", color=LINE))
        ws.row_dimensions[r].height = 20
    for status, (fg, bg) in {"Done": ("047857", "D1FAE5"), "In progress": ("1D4ED8", "DBEAFE"),
                             "Planned": ("6B7280", "F3F4F6"), "Blocked": ("BE123C", "FFE4E6")}.items():
        ws.conditional_formatting.add("D4:D10", CellIsRule(operator="equal", formula=[f'"{status}"'],
                                                           font=Font(color=fg, bold=True), fill=fill(bg)))
    ws.conditional_formatting.add("E4:E10", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                        end_value=1, color="FBBF24"))
    for col, w in zip("ABCDEF", (26, 10, 10, 13, 12, 28)):
        ws.column_dimensions[col].width = w


def build(path):
    wb = Workbook()
    dashboard(wb)
    orders(wb)
    roadmap(wb)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    print(build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "docs", "showcase.xlsx")))
